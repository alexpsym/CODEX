"""Read once, classify offline, and apply only proven Counter Trend overrides.
Private captures/ledgers must live outside journal and Git. No archive writes.
"""
from __future__ import annotations
import argparse,hashlib,io,json,os,re,time,zipfile
from pathlib import Path
from urllib.parse import unquote,urlsplit
import xml.etree.ElementTree as ET

NOTE_EXT={'.txt','.md','.markdown','.json','.yaml','.yml','.csv','.xml','.html','.htm','.docx','.pdf','.rtf','.doc','.odt','.log'}
TRADE=re.compile(r'(?i)^([CF])\s*0*(\d+)(?=$|[ _.-])')
MARKER=re.compile(r'(?i)^counter[ _-]*trend$')

def normalized_folder(value,base=None):
    value=unquote(str(value or '')).replace('\\','/')
    if value.lower().startswith('file:'):value=urlsplit(value).path.lstrip('/')
    if base and value and not re.match(r'^[A-Za-z]:/',value) and not value.startswith('/'):
        value=os.path.abspath(os.path.join(str(base),value)).replace('\\','/')
    return os.path.normpath(value).replace('\\','/').rstrip('/').casefold()

def folder_identity(path,root):
    rel=Path(path).relative_to(root);parts=rel.parts;year=next((int(x) for x in parts if re.fullmatch(r'(19|20)\d{2}',x)),None)
    for i in range(len(parts)-1,-1,-1):
        match=TRADE.match(parts[i])
        if match:
            return {'folder':str(Path(root).joinpath(*parts[:i+1])),'trade_number':match[1].upper()+str(int(match[2])),'year':year,'folder_label':parts[i]}
    return None

def extract_xlsx(data):
    from openpyxl import load_workbook
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if sum(i.file_size for i in z.infolist())>100*1024**2:raise ValueError('oversized expanded spreadsheet')
    wb=load_workbook(io.BytesIO(data),data_only=False,keep_links=False);text=[];facts=[]
    try:
        for ws in wb:
            headers={}
            for row in ws:
                literals={cell.column:cell for cell in row if cell.value is not None and cell.data_type!='f'}
                normalized={c:str(cell.value).strip().casefold().replace('_',' ') for c,cell in literals.items()}
                if any(v in {'trade number','trade','trade id','row id'} for v in normalized.values()):headers={v:c for c,v in normalized.items()}
                trade_col=next((headers[k] for k in ('trade number','trade','trade id') if k in headers),None)
                trade=str(literals[trade_col].value).strip().upper() if trade_col in literals else None
                if trade and not re.fullmatch(r'[CF]0*\d+',trade):trade=None
                stable=str(literals[headers['row id']].value).strip() if 'row id' in headers and headers['row id'] in literals else None
                counter_col=headers.get('counter trend');statements=[]
                if counter_col in literals and str(literals[counter_col].value).strip().casefold() in {'yes','no'}:statements.append((literals[counter_col].coordinate,'Counter Trend: '+str(literals[counter_col].value)))
                for c,cell in literals.items():
                    text.append(f'{ws.title}!{cell.coordinate}: {cell.value}')
                    if normalized[c]=='counter trend' and c+1 in literals and str(literals[c+1].value).strip().casefold() in {'yes','no'}:statements.append((literals[c+1].coordinate,'Counter Trend: '+str(literals[c+1].value)))
                    elif re.search(r'(?i)counter[ -]*trend',str(cell.value)) and normalized[c]!='counter trend':statements.append((cell.coordinate,str(cell.value)))
                for cell in row:
                    if cell.comment:
                        text.append(f'{ws.title}!{cell.coordinate} comment: {cell.comment.text}');statements.append((cell.coordinate+'#comment',cell.comment.text))
                for coordinate,statement in statements:facts.append({'trade_number':trade,'stable_id':stable,'sheet':ws.title,'cell':coordinate,'text':statement})
        return {'text':'\n'.join(text),'xlsx_facts':facts,'xlsx_trade_numbers':sorted({f['trade_number'] for f in facts if f['trade_number']})}
    finally:wb.close()


def extract_text(data,suffix):
    if suffix=='.xlsx':return extract_xlsx(data)['text']
    if suffix in {'.docx','.odt'}:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            name='word/document.xml' if suffix=='.docx' else 'content.xml'
            info=z.getinfo(name)
            if info.file_size>20*1024**2:raise ValueError('oversized expanded XML')
            root=ET.fromstring(z.read(name))
            return '\n'.join(' '.join(n.itertext()) for n in root.iter() if n.tag.rsplit('}',1)[-1] in {'p','h'})
    if suffix=='.pdf':
        try:from pypdf import PdfReader
        except ImportError:
            from pdfminer.high_level import extract_text as pdf_text
            return pdf_text(io.BytesIO(data))
        reader=PdfReader(io.BytesIO(data));return '\n'.join('[page %d]\n%s'%(i+1,page.extract_text() or '') for i,page in enumerate(reader.pages))
    if suffix=='.doc':raise ValueError('unsupported binary Word; use existing text/DOCX companion')
    text=None
    for codec in ('utf-8-sig','utf-16','cp1252'):
        try:text=data.decode(codec);break
        except UnicodeError:continue
    if text is None:raise ValueError('undecodable text')
    if suffix=='.rtf':
        try:
            from striprtf.striprtf import rtf_to_text
            return rtf_to_text(text)
        except ImportError:
            # Conservative RTF subset: Unicode/hex and paragraph controls. Never execute objects.
            text=re.sub(r"\\'([0-9a-fA-F]{2})",lambda m:bytes.fromhex(m[1]).decode('cp1252'),text)
            text=re.sub(r'\\u(-?\d+)\??',lambda m:chr(int(m[1])%65536),text)
            text=re.sub(r'\\(?:par|line)\b','\n',text)
            text=re.sub(r'\\[a-zA-Z]+-?\d* ?','',text).replace('{','').replace('}','')
    return text

def capture_archive(roots,cache,*,seconds=290,max_entries=100000,max_notes=10000,max_text=100*1024**2,max_doc=20*1024**2):
    cache=Path(cache);cache.mkdir(parents=True,exist_ok=True);(cache/'bytes').mkdir(exist_ok=True)
    state={'complete':True,'entries':0,'notes':0,'text_bytes':0,'roots':[str(p) for p in roots],'errors':[]};start=time.monotonic()
    inventory=cache/'inventory.jsonl';notes=cache/'notes.jsonl'
    if inventory.exists() or notes.exists():raise ValueError('Capture already exists; analyze the saved input, do not rescan')
    def check():
        if time.monotonic()-start>seconds or state['entries']>=max_entries or state['notes']>=max_notes or state['text_bytes']>=max_text:raise RuntimeError('Capture bound reached')
    def walk_error(exc):state['complete']=False;state['errors'].append(str(exc))
    try:
        with inventory.open('w',encoding='utf-8') as inv,notes.open('w',encoding='utf-8') as out:
            for root in map(Path,roots):
                if not root.is_dir():walk_error(FileNotFoundError(str(root)));continue
                for parent,dirs,files in os.walk(root,followlinks=False,onerror=walk_error):
                    for name in list(dirs):
                        check();state['entries']+=1;path=Path(parent)/name
                        inv.write(json.dumps({'path':str(path),'kind':'directory'})+'\n')
                        if path.is_symlink():dirs.remove(name);walk_error(RuntimeError('symlink skipped: '+str(path)))
                    for name in files:
                        check();state['entries']+=1;path=Path(parent)/name;identity=folder_identity(path.parent,root)
                        inv.write(json.dumps({'path':str(path),'kind':'file','trade':identity})+'\n')
                        if not identity:continue
                        suffix=path.suffix.lower()
                        if suffix not in NOTE_EXT and not MARKER.fullmatch(path.stem):continue
                        state['notes']+=1;record=dict(identity,path=str(path),root=str(root),filename=name,status='captured')
                        try:
                            if path.stat().st_size>max_doc:raise ValueError('oversized document')
                            data=path.read_bytes() # The sole original-content read.
                            digest=hashlib.sha256(data).hexdigest();(cache/'bytes'/digest).write_bytes(data);record.update(sha256=digest,bytes=len(data))
                            text=extract_text(data,suffix);size=len(text.encode('utf-8'))
                            if state['text_bytes']+size>max_text:raise RuntimeError('Extracted-text bound reached')
                            state['text_bytes']+=size;record['text']=text
                        except Exception as exc:record.update(status='unreadable_or_unsupported',error=str(exc));state['complete']=False
                        out.write(json.dumps(record,ensure_ascii=False)+'\n');out.flush();inv.flush()
                        (cache/'coverage.json').write_text(json.dumps(state,indent=2))
    except Exception as exc:state['complete']=False;state['errors'].append(str(exc))
    finally:state['elapsed']=time.monotonic()-start;(cache/'coverage.json').write_text(json.dumps(state,indent=2))
    return state

def capture_supplement(paths,root,cache,original_cache,*,seconds=285):
    cache=Path(cache);cache.mkdir(exist_ok=True);(cache/'bytes').mkdir(exist_ok=True)
    original_cache=Path(original_cache);seen={json.loads(line)['path'] for line in (original_cache/'notes.jsonl').read_text(encoding='utf-8').splitlines()};state={'complete':True,'entries':0,'notes':0,'text_bytes':0,'errors':[]};start=time.monotonic()
    if (cache/'notes.jsonl').exists():raise ValueError('Supplement already captured; never reread originals')
    def error(exc):state['complete']=False;state['errors'].append(str(exc))
    def selected():
        for entry in paths:yield entry['path'],entry.get('trade'),entry.get('root','inventoried root')
        if not Path(root).is_dir():error(FileNotFoundError(str(root)));return
        for parent,dirs,files in os.walk(root,followlinks=False,onerror=error):
            for name in list(dirs):
                state['entries']+=1
                if (Path(parent)/name).is_symlink():dirs.remove(name);error(RuntimeError('symlink skipped: '+str(Path(parent)/name)))
            for name in files:
                path=Path(parent)/name;state['entries']+=1;identity=folder_identity(path.parent,root)
                if identity and path.suffix.lower() in NOTE_EXT|{'.xlsx'}:yield str(path),identity,str(root)
    try:
        with (cache/'notes.jsonl').open('w',encoding='utf-8') as out:
            for path,identity,origin in selected():
                if time.monotonic()-start>seconds or state['entries']>20000 or state['notes']>=2531 or state['text_bytes']>=50*1024**2:raise RuntimeError('Supplemental capture bound reached')
                if path in seen:continue
                seen.add(path);state['notes']+=1;r=dict(identity or {},path=path,root=origin,filename=Path(path).name,status='captured')
                try:
                    file=Path(path)
                    if file.stat().st_size>20*1024**2:raise ValueError('oversized document')
                    data=file.read_bytes();digest=hashlib.sha256(data).hexdigest();(cache/'bytes'/digest).write_bytes(data);r.update(sha256=digest,bytes=len(data))
                    if file.suffix.lower()=='.xlsx':r.update(extract_xlsx(data))
                    else:r['text']=extract_text(data,file.suffix.lower())
                    size=len(r['text'].encode('utf-8'))
                    if state['text_bytes']+size>50*1024**2:raise RuntimeError('Supplemental text bound reached')
                    state['text_bytes']+=size
                except Exception as exc:r.update(status='unreadable_or_unsupported',error=str(exc));state['complete']=False
                out.write(json.dumps(r,ensure_ascii=False)+'\n');out.flush();(cache/'coverage.json').write_text(json.dumps(state,indent=2))
    except Exception as exc:error(exc)
    finally:state['elapsed']=time.monotonic()-start;(cache/'coverage.json').write_text(json.dumps(state,indent=2))
    return state


def classify_note(record):
    evidence=[];filename=record.get('filename','');stem=Path(filename).stem
    symbol=record.get('context_symbol','')
    named=bool(symbol and re.fullmatch(r'(?i)counter[ _-]*trend[ _-]+'+re.escape(symbol),stem))
    if re.search(r'(?i)template|example|lesson',filename):return None,[],'generic template/discussion'
    if MARKER.fullmatch(stem) or named:evidence.append({'value':'Yes','excerpt':filename,'line':0,'reason':'trade-specific marker filename'})
    for n,line in enumerate(record.get('text','').splitlines(),1):
        normalized=re.sub(r'[-_]+',' ',line).strip().lstrip('\u2022\ufffd\t *')
        if not re.search(r'(?i)counter\s*trend',normalized):continue
        if re.search(r'(?i)\b(template|example|lesson|could|might|maybe|perhaps|possible|consider|if|would|question)\b|\?|^[>\"?]',normalized):continue
        no=re.search(r'(?i)^(?:counter\s*trend\s*[:=]\s*no\b|(?:this|the)\s+trade\s+(?:(?:was|is)\s+)?(?:not|isn.t|wasn.t)\s+(?:a\s+)?counter\s*trend)',normalized)
        yes=re.search(r'(?i)^(?:counter\s*trend\s*[:=]\s*yes\b|(?:this|the)\s+trade\s+(?:was|is)\s+(?:a\s+)?counter\s*trend\b|(?:this|the)\s+(?:was|is)\s+(?:a\s+)?counter\s*trend\s+trade\b|counter\s*trend(?:\s+trade(?:\s+on\b.*)?|[.!]?)$)',normalized)
        if not yes and re.fullmatch(r'(?i)counter\s*trend',normalized):yes=True
        if no or yes:evidence.append({'value':'No' if no else 'Yes','excerpt':line[:1200],'line':n,'reason':'explicit trade statement'})
    values={e['value'] for e in evidence}
    return (next(iter(values)) if len(values)==1 else None),evidence,('conflicting evidence' if len(values)>1 else 'clear' if values else 'no explicit classification')

def match_folder(record,rows,workbook_dir):
    exact=[r for r in rows if r.get('link') and normalized_folder(r['link'],workbook_dir)==normalized_folder(record['folder'])]
    if exact:return (exact[0] if len(exact)==1 else None),'exact hyperlink' if len(exact)==1 else 'ambiguous exact hyperlinks'
    number=record['trade_number'];year=record.get('year');label=record.get('folder_label','').upper()
    candidates=[r for r in rows if str(r.get('trade_number') or '').upper()==number and str(year) in str(r.get('open_time') or '')[:4]]
    # Number/year alone is insufficient: symbol must be explicit in folder context.
    candidates=[r for r in candidates if re.sub('[^A-Z0-9]','',str(r.get('symbol') or '').upper()) in re.sub('[^A-Z0-9]','',label) and r.get('symbol')]
    return (candidates[0] if len(candidates)==1 else None),'unique number/year/symbol' if len(candidates)==1 else 'unmatched or ambiguous number/year/symbol/account'

def analyze_capture(cache,rows,workbook_dir):
    cache=Path(cache);ledger=[];groups={}
    note_path=cache/'parsed_notes.jsonl' if (cache/'parsed_notes.jsonl').exists() else cache/'notes.jsonl'
    for line in note_path.read_text(encoding='utf-8').splitlines():
        record=json.loads(line);row,match=match_folder(record,rows,workbook_dir)
        if 'xlsx_facts' in record:
            evidence=[]
            for fact in record['xlsx_facts']:
                if not row:continue
                if fact.get('stable_id') and fact['stable_id']!=str(row['id']):continue
                if fact.get('trade_number') and fact['trade_number']!=str(row.get('trade_number') or '').upper():continue
                if len(record.get('xlsx_trade_numbers',[]))>1 and not fact.get('trade_number') and not fact.get('stable_id'):continue
                value,items,_=classify_note(dict(record,filename=record['filename'],text=fact['text']))
                evidence.extend(dict(item,sheet=fact['sheet'],cell=fact['cell']) for item in items)
            values={item['value'] for item in evidence};value=next(iter(values)) if len(values)==1 else None
            reason='conflicting evidence' if len(values)>1 else 'clear' if values else 'no bound explicit spreadsheet classification'
        else:
            value,evidence,reason=classify_note(dict(record,context_symbol=row.get('symbol','') if row else ''))
        entry={k:v for k,v in record.items() if k!='text'};entry.update(classification=value,evidence=evidence,reason=reason,match_reason=match,row_id=row.get('id') if row else None,original=row.get('counter_trend') if row else None)
        ledger.append(entry)
        if row:
            groups.setdefault(row['id'],[]).append(entry)
    changes=[]
    for rid,entries in groups.items():
        values={e['classification'] for e in entries if e['classification']};conflict=any(e['reason']=='conflicting evidence' for e in entries) or len(values)>1
        current=next(r for r in rows if r['id']==rid).get('counter_trend') or ''
        for e in entries:
            e['proposed']=next(iter(values)) if len(values)==1 and not conflict else None
            e['decision']='conflict' if conflict or current and values and current not in values else 'preserve explicit' if current else 'apply' if values else 'unknown'
        if not current and len(values)==1 and not conflict:changes.append({'id':rid,'counter_trend':next(iter(values)),'expected_counter_trend':''})
    report={'coverage':json.loads((cache/'coverage.json').read_text()),'inspected_notes':len(ledger),'changes':changes,'ledger':ledger,'c398':[e for e in ledger if e['trade_number']=='C398']}
    (cache/'evidence_ledger.json').write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf-8');return report

def prepare_backfill(workbook,candidate,changes):
    from tools import master_journal_workbook as m
    from openpyxl import load_workbook
    p=Path(workbook);wb=load_workbook(p,data_only=False)
    try:
        ws=wb['Trade Log'];h=m._trade_log_header_map(ws);current={str(ws.cell(r,h['Row ID']).value):m.normalize_counter_trend(ws.cell(r,h['Counter Trend']).value) if h.get('Counter Trend') else '' for r in range(m._trade_log_data_start_row(ws),ws.max_row+1)}
        for change in changes:
            if change['id'] not in current or current[change['id']]!=change.get('expected_counter_trend',''):raise RuntimeError('Classification changed or identity absent: '+change['id'])
    finally:wb.close()
    return m.prepare_counter_trend_package(p,Path(candidate),expected_sha256=hashlib.sha256(p.read_bytes()).hexdigest(),rows=changes)

def main():
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--capture',action='store_true');ap.add_argument('--root',action='append',default=[]);ap.add_argument('--cache',required=True);ap.add_argument('--rows');ap.add_argument('--workbook');ap.add_argument('--candidate');ap.add_argument('--apply',action='store_true');args=ap.parse_args()
    if args.capture:print(json.dumps(capture_archive(args.root,args.cache)));return
    rows=json.loads(Path(args.rows).read_text());report=analyze_capture(args.cache,rows,Path(args.workbook).parent)
    if args.apply:
        if not args.candidate:ap.error('--apply requires --candidate; canonical installer owns final replacement')
        report['candidate']=prepare_backfill(args.workbook,args.candidate,report['changes'])
    print(json.dumps({k:v for k,v in report.items() if k not in {'ledger','c398'}}))
if __name__=='__main__':main()
