"""Conservative, explicit local filing. No scans/moves occur on import.
The UI supplies a request identity only. Paths are fixed by the local adapter.
Audit files are private runtime state, never journal metadata.
"""
from __future__ import annotations
import argparse, contextlib, ctypes, hashlib, io, json, os, re, stat, subprocess, sys, time, uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import unquote, urlsplit

SOURCE = Path(r'C:\Users\User\Downloads\CHARTS')
DROPBOX = Path(r'C:\Users\User\Dropbox')
BRISBANE = timezone(timedelta(hours=10))
MAX_FILES, MAX_BYTES, MAX_DOCUMENT = 1000, 1024**3, 20*1024**2
REPARSE = 0x400
FINISHED = {'complete', 'partial', 'failed'}

class Unsafe(ValueError): pass
class BatchLimit(RuntimeError): pass

def check_deadline(deadline):
    if time.monotonic() >= deadline: raise BatchLimit('Batch time limit reached; remaining files retained')

def safe_path(path, root, missing=False):
    path, root = Path(os.path.abspath(path)), Path(os.path.abspath(root))
    if not path.is_relative_to(root): raise Unsafe('Path is outside the approved root')
    for part in [root, *[root.joinpath(*path.relative_to(root).parts[:i]) for i in range(1,len(path.relative_to(root).parts)+1)]]:
        try: info=part.lstat()
        except FileNotFoundError:
            if missing and part==path: continue
            raise
        if stat.S_ISLNK(info.st_mode) or getattr(info,'st_file_attributes',0)&REPARSE: raise Unsafe('Symlink/junction/reparse path refused')
    # Ancestor junctions above the root must not redirect the approved namespace.
    for part in root.parents:
        info=part.lstat()
        if stat.S_ISLNK(info.st_mode) or getattr(info,'st_file_attributes',0)&REPARSE: raise Unsafe('Approved-root ancestor is a reparse path')
    if not path.resolve(strict=not missing).is_relative_to(root.resolve()): raise Unsafe('Resolved path escapes approved root')
    return path

def identity(info): return (info.st_dev, info.st_ino, info.st_size, info.st_mtime_ns)

def number(value):
    match=re.fullmatch(r'([CF])\s*0*(\d+)',str(value or '').strip(),re.I)
    return match[1].upper()+str(int(match[2])) if match else ''

def symbol(value):
    text=str(value or '').upper().strip().split(':')[-1]
    text=re.sub(r'\.P$|PERPETUAL$|PERP$','',text)
    return re.sub(r'[/_\-\s]','',text)

def account(value): return re.sub(r'[^A-Z0-9]','',str(value or '').upper())

def parse_time(value):
    text=str(value or '').strip()
    try:
        result=datetime.fromisoformat(text.replace('Z','+00:00'))
    except ValueError: return None,False
    unzoned=result.tzinfo is None
    return (result.replace(tzinfo=BRISBANE) if unzoned else result).astimezone(timezone.utc),unzoned

def folder_from_link(value,workbook):
    value=unquote(str(value or ''))
    if value.startswith('file:'): value=urlsplit(value).path.lstrip('/')
    path=Path(value)
    return Path(os.path.abspath(path if path.is_absolute() else Path(workbook).parent/path))

def filename_evidence(name):
    stem=Path(name).stem.replace('_',' '); ids=set()
    first=re.match(r'^([CF]\d+)(?=$|[ _.-])',stem,re.I)
    if first:ids.add(number(first[1]))
    ids.update(number(x) for x in re.findall(r'\btrade[ _-]*(?:id|number|no)?[ _:=#-]*([CF]\d+)\b',stem,re.I))
    exchange=re.search(r'\b(BINANCE|BYBIT|OANDA|PEPPERSTONE|COINSPOT)[_: -]+([A-Z0-9/.-]+)',stem,re.I)
    sym=exchange[2] if exchange and exchange[2].upper() not in {'DEMO','LIVE'} else ''
    if not sym:
        token=re.search(r'(?<![A-Z0-9])([A-Z0-9]+(?:USDT|USDC)(?:\.P)?|(?:EUR|USD|GBP|AUD|NZD|CAD|CHF|JPY|XAU|XAG)[/_-]?(?:USD|JPY|GBP|AUD|NZD|CAD|CHF|EUR))(?![A-Z0-9])',stem,re.I)
        if token:sym=token[1]
    acc=re.search(r'\b(BYBIT|OANDA|PEPPERSTONE|COINSPOT)[ _-]+(DEMO|LIVE)\b',stem,re.I)
    capture=re.search(r'(20\d{2}-\d{2}-\d{2})[_ T](\d{2})[-:](\d{2})(?:[-:](\d{2}))?(Z|[+-]\d{2}:?\d{2})?',stem,re.I)
    timestamp=''
    if capture:timestamp=f'{capture[1]}T{capture[2]}:{capture[3]}:{capture[4] or "00"}{capture[5] or ""}'
    return {'ids':ids,'symbol':symbol(sym),'account':account(' '.join(acc.groups())) if acc else '', 'capture_time':timestamp,'entry_time':'','exit_time':'','assumptions':[], 'market':('crypto' if exchange and exchange[1].upper() in {'BINANCE','BYBIT','COINSPOT'} else 'fx' if exchange else '')}

def evidence(name,data=None):
    result=filename_evidence(name);suffix=Path(name).suffix.lower();text='';warnings=[]
    if data is not None:
        try:
            if suffix=='.xlsx':
                from openpyxl import load_workbook
                import zipfile
                with zipfile.ZipFile(io.BytesIO(data)) as z:
                    if sum(x.file_size for x in z.infolist())>80*1024**2:raise Unsafe('Expanded spreadsheet exceeds evidence bound')
                wb=load_workbook(io.BytesIO(data),data_only=False,keep_links=False);lines=[]
                try:
                    count=0
                    for ws in wb:
                        header={}
                        for row in ws:
                            values={c.column:str(c.value).strip() for c in row if c.value is not None and c.data_type!='f'};count+=len(values)
                            if count>100000:raise Unsafe('Spreadsheet evidence cell limit reached')
                            labels={c:v.lower().replace('_',' ') for c,v in values.items()}
                            known={'trade number','trade id','symbol','account','entry time','exit time','capture time','timezone'}
                            lines.extend(c.comment.text for c in row if c.comment)
                            is_header=any(v in {'trade number','trade id'} for v in labels.values()) and sum(v in known for v in labels.values())>1
                            if is_header:
                                header={c:v for c,v in labels.items() if v in known};continue
                            table_row=any(header.get(c) in {'trade number','trade id'} and re.fullmatch(r'[CF]\s*\d+',v,re.I) for c,v in values.items())
                            for c,v in values.items():
                                if table_row and c in header:lines.append(header[c]+': '+v)
                                if v.lower() in known and c+1 in values and values[c+1].lower() not in known:lines.append(v+': '+values[c+1])
                                if ':' in v:lines.append(v)
                    text='\n'.join(lines)
                finally:wb.close()
            elif suffix in {'.txt','.md','.json','.csv','.docx','.pdf','.rtf'}:
                from tools.backfill_counter_trend import extract_text
                text=extract_text(data,suffix)
                if suffix=='.json':
                    obj=json.loads(text)
                    if isinstance(obj,dict):text='\n'.join(str(k).replace('_',' ')+': '+str(v) for k,v in obj.items() if not isinstance(v,(dict,list)))
            elif suffix=='.png':
                from PIL import Image
                with Image.open(io.BytesIO(data)) as image:text='\n'.join(str(v) for v in image.info.values() if isinstance(v,str))
            if len(text)>1024**2:raise Unsafe('Extracted evidence exceeds text bound')
        except Exception as exc:
            # Failed parsing is not proof of identity. Filename evidence can still stand.
            warnings.append('Embedded evidence unavailable: '+type(exc).__name__);text=''
    lines=text.splitlines()
    lines += [line+': '+lines[i+1] for i,line in enumerate(lines[:-1]) if line.strip().lower() in {'trade number','trade id','symbol','account','entry time','exit time','capture time'}]
    for line in lines:
        match=re.search(r'(?i)\btrade\s*(?:number|id|no\.?)\s*[:=#]\s*([CF]\s*\d+)\b',line)
        if match:result['ids'].add(number(match[1]))
        for key,pattern in [('symbol',r'(?:symbol|instrument|pair)'),('account',r'account'),('entry_time',r'(?:entry|open) time'),('exit_time',r'(?:exit|close) time'),('capture_time',r'(?:capture|export) time')]:
            found=re.match(r'\s*'+pattern+r'\s*[:=]\s*(.+?)\s*$',line,re.I)
            if found:
                value=symbol(found[1]) if key=='symbol' else account(found[1]) if key=='account' else found[1]
                if result.get(key) and result[key]!=value:result.setdefault('contradictions',[]).append(key)
                result[key]=value
    result['warnings']=warnings
    if result['capture_time']:
        _,naive=parse_time(result['capture_time'])
        if naive:result['assumptions'].append('Unzoned capture time would be Australia/Brisbane; not used alone for filing')
    return result

def match_file(ev,candidates):
    if ev.get('contradictions'):return None,'Contradictory filename/document metadata'
    if len(ev['ids'])>1:return None,'Multiple labelled trades; file left in place'
    choices=[]
    for row in candidates:
        if ev['ids'] and row['trade_number'] not in ev['ids']:continue
        if ev['symbol'] and row.get('symbol') and symbol(row.get('symbol'))!=ev['symbol']:continue
        if ev['account'] and row.get('account') and account(row.get('account'))!=ev['account']:continue
        if ev.get('market') and row.get('asset_class') and row['asset_class']!=ev['market']:continue
        opened,_=parse_time(row.get('open_time'));closed,_=parse_time(row.get('close_time'))
        entry,entry_naive=parse_time(ev['entry_time']);exit_time,_=parse_time(ev['exit_time']);captured,capture_naive=parse_time(ev['capture_time'])
        if ev['entry_time'] and (entry is None or opened!=entry):continue
        if ev['exit_time'] and (exit_time is None or closed!=exit_time):continue
        if captured and opened and captured<opened:continue
        if not ev['ids']:
            if not ev['symbol'] or not ev['account'] or symbol(row.get('symbol'))!=ev['symbol'] or account(row.get('account'))!=ev['account']:continue
            exact_entry=entry is not None and opened==entry
            interval=captured is not None and not capture_naive and opened is not None and closed is not None and opened<=captured<=closed
            if not exact_entry and not interval:continue
        choices.append(row)
    # Repeated accounts/years/identities are never silently collapsed to first/last.
    if len(choices)!=1:return None,('No uniquely corroborated trade/folder' if not choices else 'Ambiguous trade/account/year; more context required')
    return choices[0],('Explicit labelled trade and verified existing folder' if ev['ids'] else 'Exact symbol/account and trade time evidence')

def approved_roots():
    from tools.master_journal_workbook import _trade_folder_roots
    roots=[]
    for prefix,market in [('C','CRYPTO'),('F','FOREX')]:
        selected=Path(os.getenv('TRADING_JOURNAL_'+market+'_ROOT') or DROPBOX/market).absolute()
        if not selected.is_relative_to(DROPBOX):raise Unsafe('Configured chart root is not under the approved local Dropbox')
        # The shared helper provides the configured first choice, never its legacy fallbacks.
        selected=_trade_folder_roots(prefix,selected)[0];roots.append(selected)
    return roots

def read_candidates(workbook,roots,deadline):
    from tools import master_journal_workbook as m
    from openpyxl import load_workbook
    before=identity(Path(workbook).stat());items=m.read_master_journal_source(workbook)['items'];wb=load_workbook(workbook,read_only=False,data_only=False,keep_links=False);links={}
    try:
        ws=wb['Trade Log'];h=m._trade_log_header_map(ws)
        for row in ws.iter_rows(min_row=m._trade_log_data_start_row(ws)):
            check_deadline(deadline);cell=row[h['Trade Number']-1];rid=str(row[h['Row ID']-1].value or '')
            links[rid]=cell.hyperlink.target if cell.hyperlink else ''
    finally:wb.close()
    if identity(Path(workbook).stat())!=before:raise Unsafe('Journal changed during read-only snapshot')
    folders={};entries=0
    for root in roots:
        safe_path(root,root)
        for parent,dirs,_ in os.walk(root,followlinks=False):
            check_deadline(deadline)
            for name in list(dirs):
                entries+=1
                if entries>20000:raise BatchLimit('Destination metadata bound reached')
                path=Path(parent)/name
                try:safe_path(path,root)
                except (OSError,Unsafe):dirs.remove(name);continue
                matched=re.match(r'^([CF])\s*0*(\d+)(?=$|[ _.-])',name,re.I)
                if not matched and Path(parent).name.upper() in {'C','F'} and name.isdigit():matched=re.match(r'([CF])(\d+)',Path(parent).name.upper()+name)
                if matched:folders.setdefault(number(''.join(matched.groups())),set()).add(path)
    result=[];associated=set()
    for row in items:
        check_deadline(deadline)
        if row.get('row_type','trade')!='trade':continue
        trade=number(row.get('trade_number'))
        if not trade:continue
        linked=folder_from_link(links.get(str(row.get('id'))),workbook) if links.get(str(row.get('id'))) else None
        possibilities=[]
        if linked:
            for root in roots:
                try:
                    path=safe_path(linked,root)
                    if path.is_dir() and any(number(x)==trade for x in re.findall(r'(?i)([CF]\s*\d+)',path.name)):possibilities=[path];break
                except (OSError,Unsafe):continue
        if not possibilities:possibilities=list(folders.get(trade,()))
        opened,_=parse_time(row.get('open_time'))
        if len(possibilities)>1 and opened:
            possibilities=[p for p in possibilities if str(opened.astimezone(BRISBANE).year) in p.parts]
        for folder in possibilities:
            result.append(dict(row,trade_number=trade,folder=str(folder)));associated.add((trade,str(folder)))
    # Directly labelled files may identify a unique verified folder even if it has
    # no current journal row; no row is created. Context requiring absent facts fails.
    for trade,paths in folders.items():
        for folder in paths:
            if (trade,str(folder)) not in associated:result.append({'id':'folder:'+str(folder),'trade_number':trade,'folder':str(folder)})
    return result

@contextlib.contextmanager
def guarded_open(path,delete=False):
    """Deny concurrent writes/renames while verifying source/destination bytes."""
    if os.name=='nt':
        import msvcrt
        k=ctypes.WinDLL('kernel32',use_last_error=True);k.CreateFileW.argtypes=[ctypes.c_wchar_p,ctypes.c_uint32,ctypes.c_uint32,ctypes.c_void_p,ctypes.c_uint32,ctypes.c_uint32,ctypes.c_void_p];k.CreateFileW.restype=ctypes.c_void_p;k.CloseHandle.argtypes=[ctypes.c_void_p]
        handle=k.CreateFileW(str(path),0x80000000|(0x10000 if delete else 0),1,None,3,0,None)
        if handle==ctypes.c_void_p(-1).value:raise ctypes.WinError(ctypes.get_last_error())
        try:fd=msvcrt.open_osfhandle(handle,os.O_RDONLY|os.O_BINARY)
        except Exception:k.CloseHandle(handle);raise
        stream=os.fdopen(fd,'rb')
        k.GetFinalPathNameByHandleW.argtypes=[ctypes.c_void_p,ctypes.c_wchar_p,ctypes.c_uint32,ctypes.c_uint32];k.GetFinalPathNameByHandleW.restype=ctypes.c_uint32
        name=ctypes.create_unicode_buffer(32768);length=k.GetFinalPathNameByHandleW(handle,name,len(name),0)
        actual=name.value
        if actual.startswith('\\\\?\\'):actual=actual[4:]
        if not length or length>=len(name) or os.path.normcase(actual)!=os.path.normcase(os.path.abspath(path)):
            stream.close();raise Unsafe('Opened file resolves outside its checked path')
    else:
        import fcntl
        fd=os.open(path,os.O_RDONLY|getattr(os,'O_NOFOLLOW',0));stream=os.fdopen(fd,'rb');fcntl.flock(stream,fcntl.LOCK_EX|fcntl.LOCK_NB);handle=None
    try:yield stream,handle
    finally:stream.close()

def remove_source(stream,handle,path):
    if os.name=='nt':
        class Disposition(ctypes.Structure):_fields_=[('DeleteFile',ctypes.c_int)]
        k=ctypes.WinDLL('kernel32',use_last_error=True);k.SetFileInformationByHandle.argtypes=[ctypes.c_void_p,ctypes.c_int,ctypes.c_void_p,ctypes.c_uint32];flag=Disposition(1)
        if not k.SetFileInformationByHandle(handle,4,ctypes.byref(flag),ctypes.sizeof(flag)):raise ctypes.WinError(ctypes.get_last_error())
        stream.close() # The verified, exclusively owned source handle performs deletion.
    else:Path(path).unlink()

def stream_hash(stream,deadline):
    stream.seek(0);digest=hashlib.sha256()
    while True:
        check_deadline(deadline);chunk=stream.read(1024**2)
        if not chunk:break
        digest.update(chunk)
    stream.seek(0);return digest.hexdigest()

def exact_bytes(a,b,deadline):
    a.seek(0);b.seek(0)
    while True:
        check_deadline(deadline);x=a.read(1024**2);y=b.read(1024**2)
        if x!=y:return False
        if not x:return True

def atomic_json(path,value):
    path=Path(path);temp=path.with_name(path.name+'.'+uuid.uuid4().hex+'.tmp')
    with temp.open('x',encoding='utf-8') as f:json.dump(value,f,ensure_ascii=False);f.flush();os.fsync(f.fileno())
    os.replace(temp,path)

@contextlib.contextmanager
def batch_lock(directory):
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True);safe_path(directory,directory);safe_path(directory/'batch.lock',directory,missing=True);handle=(directory/'batch.lock').open('a+b')
    try:
        handle.seek(0,2)
        if handle.tell()==0:handle.write(b'0');handle.flush()
        handle.seek(0)
        if os.name=='nt':
            import msvcrt;msvcrt.locking(handle.fileno(),msvcrt.LK_NBLCK,1)
        else:
            import fcntl;fcntl.flock(handle,fcntl.LOCK_EX|fcntl.LOCK_NB)
        yield
    finally:handle.close()

def discover_files(source,deadline):
    safe_path(source,source);files=[];entries=0
    for parent,dirs,names in os.walk(source,followlinks=False):
        check_deadline(deadline)
        for name in list(dirs):
            entries+=1
            try:safe_path(Path(parent)/name,source)
            except (OSError,Unsafe):dirs.remove(name);files.append((Path(parent)/name,'Reparse directory not traversed'))
        for name in names:
            entries+=1
            if entries>20000:return files,True
            files.append((Path(parent)/name,''))
    return sorted(files,key=lambda x:str(x[0]).casefold()),False

def collision_path(folder,name,digest,deadline):
    def existing(filename):
        found=[]
        for i,path in enumerate(folder.iterdir()):
            check_deadline(deadline)
            if i>=20000:raise BatchLimit('Destination directory bound reached')
            if path.name.casefold()==filename.casefold():found.append(path)
        if len(found)>1:raise Unsafe('Ambiguous case-insensitive destination names')
        return found[0] if found else None
    first=existing(name)
    if first is None:return folder/name,None
    safe_path(first,folder)
    with guarded_open(first) as (f,_):same=stream_hash(f,deadline)==digest
    if same:return first,first
    path=Path(name);renamed=f'{path.stem}-{digest[:12]}{path.suffix}';other=existing(renamed)
    if other is None:return folder/renamed,None
    safe_path(other,folder)
    with guarded_open(other) as (f,_):same=stream_hash(f,deadline)==digest
    if not same:raise Unsafe('Deterministic collision filename already has conflicting contents')
    return other,other

def publish(stage,target):
    if os.name=='nt':os.rename(stage,target) # Windows rename fails if target exists.
    else:os.link(stage,target);stage.unlink() # Atomic no-clobber publication.

def report_counts(report):
    counts={k:0 for k in ('moved','already_filed','unresolved','failed','copied_source_retained','unprocessed')}
    for row in report['files']:counts[row['status']]+=1
    report['counts']=counts
    report['ok']=not any(counts[k] for k in ('unresolved','failed','copied_source_retained','unprocessed')) and not report.get('error')
    return report

def run_batch(request_id,source,roots,audit,candidates=None,loader=None,*,seconds=285,max_files=MAX_FILES,max_bytes=MAX_BYTES,hook=None):
    if not re.fullmatch(r'[a-f0-9-]{32,36}',request_id):raise Unsafe('Invalid request identity')
    source=Path(source).absolute();roots=[Path(p).absolute() for p in roots];audit=Path(audit);deadline=time.monotonic()+seconds;started=time.monotonic()
    try:
        with batch_lock(audit):
            result_path=audit/(request_id+'.json');active=audit/'active.json'
            safe_path(result_path,audit,missing=True);safe_path(active,audit,missing=True)
            if result_path.exists():
                existing=json.loads(result_path.read_text(encoding='utf-8'))
                if existing.get('request_id')!=request_id:raise Unsafe('Audit identity mismatch')
                if existing.get('status') not in FINISHED:
                    report_counts(existing);existing.update(ok=False,status='inconclusive',message='Previous worker outcome requires audit review; no automatic replay.')
                return existing
            if active.exists():return {'ok':False,'status':'inconclusive','message':'Previous batch is unfinished. Review its audit; no automatic retry.', 'request_id':request_id,'files':[],'counts':{k:0 for k in ('moved','already_filed','unresolved','failed','copied_source_retained','unprocessed')}}
            report={'request_id':request_id,'status':'running','ok':False,'files':[],'source':str(source),'transferred_bytes':0}
            atomic_json(active,{'request_id':request_id});atomic_json(result_path,report)
            try:
                files,truncated=discover_files(source,deadline)
                if not files:
                    report['message']='No downloaded files to sort.'
                else:
                    for root in roots:safe_path(root,root)
                    if not roots:raise Unsafe('No approved local Dropbox roots available')
                    if candidates is None:candidates=loader(deadline)
                for index,(path,issue) in enumerate(files):
                    row={'file':str(path.relative_to(source)),'status':'unresolved','reason':issue or ''};report['files'].append(row)
                    if index>=max_files:row.update(status='unprocessed',reason='File count limit reached');continue
                    try:
                        check_deadline(deadline)
                        if issue:continue
                        safe_path(path,source);info=path.stat()
                        if not stat.S_ISREG(info.st_mode):raise Unsafe('Not an ordinary file')
                        if re.search(r'(?i)(\.crdownload|\.part|\.partial|\.download|\.tmp)$',path.name) or path.name.startswith('~$') or time.time()-info.st_mtime<2:
                            row['reason']='Incomplete or still-changing download; retained';continue
                        if info.st_size+report['transferred_bytes']>max_bytes:raise BatchLimit('Transfer byte limit reached')
                        with guarded_open(path,delete=True) as (src,handle):
                            original=identity(os.fstat(src.fileno()))
                            if identity(path.stat())!=original:raise Unsafe('Source identity changed')
                            suffix=path.suffix.lower();data=None
                            if info.st_size<=MAX_DOCUMENT and suffix in {'.txt','.md','.json','.csv','.docx','.xlsx','.pdf','.rtf','.png'}:data=src.read(MAX_DOCUMENT+1);src.seek(0)
                            ev=evidence(path.name,data);target,reason=match_file(ev,candidates or [])
                            row.update(reason=reason,evidence={'trade_numbers':sorted(ev['ids']),'symbol':ev['symbol'],'account':ev['account'],'assumptions':ev['assumptions'],'warnings':ev['warnings']})
                            if target is None:continue
                            folder=Path(target['folder']).absolute();root=next((r for r in roots if folder.is_relative_to(r)),None)
                            if root is None:raise Unsafe('Trade destination is outside approved Dropbox roots')
                            safe_path(folder,root)
                            if not folder.is_dir():raise Unsafe('Trade folder no longer exists')
                            digest=stream_hash(src,deadline);destination,duplicate=collision_path(folder,path.name,digest,deadline);safe_path(destination,root,missing=True)
                            row.update(source_identity=list(original),sha256=digest,destination=str(destination),renamed_for_collision=destination.name.casefold()!=path.name.casefold(),trade_number=target['trade_number'],row_id=target.get('id'),phase='planned');atomic_json(result_path,report)
                            if hook:hook('planned',row,path)
                            if identity(path.stat())!=original:raise Unsafe('Source changed after planning')
                            if duplicate is None:
                                stage=folder/('.ltt-sort-'+request_id+'-'+digest[:16]+'.part');row['staging']=str(stage);safe_path(stage,root,missing=True);atomic_json(result_path,report)
                                with stage.open('xb') as out:
                                    src.seek(0)
                                    while True:
                                        check_deadline(deadline);chunk=src.read(1024**2)
                                        if not chunk:break
                                        out.write(chunk);report['transferred_bytes']+=len(chunk)
                                        if report['transferred_bytes']>max_bytes:raise BatchLimit('Transfer byte limit reached')
                                    out.flush();os.fsync(out.fileno())
                                os.utime(stage,ns=(info.st_atime_ns,info.st_mtime_ns));row['staging_identity']=list(identity(stage.stat()))
                                with guarded_open(stage) as (copy,_):
                                    if stream_hash(copy,deadline)!=digest or os.fstat(copy.fileno()).st_size!=info.st_size:raise OSError('Staging size/hash verification failed')
                                row.update(phase='staged',copy_verified=True);atomic_json(result_path,report)
                                if hook:hook('copied',row,path)
                                safe_path(stage,root);safe_path(destination,root,missing=True)
                                if identity(stage.stat())!=tuple(row['staging_identity']):raise Unsafe('Owned staging identity changed')
                                if collision_path(folder,destination.name,digest,deadline)[1] is not None:raise FileExistsError('Concurrent destination publication; source retained')
                                publish(stage,destination);row['phase']='copied';row['status']='copied_source_retained';atomic_json(result_path,report)
                            else:row.update(phase='copied',status='copied_source_retained');atomic_json(result_path,report)
                            # Lock the final destination through source cleanup. Hash AND exact
                            # bytes are verified, not just a matching filename or size.
                            with guarded_open(destination) as (dest,_):
                                if stream_hash(dest,deadline)!=digest or not exact_bytes(src,dest,deadline):raise OSError('Final destination differs; source retained')
                                safe_path(path,source);safe_path(destination,root)
                                if identity(path.stat())!=original or stream_hash(src,deadline)!=digest:raise Unsafe('Source changed before cleanup')
                                if hook:hook('before_cleanup',row,path)
                                remove_source(src,handle,path)
                            if path.exists():raise OSError('Source cleanup incomplete')
                            row.update(phase='committed',status='already_filed' if duplicate else 'moved',reason='Exact existing destination verified; source removed' if duplicate else ('Name collision; deterministic hash suffix applied. ' if row['renamed_for_collision'] else '')+'Verified destination; source removed')
                            atomic_json(result_path,report)
                    except BatchLimit as exc:
                        row.update(status='unprocessed',reason=str(exc));raise
                    except (Unsafe,PermissionError) as exc:
                        row['reason']=str(exc)
                        if row.get('phase')=='copied':row['status']='copied_source_retained';raise
                        # An unmatchable/locked source is not authority for any move.
                        if row.get('phase'):row['status']='failed';raise
                    except Exception as exc:
                        if row.get('phase')!='copied':row['status']='failed'
                        row['reason']=type(exc).__name__+': '+str(exc);raise
                if truncated:report['files'].append({'file':'Remaining source subtree','status':'unprocessed','reason':'Source metadata bound reached; remaining count unknown'})
                report['status']='complete' if report_counts(report)['ok'] else 'partial'
            except Exception as exc:
                report['error']=type(exc).__name__+': '+str(exc);report['status']='failed'
                known={x['file'] for x in report['files']}
                for path,issue in locals().get('files',[]):
                    rel=str(path.relative_to(source))
                    if rel not in known:report['files'].append({'file':rel,'status':'unprocessed','reason':'Batch stopped after structural/I/O failure'})
            report['elapsed_seconds']=time.monotonic()-started;report_counts(report);atomic_json(result_path,report)
            if json.loads(active.read_text())['request_id']==request_id:active.unlink()
            return report
    except (OSError,Unsafe) as exc:
        retained=locals().get('report',{'request_id':request_id,'files':[]})
        report_counts(retained);retained.update(ok=False,status='inconclusive' if retained['files'] else 'failed',error='Sorter lock/audit unavailable: '+str(exc))
        return retained

def run_local_batch(request_id,workbook,audit):
    """Bounded child; journal coordination locks are held by the parent adapter."""
    audit=Path(audit);audit.mkdir(parents=True,exist_ok=True);stdout=audit/(request_id+'.stdout.log');stderr=audit/(request_id+'.stderr.log');start=time.monotonic()
    command=[sys.executable,'-u',str(Path(__file__).resolve()),'--worker',request_id,'--workbook',str(workbook),'--audit',str(audit)]
    try:
        with stdout.open('w') as out,stderr.open('w') as err:process=subprocess.run(command,stdout=out,stderr=err,timeout=295)
        result_path=audit/(request_id+'.json');result=json.loads(result_path.read_text()) if result_path.exists() else report_counts({'ok':False,'status':'inconclusive','request_id':request_id,'files':[],'error':'Worker produced no durable result'})
        result['worker_returncode']=process.returncode
        if process.returncode!=0:result.update(ok=False,status='inconclusive',error='Worker failed; review audit, no automatic retry')
    except subprocess.TimeoutExpired:
        result_path=audit/(request_id+'.json')
        result=json.loads(result_path.read_text()) if result_path.exists() else {'request_id':request_id,'files':[]}
        report_counts(result);result.update(ok=False,status='inconclusive',worker_returncode=None,error='Worker timed out; inspect durable audit before another batch')
    result['worker_elapsed_seconds']=time.monotonic()-start;return result

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--worker',required=True);parser.add_argument('--workbook',required=True);parser.add_argument('--audit',required=True);args=parser.parse_args()
    sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
    # Resolve/scan only the two configured approved Dropbox roots, no fallback writes.
    roots=approved_roots()
    result=run_batch(args.worker,SOURCE,roots,Path(args.audit),loader=lambda deadline:read_candidates(Path(args.workbook),roots,deadline))
    print(json.dumps({'status':result['status'],'ok':result['ok']}),flush=True)
if __name__=='__main__':main()
