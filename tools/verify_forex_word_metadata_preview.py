"""Independent saved-preview verifier. No adapter planner/matcher/mapping calls.

The shared utility is restricted to immutable, path-safe one-read/hash/parse
storage. Eligibility, conversions, row comparisons and XML preconditions below
are independently derived, including the complete expected proposal set.
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
import ntpath
from pathlib import Path
import re
import sys
import time
from urllib.parse import unquote, urlsplit
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries
if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from tools.backfill_forex_word_metadata import JsonStore, stable_read_handle

N = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
OLD_HASH = "62e7b1cca79ef12f87460673ca5b8016d068e02b0400e2d0f0b7666c60bb9872"
SNAPSHOT_HASH = "93c11f6783282c5397ba7edbb23ab1229eb5411b4708cf8b8b2a961318e55b03"
TEMP = Path("C:/Users/User/AppData/Local/Temp")
RECEIPTS = [TEMP / s for s in (
    "forex_word_7b1_resume_4eb128350c0442ae8ba8e305c478497e",
    "forex_word_7b2_equivalence_dc2b5da560514669b8d8ae78a849321c",
    "forex_word_7b3_529338dadec748699802b92e0627af27",
    "forex_word_7b4_6b115ab1175c40c7ad1a6922e5a4e7e5",
    "forex_word_7b5_resume_ec2925baca0d4414aeed854c81c7b0e0",
    "forex_word_7b6_b01b253eb1bf45fa94485af0c198d146")]
INV = TEMP / "forex_word_7a_host_46efeff728734c77be80a4f196803cff"
ROOT = r"C:\Users\User\Dropbox\FOREX"
ORIGIN = r"C:\GPT\CODEX-master\journal\Trading Journal.xlsx"


def check(ok, message):
    if not ok:
        raise ValueError(message)


def sha(raw):
    return hashlib.sha256(raw).hexdigest()


def norm(value):
    return " ".join(str(value).strip().upper().split())


def pointer(value, path):
    check(isinstance(path, str) and path.startswith("/") and "//" not in path, "malformed pointer")
    check(not re.search(r"~(?![01])", path), "malformed pointer escape")
    try:
        for key in path[1:].split("/"):
            check(not re.search(r"~(?![01])", key), "malformed pointer escape")
            key = key.replace("~1", "/").replace("~0", "~")
            if isinstance(value, list):
                check(bool(re.fullmatch(r"0|[1-9][0-9]*", key)), "malformed list pointer")
                value = value[int(key)]
            else:
                value = value[key]
    except (IndexError, KeyError, TypeError) as exc:
        raise ValueError("pointer mismatch") from exc
    return value


def source_key(path):
    s = path.replace("\\", "/")
    check(not s.startswith("/") and not any(p in ("", ".", "..") or ":" in p for p in s.split("/")), "unsafe source path")
    return s.casefold()


def link_target(raw, origin, base):
    """Independent lexical interpretation; never touch a destination path."""
    def local(text):
        check(isinstance(text, str) and text and text == text.strip(), "invalid literal")
        check(not re.search(r"%(?![a-fA-F0-9]{2})", text), "bad percent")
        text = unquote(text, errors="strict")
        check("%" not in text and not any(ord(c) < 32 for c in text) and not any(c in text for c in '?#*"<>|'), "unsafe encoded literal")
        if text.lower().startswith("file:"):
            uri = urlsplit(text)
            check(not uri.netloc and re.fullmatch(r"/[A-Za-z]:/.*", uri.path, re.S) is not None, "file authority")
            text = uri.path[1:]
        text = text.replace("/", "\\")
        check(not text.startswith("\\") and (":" not in text or re.match(r"^[A-Za-z]:\\", text)), "ambiguous path")
        check(":" not in text[2:] and all(p in (".", "..") or not p.endswith((".", " ")) for p in text.split("\\") if p), "ambiguous component")
        return text
    try:
        value, workbook = local(raw), local(origin)
        check(re.match(r"^[A-Za-z]:\\", workbook), "missing absolute origin")
        if not re.match(r"^[A-Za-z]:\\", value):
            root = local(base) if base else ntpath.dirname(workbook)
            check(re.match(r"^[A-Za-z]:\\", root), "ambiguous base")
            value = ntpath.join(root, value)
        result = ntpath.normpath(value)
        check(ntpath.normcase(result).startswith(ntpath.normcase(ROOT) + "\\"), "outside FOREX")
        return result
    except (ValueError, UnicodeError):
        return None


def fnumber(value):
    link = None
    if not isinstance(value, str):
        return None, None
    value = value.strip()
    if value.startswith("="):
        f = re.fullmatch(r'=HYPERLINK\("((?:[^"\r\n]|"")*)"\s*[,;]\s*"(F[0-9]+)"\)', value, re.I)
        if not f:
            return None, None
        link, value = f[1].replace('""', '"'), f[2]
    return ("F" + str(int(value[1:])), link) if re.fullmatch("F[0-9]+", value, re.I) else (None, link)


CCY = set("USD EUR GBP JPY AUD NZD CAD CHF NOK SEK DKK SGD HKD ZAR MXN TRY PLN CNH CNY".split())


def pair(value):
    p = re.sub(r"[ /]", "", norm(value))
    return p if len(p) == 6 and p[:3] in CCY and p[3:] in CCY and p[:3] != p[3:] else None


def direction(value):
    return {"BUY": "Long", "LONG": "Long", "SELL": "Short", "SHORT": "Short"}.get(norm(value))


def date(value):
    s = str(value).strip()
    for fmt in (("%d%m%y",) if re.fullmatch(r"\d{6}", s) else ("%d/%m/%y", "%d/%m/%Y", "%d-%m-%Y")):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def observations(doc, prefix):
    fields = []
    for ti, t in enumerate(doc["coverage"]["tables"]):
        grid = {(c["row"], c["column"]): (i, c) for i, c in enumerate(t["cells"])}
        check(len(grid) == len(t["cells"]), "duplicate origin")
        for (r, col), (i, label) in grid.items():
            if col % 2:
                continue
            other = grid.get((r, col + label["colspan"]))
            if not other and not label["text"].strip():
                continue
            vi, v = other if other else (i, label)
            fields.append(dict(label=label["text"], raw=v["text"], pointer=prefix+f"/coverage/tables/{ti}/cells/{vi}", label_pointer=prefix+f"/coverage/tables/{ti}/cells/{i}", table_id=t["id"], part=t["part"], coordinates={k:v[k] for k in ("row", "column", "rowspan", "colspan")}, label_coordinates={k:label[k] for k in ("row", "column", "rowspan", "colspan")}, unpaired=not bool(other)))
    for i, p in enumerate(doc["coverage"]["prose"]):
        if p.get("text", "").strip():
            fields.append(dict(label="[unstructured prose]", raw=p["text"], pointer=prefix+f"/coverage/prose/{i}/text", context=p.get("context"), unpaired=True))
    return fields


BOOL = {s.upper():s for s in ("Test", "EMA", "VWAP", "Round Number", "Spiked Out", "Close Stopout", "Near Perfect Entry", "Near Win", "Early Close")}


def conversion(field, headers, fmt=None):
    """Independent, closed manual-field domain; no generic numeric inference."""
    label, raw = norm(field["label"]), field["raw"]
    if field.get("unpaired"):
        return "ambiguous", None, None, None
    if raw is None or not str(raw).strip() or norm(raw) in {"N/A", "UNKNOWN", "-"}:
        return "absent", None, None, None
    header = None
    kind = None
    if label in BOOL:
        header, kind = BOOL[label], "yesno"
    elif label in {"SETUP", "TIMEFRAME", "CHART", "ORDER", "MARKET/LIMIT/PENDING ORDER", "PATTERN", "ATHS/ATLS"}:
        header, kind = {"SETUP":("Setup","text"),"TIMEFRAME":("Timeframe","timeframe"),"CHART":("Timeframe","timeframe"),"ORDER":("Order","order"),"MARKET/LIMIT/PENDING ORDER":("Order","order"),"PATTERN":("Pattern","pattern"),"ATHS/ATLS":("ATHS/ATLS","extreme")}[label]
    else:
        for group in ("Break Even", "Profit"):
            for measure in ("Trigger Price", "Distance From Entry %", "Distance From Exit %"):
                name = f"Move to {group} {measure}"
                if label == name.upper():
                    header, kind = name, "fraction" if measure.endswith("%") else "decimal"
    if not header:
        return "no_safe_destination", None, None, None
    if header not in headers:
        return "no_safe_destination", header, kind, None
    unit, value = "text", raw
    if kind == "yesno":
        value = {"YES":"Yes","NO":"No","TRUE":"Yes","FALSE":"No","0":"No","1":"Yes"}.get(norm(raw))
    elif kind == "timeframe":
        value = {"1MIN":"1MIN","5MIN":"5MIN","15MIN":"15MIN","30MIN":"30MIN","1H":"1H","4H":"4H","DAILY":"1D","WEEKLY":"1W","MONTHLY":"1MO"}.get(re.sub(r"\s", "", norm(raw)))
    elif kind == "order":
        value = {"MARKET":"Market","MARKET ORDER":"Market","LIMIT":"Limit","BUY LIMIT":"Limit","SELL LIMIT":"Limit","LIMIT ORDER":"Limit"}.get(norm(raw))
    elif kind == "pattern":
        value = str(raw).strip().lower() if str(raw).strip().lower() in {"range", "channel"} else None
    elif kind == "extreme":
        value = {"ATH":"All-Time High","ATHS":"All-Time High","ALL-TIME HIGH":"All-Time High","ATL":"All-Time Low","ATLS":"All-Time Low","ALL-TIME LOW":"All-Time Low"}.get(norm(raw))
    elif kind in ("decimal", "fraction"):
        s = str(raw).strip()
        ok = kind != "fraction" or s.endswith("%") and (fmt is None or "%" in fmt)
        s = s[:-1].strip() if kind == "fraction" and s.endswith("%") else s
        ok = ok and bool(re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)", s))
        value = str(Decimal(s) / 100) if ok and kind == "fraction" else str(Decimal(s)) if ok else None
        unit = "fraction_of_one" if kind == "fraction" else "price"
    typed = dict(type="decimal" if kind in ("fraction", "decimal") else "string", value=value, unit=unit) if value is not None else None
    return "supported" if typed else "ambiguous", header, kind, typed


def identity_audit(doc, fields, ws, headers, rows, rowids, origin, base, legacy):
    src = doc["source"]
    numbers = set()
    for text in [src["relative_path"], *src.get("f_number_evidence", [])]:
        numbers.update("F"+str(int(n)) for n in re.findall(r"(?i)(?<!\w)F(\d+)(?!\w)", text.replace("\\", "/")))
    text = "\n".join(c["text"] for t in doc["coverage"]["tables"] for c in t["cells"]) + "\n" + "\n".join(p.get("text", "") for p in doc["coverage"]["prose"])
    content = {"F"+str(int(n)) for n in re.findall(r"(?i)(?<!\w)F(\d+)(?!\w)", text)}
    number = next(iter(numbers)) if len(numbers)==1 else None
    candidates = rows.get(number, []) if number else []
    reasons = []
    if len(numbers)!=1 or content-numbers:
        reasons.append("missing_or_conflicting_source_f_number")
    if len(candidates)!=1:
        reasons.append("missing_or_duplicate_journal_f_number")
    evidence = []
    row = candidates[0] if len(candidates)==1 else None
    rid = None
    if row is not None:
        v = {h: ws.cell(row,c).value for h,c in headers.items()}
        rid = v["Row ID"]
        if not isinstance(rid,str) or not rid.strip() or rowids[str(rid)]!=1 or ws.cell(row,headers["Row ID"]).data_type=="f":
            reasons.append("missing_duplicate_or_formula_row_id")
        for h in ("Row Type","Symbol","Open Time","Side","Entry Price"):
            if h in headers and ws.cell(row,headers[h]).data_type=="f":
                reasons.append("formula_identity:"+h)
        if norm(v.get("Row Type"))!="TRADE" or pair(v.get("Symbol")) is None:
            reasons.append("not_qualifying_forex_trade")
        observed = [x for x in fields if norm(x["label"]) in {"PAIR","SYMBOL"} and x["raw"] is not None and str(x["raw"]).strip()]
        pairs = {pair(x["raw"]) for x in observed}|{pair(s) for s in src.get("symbol_evidence",[])}
        if pairs:
            agrees = pairs=={pair(v.get("Symbol"))} and None not in pairs
            evidence.append(("symbol", agrees))
            if not agrees:reasons.append("symbol_conflict_or_ambiguous")
        opening = v.get("Open Time")
        if isinstance(opening,str):
            try:opening=dt.datetime.fromisoformat(opening)
            except ValueError:opening=None
        if isinstance(opening,dt.datetime):opening=opening.date()
        year = str(src.get("year",""))
        if year:
            agrees=isinstance(opening,dt.date) and str(opening.year)==year
            evidence.append(("year",agrees))
            if not agrees:reasons.append("year_conflict_or_unknown")
        for field in fields:
            label, raw = norm(field["label"]),field["raw"]
            if raw is None or not str(raw).strip() or norm(raw) in {"N/A","-","UNKNOWN"}:continue
            if label in {"SIDE","SHORT/LONG"}:canonical,current=direction(raw),direction(v.get("Side"))
            elif label in {"ENTRY","ENTRY PRICE"}:
                try:canonical,current=Decimal(str(raw).strip()),Decimal(str(v.get("Entry Price")))
                except InvalidOperation:canonical,current=None,None
            elif label=="DATE":canonical,current=date(raw),opening
            else:continue
            agrees=canonical is not None and canonical==current
            evidence.append((label,agrees))
            if not agrees:reasons.append("identity_conflict_or_ambiguous:"+label)
        cell=ws.cell(row,headers["Trade Number"])
        target=fnumber(cell.value)[1] or (cell.hyperlink.target if cell.hyperlink else None)
        agrees=False
        normalized=None
        if target:
            expected=ntpath.normpath(ROOT+"\\"+src["parent_relative_path"].replace("/","\\"))
            if legacy:
                actual=str(target).replace("\\","/").casefold().rstrip("/")
                if actual.startswith("file:///"):
                    uri=urlsplit(actual);actual=unquote(uri.path).lstrip("/").rstrip("/") if not uri.netloc else ""
                agrees=actual==expected.replace("\\","/").casefold()
            else:
                normalized=link_target(target,origin,base)
                agrees=normalized is not None and ntpath.normcase(normalized)==ntpath.normcase(expected)
            evidence.append(("folder_hyperlink",agrees))
            if not agrees:reasons.append("folder_hyperlink_conflict_or_unknown")
        symbol_year=bool(pairs) and pairs=={pair(v.get("Symbol"))} and isinstance(opening,dt.date) and year==str(opening.year)
        if not(symbol_year or agrees):reasons.append("insufficient_non_f_corroboration")
    return dict(f_number=number,row=row,row_id=rid,candidate_rows=candidates,reasons=sorted(set(reasons)),status="unresolved" if reasons else "proven",comparisons=evidence,normalized_link=normalized if row is not None else None)


def target_state(ws, cells, ranges, row, header, headers):
    cell=ws.cell(row,headers[header]);node=cells.get(cell.coordinate);reason=None
    if any(l<=cell.column<=r and t<=row<=b for l,t,r,b in ranges):reason="formula_range"
    for merged in ws.merged_cells.ranges:
        if cell.coordinate in merged and cell.coordinate!=merged.start_cell.coordinate:reason="merged_continuation"
    if cell.data_type=="f" or node is not None and node.find(N+"f") is not None:reason="formula"
    if reason is None and node is not None and (list(node) or node.text not in (None,"") or node.get("t") not in (None,"n") or any(k not in ("r","s","t") for k in node.attrib)):reason="nonblank"
    if reason is None and (cell.value is not None or cell.data_type!="n"):reason="nonblank"
    xml=ET.tostring(node,encoding="unicode") if node is not None else None
    return dict(coordinate=cell.coordinate,xml=xml,xml_sha256=sha(xml.encode()) if xml else None,data_type=cell.data_type,style_id=cell.style_id,number_format=cell.number_format,hyperlink=str(cell.hyperlink.target) if cell.hyperlink else None,comment=cell.comment.text if cell.comment else None,structurally_blank=reason is None,reason=reason)


def verify(result_directory, report, expected_adapter_hash, *, legacy=False, approved=None, inventory_directory=None, expected_count=102):
    start=time.monotonic();directory=Path(result_directory).absolute();approved=[Path(p) for p in (approved or RECEIPTS)];inv=Path(inventory_directory or INV)
    store=JsonStore([directory,*approved,inv]);snapshot_handle=None
    try:
        receipt=store.load(directory/"input_receipt.json");proposal=store.load(directory/"proposed_cell_patch.json");audit=store.load(directory/"document_match_audit.json");catalog=store.load(directory/"field_mapping_catalog.json");result=store.load(directory/"preview_result.json")
        check(receipt["adapter_sha256"]==proposal["adapter_sha256"]==expected_adapter_hash,"adapter identity mismatch")
        check(proposal["mapping_sha256"]==store.hashes[str(directory/"field_mapping_catalog.json")],"mapping hash mismatch")
        verifier_hash=sha(Path(__file__).read_bytes())
        if not legacy:check(receipt["verifier_sha256"]==proposal["verifier_sha256"]==verifier_hash,"verifier identity mismatch")
        for name,expected in receipt["saved_json_hashes"].items():
            check(any(Path(name)==root or root in Path(name).parents for root in [*approved,inv]),"unexpected saved input")
            if Path(name).suffix==".jsonl":
                store.jsonl(name);check(store.hashes[str(Path(name))]==expected,"inventory hash mismatch")
            else:store.load(name,expected)
        snapshot_path=store.checked(receipt["snapshot"]["path"],directory)
        check(snapshot_path.parent==directory and snapshot_path.suffix==".xlsx","snapshot ownership")
        snapshot_handle=stable_read_handle(snapshot_path);raw=snapshot_handle.read();snapshot_hash=sha(raw)
        check(snapshot_hash==receipt["snapshot"]["sha256"]==proposal["snapshot_sha256"]==result["snapshot_sha256"],"snapshot mismatch")
        capture_path=directory/"snapshot_receipt.json"
        capture=store.load(capture_path) if capture_path.is_file() else None
        if capture:
            check(capture==receipt["snapshot"]["capture"] and store.hashes[str(capture_path)]==receipt["snapshot"]["capture_receipt_sha256"],"origin receipt mismatch")
            check(capture["sha256"]==snapshot_hash and capture["before"]==capture["after"] and capture["successful_copies"]==1 and capture["share_read_only"],"unproven original snapshot")
            origin=capture["source"]
            if expected_count==102:check(origin==ORIGIN and snapshot_hash==SNAPSHOT_HASH,"unexpected historical snapshot/origin")
        else:origin=None
        # Independent receipt/source set verification; only JSON is read.
        context=store.load(approved[-1]/"final_audit_context.json");refs=context["prior_completions"]
        check([Path(r["directory"]) for r in refs]==approved,"receipt directory mismatch")
        inventory=store.jsonl(inv/"manifest.jsonl");check(store.hashes[str(inv/"manifest.jsonl")]==context["inventory_sha256"],"inventory linkage")
        manifest={source_key(x["relative_path"]):x for x in inventory if x.get("kind")=="document"};check(len(manifest)==expected_count,"source count")
        handoff=store.load(approved[-1]/"final_recovery_handoff.json");coverage=store.load(handoff["coverage_audit"]["path"],handoff["coverage_audit"]["sha256"]);progress=store.load(approved[-1]/"cumulative_progress.json")
        check(coverage["status"]=="complete" and coverage["exact_set_equality"] and coverage["completed_receipts"]==progress["completed_receipts"]==refs and progress["remaining"]==0,"incomplete aggregate receipt")
        docs={};linkage={}
        for ref in refs:
            root=Path(ref["directory"]);r={name:store.load(root/name,h,root) for name,h in ref["records"].items()}
            check(set(r)=={"preparation.json","selection.validated.json","publication.json","published_verification.json","preflight.json"},"receipt record set")
            prep,selection,pub,verified=r["preparation.json"],r["selection.validated.json"],r["publication.json"],r["published_verification.json"]
            check(prep["status"]==pub["status"]==verified["status"]=="complete" and prep["tool_sha256"]==ref["tool_sha256"]==r["preflight.json"]["tool_sha256"],"incomplete receipt")
            check(prep["selection_sha256"]==pub["selection_sha256"]==ref["records"]["selection.validated.json"],"selection binding")
            check(selection["root"]==context["root"] and selection["inventory_sha256"]==context["inventory_sha256"],"root/inventory binding")
            check(len(prep["documents"])==prep["source_count"]==verified["documents"]==len(pub["source_checks"]) and [d["source"] for d in prep["documents"]]==selection["documents"],"document count/selection")
            outputs={o["destination"]:(o["sha256"],o["kind"]) for d in prep["documents"] for o in d["outputs"]};published={o["path"]:(o["sha256"],o["kind"]) for o in pub["published"]}
            check(outputs==published and len(pub["published"])==len(outputs)==verified["published_files"],"published receipt coverage")
            check(all(verified["counts"].get(k,0)==sum(v[1]==k for v in outputs.values()) for k in ("table","picture","text")),"publication counts")
            checks={ntpath.normcase(c["source"]):c for c in pub["source_checks"]}
            for i,doc in enumerate(prep["documents"]):
                key=source_key(doc["source"]["relative_path"]);check(key not in docs and key in manifest and all(doc["source"].get(k)==v for k,v in manifest[key].items()),"duplicate/foreign source")
                check(doc["complete"] and not doc["coverage"]["issues"],"incomplete source")
                path=str(Path(context["root"]).joinpath(*doc["source"]["relative_path"].replace("\\","/").split("/")));c=checks[ntpath.normcase(path)]
                check(c["sha256"]==doc["source_sha256"] and all(c["before"][k]==c["after"][k]==doc["source"][k] for k in ("size","mtime_ns")),"source preservation linkage")
                check(store.load(root/f"document_{i+1:02d}"/"provenance.json",directory=root)==doc,"provenance mismatch")
                docs[key]=doc;linkage[key]=(root,ref,i)
        check(set(docs)==set(manifest) and len(handoff["documents"])==expected_count,"exact source set")
        entries={}
        for h in handoff["documents"]:
            key=h["source_identity"];check(key not in entries and key in docs,"handoff duplicate");entries[key]=h;d=docs[key];root,ref,i=linkage[key];prefix=f"/documents/{i}"
            check(h["source"]==d["source"] and h["source_sha256"]==d["source_sha256"] and Path(h["receipt_directory"])==root and h["receipt_hashes"]==ref["records"],"handoff identity/hash mismatch")
            check(h["preparation_reference"]==dict(path=str(root/"preparation.json"),sha256=ref["records"]["preparation.json"],document_index=i,json_pointer=prefix),"handoff pointer mismatch")
            check(Path(h["provenance_reference"]["path"])==root/f"document_{i+1:02d}"/"provenance.json" and h["structured_tables"]==[dict(table_id=t["id"],json_pointer=prefix+f"/coverage/tables/{j}") for j,t in enumerate(d["coverage"]["tables"])],"table/provenance pointer mismatch")
        # Workbook data/formulas and relationship targets come solely from bytes.
        wb=load_workbook(io.BytesIO(raw),data_only=False,keep_links=False);ws=wb["Trade Log"];headers=receipt["schema"]["headers"];startrow=receipt["schema"]["data_start"]
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            sheet=ET.fromstring(z.read(receipt["schema"]["worksheet_part"]));props=ET.fromstring(z.read("docProps/app.xml"));b=props.find("{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}HyperlinkBase");base=b.text if b is not None else None
            part=receipt["schema"]["worksheet_part"];rp=str(Path(part).parent/"_rels"/(Path(part).name+".rels")).replace("\\","/");relationships={x.get("Id"):x.get("Target") for x in ET.fromstring(z.read(rp))} if rp in z.namelist() else {}
            counts={}
            for p in z.namelist():
                if re.fullmatch(r"xl/worksheets/sheet[0-9]+\.xml",p):
                    nodes=list(ET.fromstring(z.read(p)).iter(N+"c"));counts[p]=dict(explicit_cells=len(nodes),formula_cells=sum(c.find(N+"f") is not None for c in nodes),literal_or_formula_nonblank=sum(bool(list(c)) or c.get("t") not in (None,"n") for c in nodes))
        cells={c.get("r"):c for c in sheet.iter(N+"c")};ranges=[range_boundaries(f.get("ref")) for c in cells.values() for f in [c.find(N+"f")] if f is not None and f.get("ref")]
        rows=collections.defaultdict(list);rowids=collections.Counter()
        for row in range(startrow,ws.max_row+1):
            number,_=fnumber(ws.cell(row,headers["Trade Number"]).value)
            if number:rows[number].append(row)
            rid=ws.cell(row,headers["Row ID"]).value
            if rid is not None:rowids[str(rid)]+=1
        observed_audits={a["source_identity"]:a for a in audit["documents"]};check(len(observed_audits)==len(audit["documents"])==expected_count and set(observed_audits)==set(docs),"audit coverage")
        cats=collections.defaultdict(list)
        for c in catalog:cats[(c["source_identity"],c["source_pointer"])].append(c)
        expected_entries=[];decisions=[];expected_observation_count=0;matched_targets={}
        for key,h in entries.items():
            doc=docs[key];fields=observations(doc,h["preparation_reference"]["json_pointer"]);expected_observation_count+=len(fields);a=identity_audit(doc,fields,ws,headers,rows,rowids,origin,base,legacy);recorded=observed_audits[key]
            check(all(recorded[k]==a[k] for k in ("f_number","row","row_id","candidate_rows","reasons","status")) and recorded["source_sha256"]==doc["source_sha256"],"row identity/conflicts mismatch")
            check([(c["kind"],c["agrees"]) for c in recorded["comparisons"]]==a["comparisons"],"corroboration mismatch")
            if a["row"] is not None:
                cell=ws.cell(a["row"],headers["Trade Number"])
                if cell.hyperlink:check(relationships[cell.hyperlink.id]==cell.hyperlink.target,"raw relationship mismatch")
                if not legacy:
                    for c in recorded["comparisons"]:
                        if c["kind"]=="folder_hyperlink":check(c["normalization"]["normalized_target"]==a["normalized_link"] and c["normalization"]["original_workbook"]==origin and c["normalization"]["hyperlink_base"]==base and c["cell"]==cell.coordinate,"normalized link claim mismatch")
            groups=collections.defaultdict(list);blocked=set()
            for field in fields:
                status,header,kind,value=conversion(field,headers)
                if header and a["status"]=="proven":status,header,kind,value=conversion(field,headers,ws.cell(a["row"],headers[header]).number_format)
                records=cats[(key,field["pointer"])];check(len(records)==1,"catalog duplicate/missing pointer");c=records[0]
                check(c["raw"]==field["raw"] and c["label"]==field["label"] and c["status"]==status and c.get("header")==header and c.get("type")==kind and c.get("canonical")==value,"catalog value/conversion mismatch")
                check(c["context"]=={k:field.get(k) for k in ("table_id","part","context","label_coordinates")},"catalog context mismatch")
                pointed=pointer(store.load(h["preparation_reference"]["path"]),field["pointer"]);check((pointed["text"] if isinstance(pointed,dict) else pointed)==field["raw"],"catalog raw pointer mismatch")
                if status=="supported":groups[header].append((field,value))
                else:
                    if header and status!="absent":blocked.add(header)
                    decisions.append((status,key,None,(field["pointer"],)))
            for header,items in groups.items():
                pointers=tuple(f["pointer"] for f,_ in items)
                if header in blocked:status="ambiguous_source_field"
                elif len({json.dumps(v,sort_keys=True) for _,v in items})!=1:status="conflicting_source_field"
                elif a["status"]!="proven":status="unresolved_identity"
                else:
                    target=target_state(ws,cells,ranges,a["row"],header,headers);matched_targets[target["coordinate"]]=target
                    if not target["structurally_blank"]:
                        existing=ws.cell(a["row"],headers[header]);status="existing_equal" if str(existing.value)==str(items[0][1]["value"]) and existing.data_type!="f" else "existing_conflicting"
                    else:
                        value=items[0][1];source_fields=[dict(pointer=f["pointer"],raw=f["raw"],table_id=f.get("table_id"),coordinates=f.get("coordinates"),label=f["label"]) for f,_ in items]
                        expected_entries.append(dict(sheet=ws.title,coordinate=target["coordinate"],header=header,row=a["row"],row_id=a["row_id"],f_number=a["f_number"],source_identity=key,source_sha256=doc["source_sha256"],preparation_path=h["preparation_reference"]["path"],sources=source_fields,new_value=value,target_precondition=target));continue
                decisions.append((status,key,header,pointers))
        check(len(catalog)==expected_observation_count and len(cats)==expected_observation_count,"catalog extra/missing records")
        collision=collections.defaultdict(list)
        for e in expected_entries:collision[(e["sheet"],e["coordinate"])].append(e)
        expected=[]
        for coordinate,es in sorted(collision.items()):
            if len({json.dumps(e["new_value"],sort_keys=True) for e in es})!=1:
                decisions.append(("conflicting_documents_coordinate",None,None,(coordinate[1],)));continue
            first=es[0];first["attributions"]=[dict(source_identity=e["source_identity"],source_sha256=e["source_sha256"],preparation_path=e["preparation_path"],sources=e["sources"]) for e in es];expected.append(first)
        check(len({(e["sheet"],e["coordinate"]) for e in proposal["entries"]})==len(proposal["entries"]),"duplicate proposal coordinate")
        check(proposal["entries"]==expected,"eligible proposal set/value/attribution/target mismatch")
        actual_decisions=[]
        for d in proposal["decisions"]:
            pointers=tuple(d["source_pointers"]) if "source_pointers" in d else (d["pointer"],) if "pointer" in d else (d.get("coordinate"),)
            actual_decisions.append((d["status"],d.get("source_identity"),d.get("header"),pointers))
        check(collections.Counter(actual_decisions)==collections.Counter(decisions),"decision coverage mismatch")
        check(result["counts"]["documents"]==expected_count and result["counts"]["proposed_cells"]==len(expected) and result["counts"]["proven_matches"]==sum(a["status"]=="proven" for a in observed_audits.values()),"result counts mismatch")
        store.close();snapshot_handle.close();snapshot_handle=None
        check(all(max(c.values())<=1 for c in store.counters.values()),"cache re-read/hash/parse")
        report_value=dict(status="complete",result_policy="historical conservative link policy; integrity/eligible-set checked, link interpretation remains a known limitation" if legacy else "corrected original-workbook link policy independently verified",legacy_identity_interpretation_fully_validated=not legacy,expected_adapter_sha256=expected_adapter_hash,verifier_sha256=verifier_hash,snapshot_sha256=snapshot_hash,source_count=expected_count,observed_field_occurrences=expected_observation_count,proven_matches=sum(a["status"]=="proven" for a in observed_audits.values()),unresolved_matches=sum(a["status"]!="proven" for a in observed_audits.values()),proposed_cells=len(expected),proposed_rows=len({e["row"] for e in expected}),eligible_set_exact=True,all_pointers_checked=True,decisions_checked=len(decisions),cache_counters=dict(store.counters),max_reads_per_file=max(c["reads"] for c in store.counters.values()),max_hashes_per_file=max(c["hashes"] for c in store.counters.values()),max_parses_per_file=max(c["parses"] for c in store.counters.values()),workbook_wide=dict(scope="all explicit worksheet XML cells across every worksheet part; string-typed empties count nonblank",parts=counts,total_formula_cells=sum(c["formula_cells"] for c in counts.values()),total_nonblank_cells=sum(c["literal_or_formula_nonblank"] for c in counts.values())),matched_targets=dict(scope="unique supported unconflicted manual-field targets in independently proven rows",count=len(matched_targets),nonblank=sum(not t["structurally_blank"] for t in matched_targets.values()),formula=sum(t["reason"] in {"formula","formula_range"} for t in matched_targets.values())),historical_inputs_unchanged=True,immutability_method="write/delete-denying handles through final metadata check; no JSON reread/reparse",production_writes=0,elapsed_seconds=time.monotonic()-start)
        with Path(report).open("x",encoding="utf-8") as handle:json.dump(report_value,handle,ensure_ascii=False,indent=2)
        return report_value
    finally:
        store.close()
        if snapshot_handle:snapshot_handle.close()


def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument("--result-directory",type=Path,required=True);p.add_argument("--report",type=Path,required=True);p.add_argument("--expected-adapter-sha256",required=True);p.add_argument("--legacy-link-policy",action="store_true");a=p.parse_args();r=verify(a.result_directory,a.report,a.expected_adapter_sha256,legacy=a.legacy_link_policy);print(json.dumps({k:v for k,v in r.items() if k not in {"cache_counters","workbook_wide"}}))


if __name__=="__main__":
    sys.dont_write_bytecode=True
    main()
