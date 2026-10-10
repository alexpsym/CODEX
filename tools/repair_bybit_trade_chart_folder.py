"""Dry-run-first repair for the single accepted Bybit Demo C377 chart folder.

The repair moves the existing chart directory from its stale month to the
accepted trade month and changes only that row's existing OOXML hyperlink.
It deliberately does not rebuild or save the workbook through openpyxl.
"""
from __future__ import annotations

import argparse
import errno
import hashlib
import json
import os
import re
import shutil
import sys
import uuid
import zipfile
from datetime import datetime
from pathlib import Path, PurePosixPath
from typing import Any, Callable
from urllib.parse import unquote, urlparse
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

_REPO_ROOT = Path(__file__).resolve().parents[1]
if str(_REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(_REPO_ROOT))

from tools.master_journal_workbook import (
    _TRADE_FOLDER_INDEX_CACHE,
    _select_trade_chart_root,
    _trade_log_data_start_row,
    _trade_log_header_map,
)


EXPECTED_ROW_ID = "bybit:demo:trade:BTCUSDT:cd7d7d60e40d420e"
EXPECTED_TRADE_NUMBER = "C377"
EXPECTED_ACCOUNT = "Bybit Demo"
EXPECTED_SYMBOL = "BTCUSDT"
EXPECTED_OPEN = datetime(2026, 10, 4, 13, 7, 22)
EXPECTED_CLOSE = datetime(2026, 10, 7, 12, 1, 7)
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_Q = lambda namespace, name: f"{{{namespace}}}{name}"


class RepairRefused(RuntimeError):
    """A precondition or verification failed without authorizing broader edits."""


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _canonical_json(value: Any) -> bytes:
    return json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(",", ":")).encode("utf-8")


def _file_uri_path(uri: str) -> str:
    parsed = urlparse(str(uri or ""))
    if parsed.scheme.casefold() != "file" or parsed.netloc not in ("", "localhost"):
        raise RepairRefused("C377's existing hyperlink is not a local file URI")
    path = unquote(parsed.path)
    if re.match(r"^/[A-Za-z]:/", path):
        path = path[1:]
    return path.replace("/", "\\").rstrip("\\").casefold()


def _same_path_uri(uri: str, path: Path) -> bool:
    try:
        return _file_uri_path(uri) == str(path.resolve()).rstrip("\\").casefold()
    except (OSError, RepairRefused):
        return False


def _workbook_row(workbook_path: Path) -> dict[str, Any]:
    wb = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        if "Trade Log" not in wb.sheetnames:
            raise RepairRefused("Trading Journal has no Trade Log sheet")
        ws = wb["Trade Log"]
        headers = _trade_log_header_map(ws)
        required = ("Trade Number", "Row ID", "Account", "Symbol", "Open Time", "Close Time")
        if any(name not in headers for name in required):
            raise RepairRefused("Trade Log is missing a required C377 identity/date column")
        matches = []
        start_row = _trade_log_data_start_row(ws)
        for row_num, values in enumerate(ws.iter_rows(min_row=start_row, values_only=True), start=start_row):
            row_id_index = headers["Row ID"] - 1
            row_id = str(values[row_id_index] if row_id_index < len(values) else "").strip()
            if row_id != EXPECTED_ROW_ID:
                continue
            def cell_value(header: str):
                index = headers[header] - 1
                return values[index] if index < len(values) else None
            matches.append({
                "row": row_num,
                "trade_number": str(cell_value("Trade Number") or "").strip(),
                "row_id": row_id,
                "account": str(cell_value("Account") or "").strip(),
                "symbol": str(cell_value("Symbol") or "").strip(),
                "open_time": cell_value("Open Time"),
                "close_time": cell_value("Close Time"),
                "trade_number_column": headers["Trade Number"],
            })
        if len(matches) != 1:
            raise RepairRefused(f"Expected exactly one stable C377 row; found {len(matches)}")
        item = matches[0]
        if (item["trade_number"], item["account"], item["symbol"]) != (
            EXPECTED_TRADE_NUMBER, EXPECTED_ACCOUNT, EXPECTED_SYMBOL
        ):
            raise RepairRefused("Stable row ID does not match C377 / Bybit Demo / BTCUSDT")
        if item["open_time"] != EXPECTED_OPEN or item["close_time"] != EXPECTED_CLOSE:
            raise RepairRefused("Saved C377 dates do not match the accepted October timestamps")
        return item
    finally:
        wb.close()


def _json_row(json_path: Path) -> dict[str, Any]:
    try:
        payload = json.loads(json_path.read_text(encoding="utf-8"))
    except Exception as exc:
        raise RepairRefused(f"Could not read the journal JSON trade source: {exc}") from exc
    rows = payload.get("items", []) if isinstance(payload, dict) else payload
    if not isinstance(rows, list):
        raise RepairRefused("Journal JSON does not contain an items list")
    matches = [r for r in rows if isinstance(r, dict) and str(r.get("id") or r.get("row_id") or "").strip() == EXPECTED_ROW_ID]
    if len(matches) != 1:
        raise RepairRefused(f"Expected exactly one stable C377 JSON row; found {len(matches)}")
    row = matches[0]
    opened = datetime.fromisoformat(str(row.get("open_time") or "")).replace(tzinfo=None)
    closed = datetime.fromisoformat(str(row.get("close_time") or "")).replace(tzinfo=None)
    if (
        str(row.get("trade_number") or "").strip() != EXPECTED_TRADE_NUMBER
        or str(row.get("account") or row.get("account_label") or "").strip() != EXPECTED_ACCOUNT
        or str(row.get("symbol") or "").strip() != EXPECTED_SYMBOL
        or opened != EXPECTED_OPEN
        or closed != EXPECTED_CLOSE
    ):
        raise RepairRefused("Journal JSON row does not match the accepted C377 identity and October dates")
    return {"row_id": EXPECTED_ROW_ID, "trade_number": EXPECTED_TRADE_NUMBER, "account": EXPECTED_ACCOUNT,
            "symbol": EXPECTED_SYMBOL, "open_time": opened.isoformat(), "close_time": closed.isoformat()}


def _worksheet_parts(zf: zipfile.ZipFile) -> tuple[str, str]:
    workbook_xml = ET.fromstring(zf.read("xl/workbook.xml"))
    rels_xml = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    rels = {r.attrib["Id"]: r.attrib["Target"] for r in rels_xml.findall(_Q(PKG_REL_NS, "Relationship"))}
    sheet = next((s for s in workbook_xml.findall(f"{_Q(MAIN_NS, 'sheets')}/{_Q(MAIN_NS, 'sheet')}") if s.attrib.get("name") == "Trade Log"), None)
    if sheet is None:
        raise RepairRefused("Trading Journal package has no Trade Log worksheet")
    rid = sheet.attrib.get(_Q(DOC_REL_NS, "id"), "")
    target = rels.get(rid)
    if not target:
        raise RepairRefused("Trade Log worksheet relationship is missing")
    sheet_part = target.lstrip("/") if target.startswith("/") else str(PurePosixPath("xl") / target)
    sheet_part = str(PurePosixPath(sheet_part))
    rel_part = str(PurePosixPath(sheet_part).parent / "_rels" / (PurePosixPath(sheet_part).name + ".rels"))
    if rel_part not in zf.namelist():
        raise RepairRefused("Trade Log has no worksheet relationship part for its existing hyperlink")
    return sheet_part, rel_part


def _hyperlink_info(workbook_path: Path, row: dict[str, Any]) -> dict[str, Any]:
    cell_ref = f"{get_column_letter(row['trade_number_column'])}{row['row']}"
    with zipfile.ZipFile(workbook_path, "r") as zf:
        sheet_part, rel_part = _worksheet_parts(zf)
        sheet_root = ET.fromstring(zf.read(sheet_part))
        links = [h for h in sheet_root.findall(f"{_Q(MAIN_NS, 'hyperlinks')}/{_Q(MAIN_NS, 'hyperlink')}") if h.attrib.get("ref") == cell_ref]
        if len(links) != 1:
            raise RepairRefused(f"Expected one existing hyperlink on C377 cell {cell_ref}; found {len(links)}")
        rel_id = links[0].attrib.get(_Q(DOC_REL_NS, "id"), "")
        if not rel_id:
            raise RepairRefused("C377's existing hyperlink is not backed by an external relationship")
        rel_root = ET.fromstring(zf.read(rel_part))
        rels = [r for r in rel_root.findall(_Q(PKG_REL_NS, "Relationship")) if r.attrib.get("Id") == rel_id]
        if len(rels) != 1 or rels[0].attrib.get("TargetMode") != "External":
            raise RepairRefused("C377's existing hyperlink relationship is missing or not external")
        if not str(rels[0].attrib.get("Type", "")).endswith("/hyperlink"):
            raise RepairRefused("C377's existing relationship is not a hyperlink")
        uses = sum(1 for h in sheet_root.findall(f"{_Q(MAIN_NS, 'hyperlinks')}/{_Q(MAIN_NS, 'hyperlink')}") if h.attrib.get(_Q(DOC_REL_NS, "id")) == rel_id)
        return {"cell_ref": cell_ref, "relationship_id": rel_id, "relationship_uses": uses,
                "target": rels[0].attrib.get("Target", ""), "sheet_part": sheet_part, "rels_part": rel_part}


def _inventory(folder: Path) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for path in sorted(folder.rglob("*"), key=lambda p: p.relative_to(folder).as_posix().casefold()):
        rel = path.relative_to(folder).as_posix()
        if path.is_symlink():
            raise RepairRefused(f"Refusing to move chart folder containing a symbolic link: {rel}")
        if path.is_dir():
            result.append({"relative_path": rel, "type": "directory", "size": None, "sha256": None})
        elif path.is_file():
            result.append({"relative_path": rel, "type": "file", "size": path.stat().st_size, "sha256": _sha256(path)})
        else:
            raise RepairRefused(f"Refusing to move unsupported chart-folder entry: {rel}")
    return result


def _root_for(chart_root: Path | None) -> tuple[Path, dict[str, Any]]:
    if chart_root is not None:
        return chart_root.expanduser().resolve(), {"selection": "explicit_test_or_cli_root", "root": str(chart_root.expanduser().resolve())}
    root, diagnostic = _select_trade_chart_root("C")
    if root is None:
        raise RepairRefused(f"Could not resolve configured local CRYPTO chart root: {diagnostic}")
    return Path(root).expanduser().resolve(), diagnostic


def _build_plan(workbook_path: Path, json_path: Path, chart_root: Path | None, recovery_root: Path) -> dict[str, Any]:
    workbook_path = workbook_path.expanduser().resolve()
    json_path = json_path.expanduser().resolve()
    root, root_selection = _root_for(chart_root)
    workbook_row = _workbook_row(workbook_path)
    json_row = _json_row(json_path)
    if workbook_row["row_id"] != json_row["row_id"]:
        raise RepairRefused("Workbook and JSON identify different C377 trades")
    source = root / "2026" / "APRIL" / EXPECTED_TRADE_NUMBER
    destination = root / "2026" / "OCTOBER" / EXPECTED_TRADE_NUMBER
    info = _hyperlink_info(workbook_path, workbook_row)
    if source.is_dir() and not destination.exists():
        if not _same_path_uri(info["target"], source):
            raise RepairRefused("C377's existing workbook link does not point to the verified APRIL source folder")
        action = "move_folder_and_retarget_existing_link"
        contents = _inventory(source)
    elif not source.exists() and destination.is_dir():
        if not _same_path_uri(info["target"], destination):
            raise RepairRefused("OCTOBER/C377 exists but the workbook link does not point to it; refusing to guess")
        action = "verified_no_op"
        contents = _inventory(destination)
    elif source.exists() and destination.exists():
        raise RepairRefused("Both APRIL/C377 and OCTOBER/C377 exist; refusing to merge, overwrite or duplicate")
    else:
        raise RepairRefused("Expected C377 source/destination state is incomplete or ambiguous")
    if not contents:
        raise RepairRefused("C377 chart folder is empty")
    workbook_hash = _sha256(workbook_path)
    return {
        "action": action,
        "trade": {"row_id": EXPECTED_ROW_ID, "trade_number": EXPECTED_TRADE_NUMBER,
                  "account": EXPECTED_ACCOUNT, "symbol": EXPECTED_SYMBOL,
                  "open_time": EXPECTED_OPEN.isoformat(), "close_time": EXPECTED_CLOSE.isoformat(),
                  "workbook_row": workbook_row["row"], "hyperlink_cell": info["cell_ref"]},
        "workbook_path": str(workbook_path), "workbook_sha256": workbook_hash,
        "json_path": str(json_path), "chart_root": str(root), "chart_root_selection": root_selection,
        "source_folder": str(source), "destination_folder": str(destination),
        "existing_link": info["target"], "existing_relationship_id": info["relationship_id"],
        "existing_relationship_uses": info["relationship_uses"],
        "new_link": destination.resolve().as_uri(), "inventory": contents,
        "recovery_root": str(recovery_root.expanduser().resolve()),
    }


def plan_digest(plan: dict[str, Any]) -> str:
    return hashlib.sha256(_canonical_json(plan)).hexdigest()


def _can_open_exclusively(path: Path) -> bool:
    """Test current OS sharing/lock state without changing file contents."""
    if os.name == "nt":
        import ctypes
        from ctypes import wintypes
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        create_file = kernel.CreateFileW
        create_file.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD,
                                wintypes.LPVOID, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
        create_file.restype = wintypes.HANDLE
        handle = create_file(str(path), 0x80000000, 0, None, 3, 0x80, None)
        invalid = wintypes.HANDLE(-1).value
        if handle == invalid:
            error = ctypes.get_last_error()
            if error in {5, 32, 33}:
                return False
            raise OSError(error, ctypes.FormatError(error), str(path))
        kernel.CloseHandle(handle)
        return True
    handle = path.open("rb")
    try:
        import fcntl
        try:
            fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError as exc:
            if exc.errno in {errno.EACCES, errno.EAGAIN}:
                return False
            raise
        finally:
            try:
                fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
            except OSError:
                pass
        return True
    finally:
        handle.close()


def _assert_workbook_available(path: Path) -> None:
    owner_file = path.with_name("~$" + path.name)
    if owner_file.exists() and not _can_open_exclusively(owner_file):
        raise RepairRefused(f"Excel owner file exists; close the workbook before repair: {owner_file}")
    if not os.access(path, os.R_OK | os.W_OK) or not _can_open_exclusively(path):
        raise RepairRefused("Trading Journal workbook is not readable and writable")


def _acquire_existing_gate(path: Path) -> Callable[[], None]:
    from render import master_service
    result = master_service._reserve_master_journal_workbook_sync(
        path, f"chart-folder-repair-{uuid.uuid4().hex[:12]}", "trade_folder_repair", wait_seconds=1.0,
    )
    if result is not None:
        raise RepairRefused("Journal workbook transaction is busy; no repair changes were made")
    return master_service._release_master_journal_workbook_sync


def _copy_backup(plan: dict[str, Any], recovery_root: Path) -> tuple[Path, Path, Path]:
    recovery_root.mkdir(parents=True, exist_ok=True)
    backup_dir = recovery_root / f"C377-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{uuid.uuid4().hex[:8]}"
    backup_dir.mkdir()
    workbook_backup = backup_dir / "Trading Journal.xlsx"
    folder_backup = backup_dir / "C377"
    shutil.copy2(plan["workbook_path"], workbook_backup)
    shutil.copytree(plan["source_folder"], folder_backup, symlinks=True)
    if _sha256(workbook_backup) != plan["workbook_sha256"]:
        raise RepairRefused("Workbook backup hash does not match the planned source bytes")
    if _inventory(folder_backup) != plan["inventory"]:
        raise RepairRefused("C377 backup inventory does not match the planned source contents")
    manifest = {
        "trade_row_id": EXPECTED_ROW_ID,
        "trade_number": EXPECTED_TRADE_NUMBER,
        "source_folder": plan["source_folder"],
        "destination_folder": plan["destination_folder"],
        "workbook_path": plan["workbook_path"],
        "workbook_sha256": plan["workbook_sha256"],
        "folder_inventory": plan["inventory"],
        "backup_workbook_sha256": _sha256(workbook_backup),
        "backup_folder_inventory": _inventory(folder_backup),
    }
    (backup_dir / "manifest.json").write_bytes(_canonical_json(manifest) + b"\n")
    return backup_dir, workbook_backup, folder_backup


def _patch_package(source_workbook: Path, candidate: Path, row: dict[str, Any], link_info: dict[str, Any], new_target: str) -> dict[str, Any]:
    ET.register_namespace("", MAIN_NS)
    ET.register_namespace("r", DOC_REL_NS)
    ET.register_namespace("", PKG_REL_NS)
    with zipfile.ZipFile(source_workbook, "r") as zin:
        sheet_part, rel_part = _worksheet_parts(zin)
        if sheet_part != link_info["sheet_part"] or rel_part != link_info["rels_part"]:
            raise RepairRefused("Trade Log package relationships changed after planning")
        source_sheet_bytes = zin.read(sheet_part)
        source_rels_bytes = zin.read(rel_part)
        sheet_root = ET.fromstring(source_sheet_bytes)
        rels_root = ET.fromstring(source_rels_bytes)
        links = [h for h in sheet_root.findall(f"{_Q(MAIN_NS, 'hyperlinks')}/{_Q(MAIN_NS, 'hyperlink')}") if h.attrib.get("ref") == link_info["cell_ref"]]
        if len(links) != 1 or links[0].attrib.get(_Q(DOC_REL_NS, "id")) != link_info["relationship_id"]:
            raise RepairRefused("C377 hyperlink changed after the dry-run plan")
        target_link = links[0]
        old_rid = link_info["relationship_id"]
        old_rel = next((r for r in rels_root.findall(_Q(PKG_REL_NS, "Relationship")) if r.attrib.get("Id") == old_rid), None)
        if old_rel is None or old_rel.attrib.get("Target") != link_info["target"]:
            raise RepairRefused("C377 hyperlink relationship target changed after planning")
        if link_info["relationship_uses"] == 1:
            old_rel.attrib["Target"] = new_target
            new_rid = old_rid
        else:
            ids = {r.attrib.get("Id", "") for r in rels_root.findall(_Q(PKG_REL_NS, "Relationship"))}
            n = 1
            while f"rId{n}" in ids:
                n += 1
            new_rid = f"rId{n}"
            target_link.attrib[_Q(DOC_REL_NS, "id")] = new_rid
            ET.SubElement(rels_root, _Q(PKG_REL_NS, "Relationship"), {
                "Id": new_rid,
                "Type": old_rel.attrib["Type"],
                "Target": new_target,
                "TargetMode": "External",
            })
        patched_sheet = ET.tostring(sheet_root, encoding="utf-8", xml_declaration=True)
        ET.register_namespace("", PKG_REL_NS)
        patched_rels = ET.tostring(rels_root, encoding="utf-8", xml_declaration=True)
        with zipfile.ZipFile(candidate, "w") as zout:
            zout.comment = zin.comment
            for item in zin.infolist():
                data = patched_sheet if item.filename == sheet_part else patched_rels if item.filename == rel_part else zin.read(item.filename)
                zout.writestr(item, data)
    shutil.copystat(source_workbook, candidate)
    return {"sheet_part": sheet_part, "rels_part": rel_part, "old_relationship_id": old_rid,
            "new_relationship_id": new_rid, "relationship_was_shared": link_info["relationship_uses"] > 1}


def _xml_without_cell_hyperlink(data: bytes, cell_ref: str) -> bytes:
    root = ET.fromstring(data)
    parent = root.find(_Q(MAIN_NS, "hyperlinks"))
    if parent is not None:
        for link in list(parent):
            if link.attrib.get("ref") == cell_ref:
                parent.remove(link)
    ET.register_namespace("", MAIN_NS)
    ET.register_namespace("r", DOC_REL_NS)
    return ET.tostring(root, encoding="utf-8")


def _xml_without_relationships(data: bytes, ignored_ids: set[str]) -> bytes:
    root = ET.fromstring(data)
    for rel in list(root.findall(_Q(PKG_REL_NS, "Relationship"))):
        if rel.attrib.get("Id") in ignored_ids:
            root.remove(rel)
    ET.register_namespace("", PKG_REL_NS)
    return ET.tostring(root, encoding="utf-8")


def _verify_candidate(original: Path, candidate: Path, parts: dict[str, Any], cell_ref: str, link_info: dict[str, Any], new_target: str) -> None:
    with zipfile.ZipFile(original, "r") as old, zipfile.ZipFile(candidate, "r") as new:
        if new.testzip() is not None:
            raise RepairRefused("Candidate workbook ZIP integrity check failed")
        if set(old.namelist()) != set(new.namelist()):
            raise RepairRefused("Candidate workbook changed the package-part inventory")
        for name in old.namelist():
            if name not in {parts["sheet_part"], parts["rels_part"]} and old.read(name) != new.read(name):
                raise RepairRefused(f"Candidate changed unrelated workbook package part: {name}")
        if _xml_without_cell_hyperlink(old.read(parts["sheet_part"]), cell_ref) != _xml_without_cell_hyperlink(new.read(parts["sheet_part"]), cell_ref):
            raise RepairRefused("Candidate changed Trade Log cells or presentation outside C377's hyperlink")
        old_rid = link_info["relationship_id"]
        if parts["relationship_was_shared"]:
            old_ignore: set[str] = set()
            new_ignore = {parts["new_relationship_id"]}
        else:
            old_ignore = {old_rid}
            new_ignore = {parts["new_relationship_id"]}
        if _xml_without_relationships(old.read(parts["rels_part"]), old_ignore) != _xml_without_relationships(new.read(parts["rels_part"]), new_ignore):
            raise RepairRefused("Candidate changed worksheet relationships outside C377's existing hyperlink")
        sheet = ET.fromstring(new.read(parts["sheet_part"]))
        link_nodes = [h for h in sheet.findall(f"{_Q(MAIN_NS, 'hyperlinks')}/{_Q(MAIN_NS, 'hyperlink')}") if h.attrib.get("ref") == cell_ref]
        if len(link_nodes) != 1 or link_nodes[0].attrib.get(_Q(DOC_REL_NS, "id")) != parts["new_relationship_id"]:
            raise RepairRefused("Candidate C377 hyperlink reference was not installed exactly once")
        rels = ET.fromstring(new.read(parts["rels_part"]))
        rel = next((r for r in rels.findall(_Q(PKG_REL_NS, "Relationship")) if r.attrib.get("Id") == parts["new_relationship_id"]), None)
        if rel is None or rel.attrib.get("Target") != new_target or rel.attrib.get("TargetMode") != "External":
            raise RepairRefused("Candidate C377 hyperlink does not target OCTOBER/C377")


def _verify_saved_workbook(workbook: Path, new_folder: Path) -> dict[str, Any]:
    row = _workbook_row(workbook)
    info = _hyperlink_info(workbook, row)
    if not _same_path_uri(info["target"], new_folder):
        raise RepairRefused("Saved C377 hyperlink does not resolve to OCTOBER/C377")
    return {"row": row["row"], "trade_number": row["trade_number"], "row_id": row["row_id"],
            "account": row["account"], "symbol": row["symbol"], "open_time": row["open_time"].isoformat(),
            "close_time": row["close_time"].isoformat(), "hyperlink": info["target"]}


def _acquire_gate(path: Path, reserve: Callable[[Path], Callable[[], None]] | None) -> Callable[[], None]:
    return reserve(path) if reserve is not None else _acquire_existing_gate(path)


def repair(
    workbook_path: Path,
    json_path: Path,
    *,
    chart_root: Path | None = None,
    recovery_root: Path | None = None,
    apply: bool = False,
    confirm_plan_sha256: str | None = None,
    reserve: Callable[[Path], Callable[[], None]] | None = None,
) -> dict[str, Any]:
    """Return a verified plan by default; mutate only with the matching plan digest."""
    repo = Path(__file__).resolve().parents[1]
    recovery = recovery_root or (repo / ".job7f_recovery")
    plan = _build_plan(Path(workbook_path), Path(json_path), chart_root, Path(recovery))
    digest = plan_digest(plan)
    result = {"plan": plan, "plan_sha256": digest, "applied": False}
    if not apply:
        return result
    if confirm_plan_sha256 != digest:
        raise RepairRefused("Apply requires the SHA-256 from the reviewed dry-run plan")
    if plan["action"] == "verified_no_op":
        _TRADE_FOLDER_INDEX_CACHE.clear()
        return {**result, "verified_no_op": True, "verified_workbook": _verify_saved_workbook(Path(workbook_path), Path(plan["destination_folder"]))}

    workbook = Path(workbook_path).expanduser().resolve()
    root = Path(plan["chart_root"])
    source = Path(plan["source_folder"])
    destination = Path(plan["destination_folder"])
    month = destination.parent
    release = _acquire_gate(workbook, reserve)
    backup_dir: Path | None = None
    candidate: Path | None = None
    moved = False
    installed = False
    candidate_hash = ""
    created_month = False
    rollback: dict[str, Any] = {"workbook_restored": False, "folder_restored": False, "complete": True}
    try:
        _assert_workbook_available(workbook)
        current = _build_plan(workbook, Path(plan["json_path"]), chart_root, Path(plan["recovery_root"]))
        if plan_digest(current) != digest:
            raise RepairRefused("C377, its link, chart files or workbook changed after the dry-run plan")
        backup_dir, workbook_backup, folder_backup = _copy_backup(plan, Path(plan["recovery_root"]))
        link_info = _hyperlink_info(workbook, _workbook_row(workbook))
        candidate = workbook.with_name(f".{workbook.stem}.C377-{uuid.uuid4().hex}.candidate{workbook.suffix}")
        parts = _patch_package(workbook, candidate, _workbook_row(workbook), link_info, plan["new_link"])
        _verify_candidate(workbook, candidate, parts, link_info["cell_ref"], link_info, plan["new_link"])
        candidate_hash = _sha256(candidate)
        if not month.exists():
            month.mkdir()
            created_month = True
        if destination.exists() or not source.is_dir():
            raise RepairRefused("C377 source/destination changed before the atomic folder move")
        if _inventory(source) != plan["inventory"]:
            raise RepairRefused("C377 contents changed after backup; refusing to move them")
        source.rename(destination)
        moved = True
        if _inventory(destination) != plan["inventory"]:
            raise RepairRefused("Moved C377 inventory differs from the verified backup")
        os.replace(candidate, workbook)
        installed = True
        verified = _verify_saved_workbook(workbook, destination)
        if _sha256(workbook) != candidate_hash or _inventory(destination) != plan["inventory"]:
            raise RepairRefused("Final workbook bytes or moved chart files changed during verification")
        april = source.parent
        april_removed = False
        try:
            april.rmdir()
            april_removed = True
        except OSError:
            pass
        _TRADE_FOLDER_INDEX_CACHE.clear()
        return {**result, "applied": True, "verified_no_op": False, "backup_dir": str(backup_dir),
                "workbook_backup": str(workbook_backup), "chart_folder_backup": str(folder_backup),
                "verified_workbook": verified, "inventory": plan["inventory"], "april_parent_removed_if_empty": april_removed,
                "workbook_sha256": _sha256(workbook), "candidate_package_parts": parts}
    except Exception as exc:
        rollback_errors = []
        if installed and backup_dir is not None:
            if workbook.exists() and _sha256(workbook) == candidate_hash:
                restore_candidate = workbook.with_name(f".{workbook.name}.C377-restore-{uuid.uuid4().hex}.tmp")
                shutil.copy2(backup_dir / "Trading Journal.xlsx", restore_candidate)
                os.replace(restore_candidate, workbook)
                rollback["workbook_restored"] = _sha256(workbook) == plan["workbook_sha256"]
            else:
                rollback_errors.append("workbook changed after candidate install; backup was not written over it")
        if moved:
            if destination.is_dir() and not source.exists() and _inventory(destination) == plan["inventory"]:
                source.parent.mkdir(parents=True, exist_ok=True)
                destination.rename(source)
                rollback["folder_restored"] = _inventory(source) == plan["inventory"]
            else:
                rollback_errors.append("moved folder changed or source reappeared; concurrent contents were preserved")
        if created_month:
            try:
                month.rmdir()
            except OSError:
                pass
        rollback["complete"] = not rollback_errors and (not installed or rollback["workbook_restored"]) and (not moved or rollback["folder_restored"])
        rollback["errors"] = rollback_errors
        raise RepairRefused(f"C377 repair failed: {exc}; rollback={json.dumps(rollback, sort_keys=True)}; backup={backup_dir}") from exc
    finally:
        if candidate is not None and candidate.exists():
            candidate.unlink()
        release()


def main(argv: list[str] | None = None) -> int:
    repo = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=repo / "journal" / "Trading Journal.xlsx")
    parser.add_argument("--json", type=Path, default=repo / "render" / "data" / "trading_journal.json")
    parser.add_argument("--chart-root", type=Path, default=None)
    parser.add_argument("--recovery-root", type=Path, default=repo / ".job7f_recovery")
    parser.add_argument("--apply", action="store_true", help="apply only the exact C377 mapping in a reviewed plan")
    parser.add_argument("--confirm-plan-sha256", default="")
    args = parser.parse_args(argv)
    try:
        result = repair(args.workbook, args.json, chart_root=args.chart_root, recovery_root=args.recovery_root,
                        apply=args.apply, confirm_plan_sha256=args.confirm_plan_sha256 or None)
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False, indent=2))
        return 2
    print(json.dumps({"ok": True, **result}, ensure_ascii=False, indent=2, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
