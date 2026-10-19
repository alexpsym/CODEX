#!/usr/bin/env python3
from __future__ import annotations
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SNAPSHOT_PATH = ROOT / "render" / "data" / "trading_journal_view_cache.json"
REPORT_PATH = ROOT / "render" / "data" / "trading_journal_report.html"


def legacy_main() -> int:
    if not SNAPSHOT_PATH.exists():
        print(f"Snapshot not found: {SNAPSHOT_PATH}")
        return 1
    payload = json.loads(SNAPSHOT_PATH.read_text(encoding="utf-8"))
    items = payload.get("items") or []
    stats = payload.get("stats") or {}
    diagnostics = payload.get("diagnostics") or {}
    html = f"""<!doctype html>
<html><head><meta charset='utf-8'><title>Trading Journal Report</title>
<style>body{{font-family:Arial,sans-serif;margin:20px}}table{{border-collapse:collapse;width:100%}}th,td{{border:1px solid #ddd;padding:6px;font-size:12px}}th{{background:#f5f5f5}}</style>
</head><body>
<h1>Trading Journal Report</h1>
<p>Generated snapshot: {payload.get('generated_at','')}</p>
<p>Total rows: {len(items)}</p>
<h2>Overview</h2>
<pre>{json.dumps((stats or {}).get('groups', {}).get('overview', {}), indent=2)}</pre>
<h2>Diagnostics</h2>
<pre>{json.dumps(diagnostics, indent=2)}</pre>
<h2>Rows</h2>
<table>
<tr><th>Close Time</th><th>Account</th><th>Symbol</th><th>Side</th><th>Net Profit</th><th>Result %</th><th>Balance After</th></tr>
{''.join(f"<tr><td>{r.get('close_time') or r.get('open_time') or ''}</td><td>{r.get('account_label') or r.get('account') or ''}</td><td>{r.get('symbol') or ''}</td><td>{r.get('side') or ''}</td><td>{r.get('net_profit') or ''}</td><td>{r.get('result_pct') or ''}</td><td>{r.get('analysis_balance_after_trade') or r.get('balance_after_trade') or r.get('cashflow_new_balance') or ''}</td></tr>" for r in items)}
</table>
</body></html>"""
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text(html, encoding="utf-8")
    print(f"Wrote {REPORT_PATH}")
    return 0


# Print primitives intentionally do not import the workbook generator or service.
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
import hashlib
import io
import re
import zipfile
import xml.etree.ElementTree as ET

NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
REL = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
TECHNICAL_SHEETS = {'_Target R Metadata': 'Internal target-R provenance; not a business report'}


@dataclass(frozen=True)
class Capture:
    source: str
    sha256: str
    size: int
    mtime_ns: int
    data: bytes = field(repr=False)
    temporary_path: str = ''


def capture_workbook(source: Path, temporary_directory: Path) -> Capture:
    """One attempt, one byte read; never opens the source for writing or retries."""
    before = source.stat()
    data = source.read_bytes()
    after = source.stat()
    digest = hashlib.sha256(data).hexdigest()
    target = temporary_directory / 'captured-workbook.xlsx'
    with target.open('xb') as stream:
        stream.write(data)
    if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns) or len(data) != before.st_size:
        raise RuntimeError(f'Source changed during the sole capture; preserved bytes: {target}; SHA256={digest}')
    return Capture(str(source), digest, len(data), before.st_mtime_ns, data, str(target))


@dataclass(frozen=True)
class PrintCell:
    coordinate: str
    value: object
    number_format: str
    formula: str | None = None
    missing_cache: bool = False


@dataclass
class SheetSnapshot:
    name: str
    visibility: str
    cells: list[PrintCell]
    bounds: str
    merges: list[str]
    tables: list[dict]
    objects: dict
    exclusion: str | None
    format_limitations: list[str]


def read_print_snapshot(capture: Capture) -> list[SheetSnapshot]:
    """Parse inert OOXML only; populated cells determine bounds, not dimensions.

    Resource guards fail honestly rather than scanning unbounded formatting padding.
    Formula expressions are retained as provenance, never substituted for results.
    """
    from openpyxl.styles.stylesheet import Stylesheet
    from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format
    from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
    from openpyxl.utils.datetime import from_excel, CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900
    import posixpath
    with zipfile.ZipFile(io.BytesIO(capture.data)) as archive:
        if sum(i.file_size for i in archive.infolist()) > 200_000_000:
            raise ValueError('Snapshot exceeds 200 MB expanded scan budget')
        def xml(path):
            return ET.fromstring(archive.read(path))
        def relations(path):
            folder, base = posixpath.split(path)
            relpath = f'{folder}/_rels/{base}.rels'
            if relpath not in archive.namelist():
                return {}
            return {r.attrib['Id']: (r.attrib['Type'].split('/')[-1],
                    posixpath.normpath(posixpath.join(folder, r.attrib['Target'])).lstrip('/')
                    if not r.attrib['Target'].startswith('/') else r.attrib['Target'].lstrip('/'))
                    for r in xml(relpath) if r.attrib.get('TargetMode') != 'External'}
        shared = []
        if 'xl/sharedStrings.xml' in archive.namelist():
            shared = [''.join(n.itertext()) for n in xml('xl/sharedStrings.xml')]
        styles = Stylesheet.from_tree(xml('xl/styles.xml'))
        formats = dict(BUILTIN_FORMATS)
        formats.update({n.numFmtId: n.formatCode for n in styles.numFmts.numFmt})
        workbook = xml('xl/workbook.xml')
        props = workbook.find('s:workbookPr', NS)
        epoch = CALENDAR_MAC_1904 if props is not None and props.get('date1904') in ('1', 'true') else CALENDAR_WINDOWS_1900
        links = relations('xl/workbook.xml')
        result = []
        scanned = 0
        for sheet in workbook.find('s:sheets', NS):
            path = links[sheet.attrib[REL]][1]
            root = xml(path)
            cells, limitations = [], set()
            for node in root.findall('.//s:sheetData/s:row/s:c', NS):
                scanned += 1
                if scanned > 1_000_000:
                    raise ValueError('Snapshot exceeds one million serialized cells')
                formula = node.find('s:f', NS)
                value = node.find('s:v', NS)
                inline = node.find('s:is', NS)
                if formula is None and value is None and inline is None:
                    continue
                text = value.text if value is not None else None
                kind = node.get('t', 'n')
                missing = formula is not None and text is None
                if inline is not None:
                    raw = ''.join(t.text or '' for t in inline.findall('.//s:t', NS))
                elif kind == 's' and text is not None:
                    raw = shared[int(text)]
                elif kind == 'b' and text is not None:
                    raw = text == '1'
                elif kind in ('str', 'e', 'd'):
                    raw = text
                else:
                    raw = Decimal(text) if text is not None else None
                fmt = formats.get(styles.cellXfs.xf[int(node.get('s', '0'))].numFmtId, 'General')
                if isinstance(raw, Decimal) and is_date_format(fmt):
                    raw = from_excel(float(raw), epoch=epoch)
                elif fmt != 'General':
                    limitations.add(fmt)
                cells.append(PrintCell(node.attrib['r'], raw, fmt,
                                       formula.text or '' if formula is not None else None, missing))
            coordinates = [coordinate_to_tuple(c.coordinate) for c in cells]
            bounds = ''
            if coordinates:
                rows, cols = zip(*coordinates)
                bounds = f'{get_column_letter(min(cols))}{min(rows)}:{get_column_letter(max(cols))}{max(rows)}'
            rels = relations(path)
            tables = []
            objects = {'charts': 0, 'images': 0, 'comments': 0, 'other_relationships': []}
            for kind, target in rels.values():
                if kind == 'table':
                    table = xml(target)
                    tables.append({'name': table.get('displayName'), 'ref': table.get('ref')})
                elif kind == 'drawing':
                    for drawing_kind, _ in relations(target).values():
                        if drawing_kind in ('chart', 'image'):
                            objects['charts' if drawing_kind == 'chart' else 'images'] += 1
                        else:
                            objects['other_relationships'].append(drawing_kind)
                elif kind == 'comments':
                    objects['comments'] += len(xml(target).findall('.//s:comment', NS))
                else:
                    objects['other_relationships'].append(kind)
            result.append(SheetSnapshot(sheet.get('name'), sheet.get('state', 'visible'), cells, bounds,
                [m.get('ref') for m in root.findall('s:mergeCells/s:mergeCell', NS)], tables, objects,
                TECHNICAL_SHEETS.get(sheet.get('name')), sorted(limitations)))
        return result


def display_value(cell: PrintCell) -> str:
    """Lossless fallback; complex Excel display formats require Phase 6B work."""
    if cell.missing_cache:
        return '[stored result unavailable]'
    value = cell.value
    if value is None:
        return '[missing]'
    if isinstance(value, bool):
        return 'Yes' if value else 'No'
    if isinstance(value, (date, datetime)):
        return value.isoformat(sep=' ') if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, (Decimal, int, float)):
        number = Decimal(str(value))
        if '%' in cell.number_format:
            return f'{number * 100:f}%'
        text = f'{number:f}'
        currency = next((c for c in '$£€' if c in cell.number_format), '')
        suffixes = re.findall(r'"([^"#0]*)"', cell.number_format)
        unit = ''.join(suffixes).strip()
        if not unit and re.search(r'\\R', cell.number_format):
            unit = 'R'
        return currency + text + ((' ' + unit) if unit and unit not in currency else '')
    return str(value)


@dataclass(frozen=True)
class PrintPanel:
    sheet: str
    first_column: int
    last_column: int
    first_row: int
    last_row: int
    header_rows: int
    identity_columns: int
    orientation: str
    panel: int
    panels: int
    block: int
    zoom: int = 92

    @property
    def context(self):
        from openpyxl.utils import get_column_letter as letter
        return (f'{self.sheet} | rows {self.first_row}-{self.last_row} | '
                f'panel {self.panel}/{self.panels} | '
                f'columns {letter(self.first_column)}:{letter(self.last_column)} | block {self.block}')


def column_points(sheet, column):
    """Excel character widths to conservative print points; no font scaling here."""
    from openpyxl.utils import get_column_letter
    dimension = sheet.column_dimensions.get(get_column_letter(column))
    width = dimension.width if dimension is not None else (sheet.sheet_format.defaultColWidth or 13)
    return (width * 7 + 5) * .75


def plan_panels(workbook, specifications):
    """Consecutive source columns; finish horizontal panels before next row block.

    specifications: (sheet, last-column, header-count, repeated-columns,
    orientation, [(first-data-row,last-data-row), ...]). Merged header groups
    stay whole when they fit; oversized groups explicitly repeat on continuation.
    """
    from openpyxl.utils.cell import range_boundaries
    result = []
    for name, last, headers, identities, orientation, blocks in specifications:
        sheet = workbook[name]
        usable = (841.89 if orientation == 'landscape' else 595.276) - 64
        budget = usable / .92 - sum(column_points(sheet,c) for c in range(1,identities+1))
        bands = []
        first = identities+1
        while first <= last:
            end, width = first-1, 0
            while end < last and width+column_points(sheet,end+1) <= budget:
                end += 1
                width += column_points(sheet,end)
            if end < first:
                raise ValueError(f'{name}: source column {first} exceeds readable A4 width')
            for merged in sheet.merged_cells.ranges:
                a,b,c,d = range_boundaries(str(merged))
                if b <= headers and first < a <= end < c:
                    end = a-1
            bands.append((first,end))
            first = end+1
        for block,(top,bottom) in enumerate(blocks,1):
            for panel,(left,right) in enumerate(bands,1):
                result.append(PrintPanel(name,left,right,top,bottom,headers,identities,
                                         orientation,panel,len(bands),block))
    return result


def grayscale_rgb(rgb):
    value = rgb[-6:]
    r,g,b = (int(value[i:i+2],16) for i in (0,2,4))
    gray = round(.299*r+.587*g+.114*b)
    return f'FF{gray:02X}{gray:02X}{gray:02X}'


def prepare_print_copy(capture: Capture, directory: Path) -> Path:
    """Surgical ZIP edits: preserve parts/styles, freeze exact cached results.

    No openpyxl save/rebuild and no native source open. Fonts, borders, merged
    ranges, dimensions and all non-colour style properties survive unchanged.
    """
    from openpyxl.styles.colors import COLOR_INDEX
    from xml.dom import minidom
    directory = directory.resolve()
    target = directory/'print-copy.xlsx'
    if Path(capture.source).resolve() == target:
        raise ValueError('Print-copy path must differ from source')
    with zipfile.ZipFile(io.BytesIO(capture.data)) as source:
        theme = ET.fromstring(source.read('xl/theme/theme1.xml')) if 'xl/theme/theme1.xml' in source.namelist() else None
        palette = []
        if theme is not None:
            ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
            scheme=theme.find('.//a:clrScheme',ns)
            colors={x.tag.split('}')[-1]: next(iter(x)).get('val',next(iter(x)).get('lastClr')) for x in scheme}
            # System colours use lastClr rather than symbolic window/windowText.
            for x in scheme:
                child=next(iter(x))
                if child.get('lastClr'):colors[x.tag.split('}')[-1]]=child.get('lastClr')
            palette=[colors[k] for k in ('lt1','dk1','lt2','dk2','accent1','accent2','accent3','accent4','accent5','accent6','hlink','folHlink')]
        with zipfile.ZipFile(target,'x',compression=zipfile.ZIP_DEFLATED) as output:
            for info in source.infolist():
                if info.filename=='xl/calcChain.xml':
                    continue
                data=source.read(info.filename)
                if info.filename.endswith('.xml') or info.filename.endswith('.rels'):
                    # Office compatibility prefixes also occur inside STRING
                    # values (mc:Ignorable, Choice/@Requires, etc.). ElementTree
                    # drops declarations not used in expanded element names.
                    # DOM retains original prefixes/declarations, including
                    # nested/rebound scopes, while we edit only owned nodes.
                    document=minidom.parseString(data)
                    root=document.documentElement
                    changed=False
                    if info.filename.startswith('xl/worksheets/sheet'):
                        for cell in root.getElementsByTagNameNS(NS['s'],'c'):
                            formula=next((n for n in cell.childNodes if n.nodeType==n.ELEMENT_NODE and n.localName=='f' and n.namespaceURI==NS['s']),None)
                            if formula is not None:
                                cell.removeChild(formula)
                                value=next((n for n in cell.childNodes if n.nodeType==n.ELEMENT_NODE and n.localName=='v' and n.namespaceURI==NS['s']),None)
                                if value is None or not value.childNodes:
                                    if value is not None:cell.removeChild(value)
                                    cell.setAttribute('t','inlineStr')
                                    prefix=(cell.prefix+':') if cell.prefix else ''
                                    inline=document.createElementNS(NS['s'],prefix+'is')
                                    text=document.createElementNS(NS['s'],prefix+'t')
                                    text.appendChild(document.createTextNode('[stored result unavailable]'))
                                    inline.appendChild(text);cell.appendChild(inline)
                                changed=True
                    for node in [root,*root.getElementsByTagName('*')]:
                        kind=node.localName
                        # Comments/unsupported extension objects remain byte
                        # identical; print-cell styles and CF colours are the
                        # only grayscale targets in this preparation stage.
                        if (info.filename=='xl/styles.xml' or info.filename.startswith('xl/worksheets/')) and kind in ('color','fgColor','bgColor') and node.namespaceURI==NS['s']:
                            rgb=node.getAttribute('rgb') or None
                            if rgb is None and node.hasAttribute('theme') and palette:
                                rgb=palette[int(node.getAttribute('theme'))]
                            if rgb is None and node.hasAttribute('indexed'):
                                index=int(node.getAttribute('indexed'))
                                if index<len(COLOR_INDEX):rgb=COLOR_INDEX[index]
                            if rgb is not None:
                                tint=float(node.getAttribute('tint') or '0')
                                base=[int(rgb[-6:][i:i+2],16) for i in (0,2,4)]
                                tinted=[round(v*(1+tint) if tint<0 else v*(1-tint)+255*tint) for v in base]
                                rgb=''.join(f'{v:02X}' for v in tinted)
                                for attribute in ('rgb','theme','indexed','auto','tint'):
                                    if node.hasAttribute(attribute):node.removeAttribute(attribute)
                                node.setAttribute('rgb',grayscale_rgb(rgb));changed=True
                        for flag in ('refreshOnLoad','backgroundRefresh','enableRefresh'):
                            if node.hasAttribute(flag):node.setAttribute(flag,'0');changed=True
                        if kind=='calcPr':
                            for attribute,value in {'calcMode':'manual','fullCalcOnLoad':'0','forceFullCalc':'0','calcOnSave':'0'}.items():node.setAttribute(attribute,value)
                            changed=True
                        if kind=='workbookPr':node.setAttribute('updateLinks','never');changed=True
                    for parent in [root,*root.getElementsByTagName('*')]:
                        for child in list(parent.childNodes):
                            if child.nodeType==child.ELEMENT_NODE and (child.getAttribute('Type').endswith('/calcChain') or child.getAttribute('PartName')=='/xl/calcChain.xml'):
                                parent.removeChild(child);changed=True
                    if changed:data=document.toxml(encoding='utf-8')
                    document.unlink()
                output.writestr(info,data)
    return target


def validate_virtual_context(name):
    # Only exact installed names and cached ports supplied by the host inventory.
    if not re.fullmatch(r'Microsoft (?:Print to PDF|XPS Document Writer) on (?:Ne\d+|PORTPROMPT):',name):
        raise ValueError('Only verified installed local Microsoft PDF/XPS contexts are eligible')
    return name


def page_setup_values(setup):
    values={}
    for key in ('PaperSize','Orientation','Zoom','FitToPagesWide','FitToPagesTall'):
        try:
            value=getattr(setup,key)
            values[key]={'value':value,'type':type(value).__name__}
        except Exception as error:
            values[key]={'error':repr(error)}
    return values


def apply_native_page_setup(excel, sheet, panel, ordinal, diagnostic=None):
    from openpyxl.utils import get_column_letter as letter
    setup=sheet.PageSetup
    expected={'PaperSize':9,'Orientation':2 if panel.orientation=='landscape' else 1,
              'Zoom':panel.zoom,'FitToPagesWide':False,'FitToPagesTall':False}
    evidence={'panel':f'P{ordinal:03d}','sheet':panel.sheet,
              'expected':{k:{'value':v,'type':type(v).__name__} for k,v in expected.items()},
              'before':page_setup_values(setup),'stage':'before configuration'}
    def record():
        if diagnostic:
            try:diagnostic(evidence.copy())
            except Exception:pass  # Diagnostics must not replace a native failure.
    record()
    excel.PrintCommunication=False
    try:
        setup.PaperSize=9;setup.Orientation=2 if panel.orientation=='landscape' else 1
        setup.LeftMargin=28.35;setup.RightMargin=28.35
        setup.TopMargin=36;setup.BottomMargin=36
        setup.HeaderMargin=14;setup.FooterMargin=14
        setup.FitToPagesWide=False;setup.FitToPagesTall=False;setup.Zoom=panel.zoom
        setup.PrintTitleRows=f'$1:${panel.header_rows}'
        setup.PrintTitleColumns=f'$A:${letter(panel.identity_columns)}'
        setup.PrintArea=f'${letter(panel.first_column)}${panel.first_row}:${letter(panel.last_column)}${panel.last_row}'
        setup.Order=2;setup.BlackAndWhite=False;setup.PrintHeadings=True
        setup.CenterFooter=f'&8 P{ordinal:03d} | '+panel.context.replace('&','&&')+' | page &P'
    finally:
        excel.PrintCommunication=True
    evidence.update(stage='after PrintCommunication commit',actual=page_setup_values(setup),
                    fresh=page_setup_values(sheet.PageSetup))
    record()
    actual=evidence['fresh']
    if any(actual[k].get('value')!=expected[k] or
           (k=='Zoom' and isinstance(actual[k].get('value'),bool))
           for k in ('PaperSize','Orientation','Zoom')):
        raise RuntimeError('Native A4 settings readback mismatch: '+json.dumps(evidence,default=str))


def native_preview(copy_path, directory, panels, virtual_context, factory=None, progress=None):
    """No-save native adapter. Bootstrap precedes every contextual setter.

    factory is a narrow test seam; production always uses DispatchEx, never an
    existing Excel instance. Caller runs this adapter in a bounded child.
    """
    from openpyxl import load_workbook
    directory=Path(directory).resolve();copy_path=Path(copy_path).resolve()
    if copy_path.parent!=directory or copy_path.name!='print-copy.xlsx':
        raise ValueError('Native open must target the owned print-copy.xlsx')
    context=validate_virtual_context(virtual_context)
    if factory is None:
        import win32com.client
        factory=lambda:win32com.client.DispatchEx('Excel.Application')
    excel=None;bootstrap=None;source=None;output=None
    def stage(value):
        if progress:progress(value)
    def diagnostic(value):
        with (directory/'page-setup.jsonl').open('a',encoding='utf-8') as stream:
            stream.write(json.dumps(value,default=str)+'\n')
    try:
        stage('DispatchEx and owned bootstrap')
        excel=factory()
        excel.Visible=False;excel.DisplayAlerts=False;excel.EnableEvents=False
        excel.AutomationSecurity=3;excel.AskToUpdateLinks=False
        bootstrap=excel.Workbooks.Add(-4167)
        if excel.Workbooks.Count<1 or excel.ActiveSheet is None:raise RuntimeError('Missing bootstrap context')
        excel.Calculation=-4135;excel.CalculateBeforeSave=False
        if excel.Calculation!=-4135 or excel.CalculateBeforeSave:raise RuntimeError('Unsafe calculation state')
        if excel.ActivePrinter!=context:excel.ActivePrinter=context
        if excel.ActivePrinter!=context:raise RuntimeError('Virtual context mismatch')
        stage('read-only print-copy Open')
        source=excel.Workbooks.Open(str(copy_path),UpdateLinks=0,ReadOnly=True,IgnoreReadOnlyRecommended=True,AddToMru=False)
        if excel.Calculation!=-4135:raise RuntimeError('Print copy changed manual calculation')
        # Return a native PDF via an output workbook of source-sheet copies. No
        # source values, formats, or objects are reconstructed from transactions.
        output=excel.Workbooks.Add(-4167);placeholder=output.Worksheets(1)
        stage('preserve source theme font scheme in output')
        font_scheme=directory/'source-font-scheme.xml'
        source.Theme.ThemeFontScheme.Save(str(font_scheme))
        output.Theme.ThemeFontScheme.Load(str(font_scheme))
        model=load_workbook(copy_path,data_only=True)
        try:
            for ordinal,panel in enumerate(panels,1):
                stage(f'P{ordinal:03d} source worksheet Copy')
                original=source.Worksheets(panel.sheet)
                count=output.Worksheets.Count
                # Positional arguments are required by late-bound Excel Copy;
                # a named After can be silently ignored by this COM dispatch.
                original.Copy(None,output.Worksheets(count))
                if output.Worksheets.Count!=count+1:
                    raise RuntimeError(f'P{ordinal:03d} worksheet Copy missed output workbook')
                sheet=output.Worksheets(output.Worksheets.Count)
                sheet.Name=f'P{ordinal:03d}'
                # Preserve grouped headers, repeating context if a group spans a
                # horizontal panel boundary. No trade is transposed.
                for merged in model[panel.sheet].merged_cells.ranges:
                    if merged.min_row<=panel.header_rows and merged.min_col<=panel.last_column and merged.max_col>=panel.first_column:
                        if merged.min_col<panel.first_column or merged.max_col>panel.last_column:
                            sheet.Range(str(merged)).UnMerge()
                            left=max(merged.min_col,panel.first_column);right=min(merged.max_col,panel.last_column)
                            original.Cells(merged.min_row,merged.min_col).Copy(Destination=sheet.Cells(merged.min_row,left))
                            sheet.Range(sheet.Cells(merged.min_row,left),sheet.Cells(merged.max_row,right)).Merge()
                # Output-only wrapping and local row expansion retain notes and
                # long identifiers; no global AutoFit or font replacement.
                body=sheet.Range(sheet.Cells(panel.first_row,1),sheet.Cells(panel.last_row,panel.last_column))
                body.WrapText=True;body.Rows.AutoFit()
                stage(f'P{ordinal:03d} PageSetup configuration and committed readback')
                apply_native_page_setup(excel,sheet,panel,ordinal,diagnostic)
            stage('delete output placeholder')
            placeholder.Delete()
            if len(panels)>24:raise ValueError('Proof plan exceeds 24 pages')
            native=directory/'native-preview.pdf'
            stage('ExportAsFixedFormat')
            output.ExportAsFixedFormat(0,str(native),OpenAfterPublish=False)
        finally:model.close()
        return native
    finally:
        # No Save/SaveAs/SaveCopyAs path exists; only our books are discarded.
        try:
            for book in (output,source,bootstrap):
                if book is not None:book.Close(SaveChanges=False)
        finally:
            if excel is not None:excel.Quit()


def finalize_native_pdf(native, directory, panels):
    """Validate native page geometry/content and arrange panels by explicit IDs."""
    import fitz
    target=Path(directory)/'journal_a4_excel_layout_preview.pdf'
    with fitz.open(native) as pdf:
        if len(pdf)>24:raise RuntimeError('Native pagination exceeded 24 pages')
        order=[]
        for index,page in enumerate(pdf):
            text=page.get_text()
            match=re.search(r'P(\d{3})\s*\|',text)
            if not match:raise RuntimeError('Native page lost explicit panel context')
            ordinal=int(match.group(1))
            if not 1<=ordinal<=len(panels):raise RuntimeError('Unknown native panel')
            expected=(841.89,595.276) if panels[ordinal-1].orientation=='landscape' else (595.276,841.89)
            if any(abs(a-b)>1 for a,b in zip((page.rect.width,page.rect.height),expected)):
                raise RuntimeError('Native saved-PDF dimensions are not the planned A4 orientation')
            if re.search(r'#{4,}',text):raise RuntimeError('Native PDF contains unreadable numeric overflow')
            order.append((ordinal,index))
        if {n for n,i in order}!=set(range(1,len(panels)+1)):
            raise RuntimeError('Native export omitted planned panels')
        pdf.select([i for _,i in sorted(order)])
        pdf.save(target)
        inspection=[]
        for name in ('STATS1','SYMBOLS','Trade Log','P&L Calendar'):
            candidates=[i for i,(ordinal,_) in enumerate(sorted(order)) if panels[ordinal-1].sheet==name]
            inspection.append(candidates[-1] if name=='STATS1' else candidates[0])
        for i in inspection:
            pdf[i].get_pixmap(matrix=fitz.Matrix(1.25,1.25),colorspace=fitz.csGRAY).save(Path(directory)/f'inspect-{i+1}.png')
        return {'pdf':str(target),'pages':len(pdf),'panel_order':[n for n,_ in sorted(order)],
                'inspection_pages':[i+1 for i in inspection]}


def main() -> int:
    import argparse
    parser=argparse.ArgumentParser(description='Legacy HTML or isolated native Excel print-copy proof')
    parser.add_argument('--native-preview-directory',type=Path)
    parser.add_argument('--virtual-context')
    args=parser.parse_args()
    if args.native_preview_directory is None:return legacy_main()
    directory=args.native_preview_directory.resolve()
    manifest=json.loads((directory/'plan.json').read_text(encoding='utf-8'))
    panels=[PrintPanel(**p) for p in manifest['panels']]
    native=native_preview(directory/'print-copy.xlsx',directory,panels,args.virtual_context)
    result=finalize_native_pdf(native,directory,panels)
    (directory/'result.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps(result),flush=True)
    return 0


if __name__=='__main__':
    raise SystemExit(main())
