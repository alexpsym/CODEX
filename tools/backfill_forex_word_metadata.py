"""7C1: saved Word provenance -> exact, read-only journal patch proposal.

No application mode exists. Numbers in JSON use decimal strings with an explicit
type/unit; they are never converted through binary floats for matching/planning.
Only the owned workbook snapshot is opened. No workbook is saved by this module.
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
from decimal import Decimal, InvalidOperation
import hashlib
import importlib.util
import io
import json
import os
import ntpath
from pathlib import Path, PurePosixPath
import re
import sys
import time
from types import SimpleNamespace
from urllib.parse import unquote, urlsplit
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

RECOVERY_HASH = "16c65e076563f8772263fecb221e339878bcb23709546a42023b60b351e16647"
RECOVERY_TEST_HASH = "3f66f952597169fa1afa612a5828e627ad59b20ef549ce79f733dbeced24167f"
TEMP = Path("C:/Users/User/AppData/Local/Temp")
RECEIPT_NAMES = (
    "forex_word_7b1_resume_4eb128350c0442ae8ba8e305c478497e",
    "forex_word_7b2_equivalence_dc2b5da560514669b8d8ae78a849321c",
    "forex_word_7b3_529338dadec748699802b92e0627af27",
    "forex_word_7b4_6b115ab1175c40c7ad1a6922e5a4e7e5",
    "forex_word_7b5_resume_ec2925baca0d4414aeed854c81c7b0e0",
    "forex_word_7b6_b01b253eb1bf45fa94485af0c198d146",
)
INVENTORY_DIR = TEMP / "forex_word_7a_host_46efeff728734c77be80a4f196803cff"
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
CCY = set("USD EUR GBP JPY AUD NZD CAD CHF NOK SEK DKK SGD HKD ZAR MXN TRY PLN CNH CNY".split())
IDENTITY_LABELS = {"PAIR", "SYMBOL", "SHORT/LONG", "SIDE", "DATE", "TIME", "ENTRY", "ENTRY PRICE", "TRADE NUMBER"}
YES_NO = {"YES": "Yes", "NO": "No", "TRUE": "Yes", "FALSE": "No", "0": "No", "1": "Yes"}
BOOLEAN_HEADERS = {"Test", "EMA", "VWAP", "Round Number", "Spiked Out", "Close Stopout", "Near Perfect Entry", "Near Win", "Early Close"}
# Schema-backed manual destinations only. Settled execution/accounting/identity
# columns are intentionally absent. SIGNAL CANDLE is not Pattern (range/channel).
RULES = {
    "SETUP": ("Setup", "text", "exact descriptive text; Setup is not Pattern"),
    "TIMEFRAME": ("Timeframe", "timeframe", "single explicit timeframe token, repository spelling"),
    "CHART": ("Timeframe", "timeframe", "historical chart timeframe, single explicit unit"),
    "MARKET/LIMIT/PENDING ORDER": ("Order", "order", "only Market/Limit proven by journal validation domain"),
    "ORDER": ("Order", "order", "only explicit Market or Limit"),
    "PATTERN": ("Pattern", "pattern", "only literal range/channel, not candlestick descriptions"),
    "ATHS/ATLS": ("ATHS/ATLS", "extreme", "explicit all-time high/low label"),
}
for _header in BOOLEAN_HEADERS:
    RULES[_header.upper()] = (_header, "yesno", "explicit Yes/No/True/False/0/1; period is not a boolean")
for _group in ("Break Even", "Profit"):
    for _measure in ("Trigger Price", "Distance From Entry %", "Distance From Exit %"):
        _dest = f"Move to {_group} {_measure}"
        RULES[_dest.upper()] = (_dest, "fraction" if _measure.endswith("%") else "decimal", "explicit labeled percent /100 to fractional journal percent" if _measure.endswith("%") else "explicit manual trigger price, no inference")


def require(ok, message):
    if not ok:
        raise ValueError(message)


def digest(raw):
    return hashlib.sha256(raw).hexdigest()


def encoded(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2).encode("utf-8")


def save(path, value):
    # Exclusive output files: an audit cannot silently overwrite another run.
    with Path(path).open("xb") as handle:
        handle.write(encoded(value))


def norm(text):
    return " ".join(str(text).strip().upper().split())


def identity(relative):
    relative = str(relative).replace("\\", "/")
    p = PurePosixPath(relative)
    require(not p.is_absolute() and not any(x in ("..", ".", "") or ":" in x for x in p.parts), "unsafe source identity")
    return str(p).casefold()


class JsonStore:
    """Read, hash and parse each approved saved input once; reject links/escapes."""

    def __init__(self, roots):
        self.roots = [Path(x).absolute() for x in roots]
        self.raws, self.parsed, self.hashes = {}, {}, {}
        self.handles, self.identities = {}, {}
        self.counters = collections.defaultdict(lambda: dict(reads=0, hashes=0, parses=0))

    def checked(self, path, directory=None):
        p = Path(path).absolute()
        candidates = [Path(directory).absolute()] if directory else self.roots
        root = next((r for r in candidates if p == r or r in p.parents), None)
        require(root is not None and any(root == r or r in root.parents for r in self.roots), "evidence path escape")
        require(p.resolve() == p and root.resolve() == root, "evidence link/alias")
        for q in (p, *p.parents):
            if q == root.parent:
                break
            st = q.lstat()
            require(not q.is_symlink() and not getattr(st, "st_file_attributes", 0) & 0x400, "linked evidence")
        return p

    def raw(self, path, directory=None):
        p = self.checked(path, directory)
        if p not in self.raws:
            require(p.suffix.lower() in (".json", ".jsonl"), "only saved JSON evidence allowed")
            handle = stable_read_handle(p)
            self.handles[p] = handle
            self.identities[p] = os.fstat(handle.fileno())
            b = handle.read()
            self.raws[p], self.hashes[str(p)] = b, digest(b)
            self.counters[str(p)]["reads"] += 1
            self.counters[str(p)]["hashes"] += 1
        return self.raws[p]

    def load(self, path, expected=None, directory=None):
        p = self.checked(path, directory)
        raw = self.raw(p, directory)
        require(expected is None or self.hashes[str(p)] == expected, "saved JSON hash mismatch")
        if p not in self.parsed:
            self.parsed[p] = json.loads(raw.decode("utf-8"))
            self.counters[str(p)]["parses"] += 1
        return self.parsed[p]

    def jsonl(self, path):
        p = self.checked(path)
        raw = self.raw(p)
        if p not in self.parsed:
            self.parsed[p] = [json.loads(x) for x in raw.decode("utf-8").splitlines()]
            self.counters[str(p)]["parses"] += 1
        return self.parsed[p]

    def close(self):
        for p, handle in self.handles.items():
            before, after = self.identities[p], os.fstat(handle.fileno())
            require((before.st_size, before.st_mtime_ns) == (after.st_size, after.st_mtime_ns), "saved input changed while held")
            handle.close()
        self.handles.clear()


def stable_read_handle(path):
    """Saved inputs only: hold read access denying concurrent write/delete."""
    if os.name != "nt":
        return Path(path).open("rb")
    import ctypes
    import msvcrt
    k = ctypes.WinDLL("kernel32", use_last_error=True)
    k.CreateFileW.argtypes = [ctypes.c_wchar_p, ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p, ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p]
    k.CreateFileW.restype = ctypes.c_void_p
    h = k.CreateFileW(str(path), 0x80000000, 1, None, 3, 0x00100000 | 0x08000000, None)
    if h == ctypes.c_void_p(-1).value:
        raise OSError(ctypes.get_last_error(), "saved-input read lock unavailable")
    return os.fdopen(msvcrt.open_osfhandle(h, os.O_RDONLY | os.O_BINARY), "rb")


def normalize_folder_link(raw, original_workbook, hyperlink_base=None, forex_root=r"C:\Users\User\Dropbox\FOREX"):
    """Pure Windows lexical paths, per Excel's workbook/HyperlinkBase rule.

    No filesystem access. Full addresses override HyperlinkBase. Supported bases
    are absolute local directory paths/file URIs; other forms stay unresolved.
    """
    evidence = dict(raw=raw, original_workbook=original_workbook, hyperlink_base=hyperlink_base, transformations=[])

    def decode(value):
        require(isinstance(value, str) and value and value == value.strip(), "empty/whitespace link")
        require(not re.search(r"%(?![0-9a-fA-F]{2})", value), "malformed encoding")
        decoded = unquote(value, encoding="utf-8", errors="strict")
        require("%" not in decoded, "multiply encoded/unsupported percent path")
        require(not any(ord(c) < 32 for c in decoded) and not any(c in decoded for c in '?#*"<>|'), "query/fragment/unsafe link syntax")
        if decoded.lower().startswith("file:"):
            uri = urlsplit(decoded)
            require(uri.scheme.lower() == "file" and not uri.netloc and re.match(r"^/[A-Za-z]:/", uri.path), "unsupported file authority/path")
            decoded = uri.path[1:]
        require(not re.match(r"^[A-Za-z][A-Za-z0-9+.-]*:", decoded) or re.match(r"^[A-Za-z]:[/\\]", decoded), "unsupported scheme/drive-relative path")
        decoded = decoded.replace("/", "\\")
        require(not decoded.startswith("\\") and not decoded.startswith("\\\\"), "root-relative/UNC path")
        require(":" not in decoded[2:], "alternate stream/colon")
        require(all(part in (".", "..") or not part.endswith((" ", ".")) for part in decoded.split("\\") if part), "ambiguous Windows component")
        return decoded

    try:
        target = decode(raw)
        origin = decode(original_workbook)
        require(re.match(r"^[A-Za-z]:\\", origin), "unproven original workbook location")
        absolute = bool(re.match(r"^[A-Za-z]:\\", target))
        if absolute:
            base = None
        elif hyperlink_base not in (None, ""):
            base = decode(hyperlink_base)
            require(re.match(r"^[A-Za-z]:\\", base), "unsupported relative HyperlinkBase")
        else:
            base = ntpath.dirname(origin)
        result = ntpath.normpath(target if absolute else ntpath.join(base, target))
        root = ntpath.normcase(ntpath.normpath(forex_root))
        require(ntpath.normcase(result).startswith(root + "\\"), "outside approved FOREX root")
        evidence.update(status="normalized", base=base, normalized_target=result, transformations=["one URL decode", "Windows separators", "absolute override" if absolute else "HyperlinkBase" if hyperlink_base else "original workbook directory", "lexical dot-segment normalization", "Windows case-insensitive full-path comparison"])
    except (ValueError, UnicodeError) as exc:
        evidence.update(status="unsupported", reason=str(exc), normalized_target=None)
    return evidence


def pointer(value, path):
    require(path.startswith("/"), "invalid JSON pointer")
    for key in path[1:].split("/"):
        key = key.replace("~1", "/").replace("~0", "~")
        value = value[int(key)] if isinstance(value, list) else value[key]
    return value


def _load_sources(handoff_path, *, approved=None, inventory_directory=None, expected_count=102, store):
    approved = [Path(x) for x in (approved or [TEMP / x for x in RECEIPT_NAMES])]
    inventory_directory = Path(inventory_directory or INVENTORY_DIR)
    final = approved[-1]
    require(Path(handoff_path).absolute() == final / "final_recovery_handoff.json", "unexpected handoff path")
    handoff = store.load(handoff_path, directory=final)
    coverage_ref = handoff["coverage_audit"]
    require(Path(coverage_ref["path"]) == final / "final_recovery_coverage.json", "unexpected coverage reference")
    coverage = store.load(coverage_ref["path"], coverage_ref["sha256"], final)
    context = store.load(final / "final_audit_context.json", directory=final)
    progress = store.load(final / "cumulative_progress.json", directory=final)
    refs = context["prior_completions"]
    require([Path(x["directory"]) for x in refs] == approved, "unexpected receipt directories/order")
    require(refs == coverage["completed_receipts"] == progress["completed_receipts"], "conflicting receipt references")
    require(coverage["status"] == "complete" and coverage["exact_set_equality"] and progress["remaining"] == 0, "incomplete coverage")
    inv = Path(context["inventory"])
    require(inv == inventory_directory / "manifest.jsonl", "unexpected inventory path")
    raw_inv = store.raw(inv)
    require(store.hashes[str(inv)] == context["inventory_sha256"] == handoff["inventory"]["sha256"], "inventory hash mismatch")
    records = store.jsonl(inv)
    inventory = {identity(x["relative_path"]): x for x in records if x.get("kind") == "document"}
    require(len(inventory) == expected_count, "inventory source count mismatch")
    recovery_path = Path(__file__).with_name("recover_forex_word_documents.py")
    require(digest(recovery_path.read_bytes()) == RECOVERY_HASH, "recovery tool changed")
    require(digest(recovery_path.parents[1].joinpath("tests/test_recover_forex_word_documents.py").read_bytes()) == RECOVERY_TEST_HASH, "recovery test changed")
    # Cache-backed Path read for the unchanged receipt-only helper. It is never
    # allowed to open captures or published output files. No recovery CLI/parser.
    spec = importlib.util.spec_from_file_location("word_receipt_validation", recovery_path)
    recovery = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(recovery)
    base_path = type(Path())

    class CachedPath(base_path):
        def read_bytes(self):
            return store.raw(Path(str(self)))

    recovery.Path = CachedPath
    hash_cache = {}

    def cached_digest(raw):
        if raw not in hash_cache:
            # reuse JsonStore digests where already computed
            hash_cache[raw] = next((store.hashes[str(p)] for p, b in store.raws.items() if b is raw), None) or digest(raw)
        return hash_cache[raw]

    recovery.digest = cached_digest
    for ref in refs:
        for name, sha in ref["records"].items():
            require(Path(name).name == name, "receipt name escape")
            store.load(Path(ref["directory"]) / name, sha, ref["directory"])
    parse_cache = {store.raws[p].decode("utf-8"): value for p, value in store.parsed.items() if p.suffix == ".json"}

    def cached_loads(text):
        if text not in parse_cache:
            parse_cache[text] = json.loads(text)
        return parse_cache[text]

    recovery.json = SimpleNamespace(loads=cached_loads)
    completed = recovery.completion_sources(context, records)
    require(completed == set(inventory) and len(completed) == expected_count, "receipt identity-set mismatch")
    require(handoff["document_count"] == len(handoff["documents"]) == expected_count, "handoff count mismatch")
    sources, seen = [], set()
    refs_by_dir = {str(Path(x["directory"])): x for x in refs}
    for entry in handoff["documents"]:
        key = identity(entry["source"]["relative_path"])
        require(key == entry["source_identity"] and key in completed and key not in seen, "duplicate/foreign source")
        seen.add(key)
        directory = Path(entry["receipt_directory"])
        require(str(directory) in refs_by_dir, "unexpected source receipt")
        ref = refs_by_dir[str(directory)]
        require(entry["receipt_hashes"] == ref["records"] and entry["historical_tool_sha256"] == ref["tool_sha256"], "source receipt linkage mismatch")
        pre = entry["preparation_reference"]
        require(Path(pre["path"]) == directory / "preparation.json" and pre["sha256"] == ref["records"]["preparation.json"], "preparation reference mismatch")
        prep = store.load(pre["path"], pre["sha256"], directory)
        index = pre["document_index"]
        require(pre["json_pointer"] == f"/documents/{index}", "document pointer mismatch")
        doc = pointer(prep, pre["json_pointer"])
        require(doc["source"] == entry["source"] and doc["source_sha256"] == entry["source_sha256"], "source hash/metadata linkage mismatch")
        require(all(doc["source"].get(k) == v for k, v in inventory[key].items()), "manifest metadata mismatch")
        require(doc["complete"] and not doc["coverage"]["issues"], "incomplete source")
        provenance_path = Path(entry["provenance_reference"]["path"])
        require(provenance_path == directory / f"document_{index+1:02d}" / "provenance.json", "provenance path mismatch")
        require(store.load(provenance_path, directory=directory) == doc, "provenance/preparation differs")
        require(entry["coverage_reference"]["path"] == pre["path"] and entry["coverage_reference"]["json_pointer"] == f"/documents/{index}/coverage", "coverage pointer mismatch")
        tables = doc["coverage"]["tables"]
        require(entry["structured_tables"] == [dict(table_id=t["id"], json_pointer=f"/documents/{index}/coverage/tables/{i}") for i, t in enumerate(tables)], "table pointer mismatch")
        sources.append(dict(identity=key, document=doc, preparation=pre["path"], pointer=pre["json_pointer"], provenance=str(provenance_path)))
    require(seen == completed, "missing handoff source")
    return sources, store, handoff


def load_sources(handoff_path, *, approved=None, inventory_directory=None, expected_count=102):
    roots = [Path(x) for x in (approved or [TEMP / x for x in RECEIPT_NAMES])]
    inv = Path(inventory_directory or INVENTORY_DIR)
    store = JsonStore([*roots, inv])
    try:
        return _load_sources(handoff_path, approved=roots, inventory_directory=inv, expected_count=expected_count, store=store)
    except BaseException:
        store.close()
        raise


class Snapshot:
    def __init__(self, path, capture=None):
        self.path = Path(path).absolute()
        require(self.path != Path(__file__).resolve().parents[1] / "journal" / "Trading Journal.xlsx", "live workbook prohibited")
        self.handle = stable_read_handle(self.path)
        self.raw = self.handle.read()
        self.sha256 = digest(self.raw)
        self.capture = capture
        self.original_workbook = capture["source"] if capture else None
        if capture:
            require(capture["sha256"] == self.sha256 and capture["before"] == capture["after"], "invalid origin capture receipt")
            require(capture.get("successful_copies") == 1 and capture.get("read_only") and capture.get("share_read_only"), "unproven snapshot origin")
        self.wb = load_workbook(io.BytesIO(self.raw), data_only=False, keep_links=False)
        self.ws = self.wb["Trade Log"]
        from tools import master_journal_workbook as journal
        self.headers = journal._trade_log_header_map(self.ws)
        self.start = journal._trade_log_data_start_row(self.ws)
        require(all(x in self.headers for x in ("Trade Number", "Row ID", "Row Type", "Symbol", "Open Time")), "missing identity headers")
        require(len(set(self.headers.values())) == len(self.headers), "duplicate headers")
        with zipfile.ZipFile(io.BytesIO(self.raw)) as z:
            self.part = journal._worksheet_ooxml_part_name(z, self.ws.title)
            self.xml = ET.fromstring(z.read(self.part))
            properties = ET.fromstring(z.read("docProps/app.xml")) if "docProps/app.xml" in z.namelist() else None
            base = properties.find("{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}HyperlinkBase") if properties is not None else None
            self.hyperlink_base = base.text if base is not None else None
        self.cells = {x.attrib["r"]: x for x in self.xml.iter(NS + "c")}
        self.formula_ranges = []
        for cell in self.cells.values():
            formula = cell.find(NS + "f")
            if formula is not None and formula.get("ref"):
                self.formula_ranges.append(range_boundaries(formula.get("ref")))
        self.formula_count = sum(c.find(NS + "f") is not None for c in self.cells.values())
        self.rowids = collections.Counter()
        self.rows = []
        for row in range(self.start, self.ws.max_row + 1):
            data = {h: self.ws.cell(row, col).value for h, col in self.headers.items()}
            if data["Row ID"] is not None:
                self.rowids[str(data["Row ID"])] += 1
            number, formula_unknown, link = trade_number(data["Trade Number"])
            if number or formula_unknown:
                self.rows.append(dict(row=row, values=data, f_number=number, formula_unknown=formula_unknown, folder=link))

    def literal(self, row, header):
        cell = self.ws.cell(row, self.headers[header]) if header in self.headers else None
        return cell is not None and cell.data_type != "f" and not (cell.value is not None and not isinstance(cell.value, (str, int, float, bool, dt.date, dt.datetime)))

    def target(self, row, header):
        col = self.headers[header]
        cell = self.ws.cell(row, col)
        coordinate = cell.coordinate
        node = self.cells.get(coordinate)
        xml = ET.tostring(node, encoding="unicode") if node is not None else None
        reason = None
        if any(lo <= col <= hi and top <= row <= bottom for lo, top, hi, bottom in self.formula_ranges):
            reason = "formula_range"
        for merged in self.ws.merged_cells.ranges:
            if coordinate in merged and coordinate != merged.start_cell.coordinate:
                reason = "merged_continuation"
        if cell.data_type == "f" or node is not None and node.find(NS + "f") is not None:
            reason = "formula"
        if reason is None and node is not None:
            if list(node) or node.text not in (None, "") or node.get("t") not in (None, "n") or any(k not in ("r", "s", "t") for k in node.attrib):
                reason = "nonblank"
        if reason is None and (cell.value is not None or cell.data_type not in ("n",)):
            reason = "nonblank"
        signature = dict(coordinate=coordinate, xml=xml, xml_sha256=digest(xml.encode()) if xml else None, data_type=cell.data_type, style_id=cell.style_id, number_format=cell.number_format, hyperlink=str(cell.hyperlink.target) if cell.hyperlink else None, comment=cell.comment.text if cell.comment else None, structurally_blank=reason is None, reason=reason)
        return signature


def trade_number(value):
    if not isinstance(value, str):
        return None, False, None
    text = value.strip()
    link = None
    if text.startswith("="):
        match = re.fullmatch(r'=HYPERLINK\("((?:[^"\r\n]|"")*)"\s*[,;]\s*"(F\d+)"\)', text, re.I)
        if not match:
            return None, True, None
        link, text = match.group(1).replace('""', '"'), match.group(2)
    if re.fullmatch(r"F[0-9]+", text, re.I):
        return "F" + str(int(text[1:])), False, link
    return None, False, link


def symbol(value):
    text = re.sub(r"[ /]", "", norm(value))
    return text if len(text) == 6 and text[:3] in CCY and text[3:] in CCY and text[:3] != text[3:] else None


def side(value):
    return {"BUY": "Long", "LONG": "Long", "SELL": "Short", "SHORT": "Short"}.get(norm(value))


def fields(source):
    """Use origin cells only, never fallback aliases or merge continuations."""
    doc = source["document"]
    result = []
    for ti, table in enumerate(doc["coverage"]["tables"]):
        cells = {(c["row"], c["column"]): (ci, c) for ci, c in enumerate(table["cells"])}
        require(len(cells) == len(table["cells"]), "duplicate table origin")
        for (row, col), (ci, cell) in cells.items():
            if col % 2 != 0:
                continue
            other = cells.get((row, col + cell["colspan"]))
            ptr = source["pointer"] + f"/coverage/tables/{ti}/cells/{ci}"
            base = dict(label=cell["text"], table_id=table["id"], part=table["part"], label_pointer=ptr, label_coordinates={k: cell[k] for k in ("row", "column", "rowspan", "colspan")})
            if other:
                oi, value = other
                result.append(dict(**base, raw=value["text"], pointer=source["pointer"] + f"/coverage/tables/{ti}/cells/{oi}", coordinates={k: value[k] for k in ("row", "column", "rowspan", "colspan")}))
            elif str(cell["text"]).strip():
                result.append(dict(**base, raw=cell["text"], pointer=ptr, coordinates=base["label_coordinates"], unpaired=True))
    for pi, prose in enumerate(doc["coverage"]["prose"]):
        if prose.get("text", "").strip():
            result.append(dict(label="[unstructured prose]", raw=prose["text"], pointer=source["pointer"] + f"/coverage/prose/{pi}/text", context=prose.get("context"), unpaired=True))
    return result


def observed_date(raw):
    text = str(raw).strip()
    formats = ("%d/%m/%y", "%d/%m/%Y", "%d-%m-%Y")
    if re.fullmatch(r"\d{6}", text):
        formats = ("%d%m%y",)
    for fmt in formats:
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def match_source(source, observations, snapshot):
    d = source["document"]
    src = d["source"]
    evidence = []
    reasons = []
    fset = set()
    for text in [src["relative_path"], *src.get("f_number_evidence", [])]:
        fset.update("F" + str(int(n)) for n in re.findall(r"(?i)(?<!\w)F(\d+)(?!\w)", text.replace("\\", "/")))
    content = "\n".join(c["text"] for t in d["coverage"]["tables"] for c in t["cells"]) + "\n" + "\n".join(p.get("text", "") for p in d["coverage"]["prose"])
    content_f = {"F" + str(int(n)) for n in re.findall(r"(?i)(?<!\w)F(\d+)(?!\w)", content)}
    if len(fset) != 1 or content_f - fset:
        reasons.append("missing_or_conflicting_source_f_number")
    number = next(iter(fset)) if len(fset) == 1 else None
    rows = [r for r in snapshot.rows if r["f_number"] == number] if number else []
    if len(rows) != 1:
        reasons.append("missing_or_duplicate_journal_f_number")
    selected = rows[0] if len(rows) == 1 else None
    if selected:
        values, row = selected["values"], selected["row"]
        rid = values.get("Row ID")
        if not isinstance(rid, str) or not rid.strip() or snapshot.rowids[rid] != 1 or not snapshot.literal(row, "Row ID"):
            reasons.append("missing_duplicate_or_formula_row_id")
        for h in ("Row Type", "Symbol", "Open Time", "Side", "Entry Price"):
            if h in snapshot.headers and not snapshot.literal(row, h):
                reasons.append("formula_identity:" + h)
        if norm(values.get("Row Type")) != "TRADE" or symbol(values.get("Symbol")) is None:
            reasons.append("not_qualifying_forex_trade")
        # All explicit reliable identity fields must agree. Unknown multi-value
        # identity values remain unresolved instead of picking one occurrence.
        pairs = [x for x in observations if norm(x["label"]) in ("PAIR", "SYMBOL") and x["raw"] is not None and str(x["raw"]).strip()]
        pair_values = {symbol(x["raw"]) for x in pairs} | {symbol(s) for s in src.get("symbol_evidence", [])}
        if pair_values:
            agrees = pair_values == {symbol(values.get("Symbol"))} and None not in pair_values
            evidence.append(dict(kind="symbol", raw=[x["raw"] for x in pairs], canonical=sorted(str(x) for x in pair_values), journal=values.get("Symbol"), agrees=agrees, normalization="uppercase explicit currency pairs, remove slash/space only"))
            if not agrees:
                reasons.append("symbol_conflict_or_ambiguous")
        opening = values.get("Open Time")
        if isinstance(opening, str):
            try:
                opening = dt.datetime.fromisoformat(opening)
            except ValueError:
                opening = None
        opening = opening.date() if isinstance(opening, dt.datetime) else opening
        year = str(src.get("year", ""))
        if year:
            agrees = isinstance(opening, dt.date) and str(opening.year) == year
            evidence.append(dict(kind="year", source=year, journal=str(opening), agrees=agrees))
            if not agrees:
                reasons.append("year_conflict_or_unknown")
        for item in observations:
            label, raw = norm(item["label"]), item["raw"]
            if raw is None or not str(raw).strip() or norm(raw) in {"N/A", "-", "UNKNOWN"}:
                continue
            journal_value, canonical = None, None
            if label in ("SHORT/LONG", "SIDE"):
                canonical, journal_value = side(raw), side(values.get("Side"))
            elif label in ("ENTRY PRICE", "ENTRY"):
                try:
                    canonical, journal_value = Decimal(str(raw).strip()), Decimal(str(values.get("Entry Price")))
                except InvalidOperation:
                    canonical = None
            elif label == "DATE":
                canonical, journal_value = observed_date(raw), opening
            else:
                continue
            agrees = canonical is not None and canonical == journal_value
            evidence.append(dict(kind=label, raw=raw, canonical=str(canonical), journal=str(journal_value), agrees=agrees))
            if not agrees:
                reasons.append("identity_conflict_or_ambiguous:" + label)
        folder_agrees = False
        # Optional corroboration from a literal stored link, never link evaluation.
        target = selected["folder"]
        if not target:
            cell = snapshot.ws.cell(row, snapshot.headers["Trade Number"])
            target = cell.hyperlink.target if cell.hyperlink else None
        if target:
            expected = ntpath.normpath(r"C:\Users\User\Dropbox\FOREX" + "\\" + src["parent_relative_path"].replace("/", "\\"))
            normalized = normalize_folder_link(target, snapshot.original_workbook, snapshot.hyperlink_base)
            folder_agrees = normalized["status"] == "normalized" and ntpath.normcase(normalized["normalized_target"]) == ntpath.normcase(expected)
            origin_cell = snapshot.ws.cell(row, snapshot.headers["Trade Number"])
            evidence.append(dict(kind="folder_hyperlink", raw=target, agrees=folder_agrees, normalization=normalized, expected_parent=expected, cell=origin_cell.coordinate, relationship_id=origin_cell.hyperlink.id if origin_cell.hyperlink else None, origin="literal HYPERLINK" if selected["folder"] else "worksheet relationship"))
            if not folder_agrees:
                reasons.append("folder_hyperlink_conflict_or_unknown")
        symbol_year = bool(pair_values) and pair_values == {symbol(values.get("Symbol"))} and isinstance(opening, dt.date) and year == str(opening.year)
        if not (symbol_year or folder_agrees):
            reasons.append("insufficient_non_f_corroboration")
    return dict(source_identity=source["identity"], source_sha256=d["source_sha256"], f_number=number, normalization="F/f + decimal digits -> F + integer (leading zeros explicit)", source_content_f_numbers=sorted(content_f), status="proven" if not reasons else "unresolved", reasons=sorted(set(reasons)), comparisons=evidence, candidate_rows=[r["row"] for r in rows], row=selected["row"] if selected else None, row_id=selected["values"].get("Row ID") if selected else None)


def mapping(item, headers, number_format=None):
    label, raw = norm(item["label"]), item["raw"]
    result = dict(label=item["label"], context={k: item.get(k) for k in ("table_id", "part", "context", "label_coordinates")}, raw=raw)
    if item.get("unpaired"):
        return dict(**result, status="ambiguous", reason="unstructured prose/merged or unpaired content; no inferred label, time or units")
    if raw is None or str(raw).strip() == "" or norm(raw) in {"N/A", "UNKNOWN", "-"}:
        return dict(**result, status="absent", reason="no explicit source value")
    if label not in RULES:
        reason = "identity/settled execution/accounting fact; comparison only" if label in IDENTITY_LABELS else "no safe existing manual destination; no Notes fallback"
        if label in {"MOVING AVERAGE", "MOVING AVERAGES IN SEQUENCE"}:
            reason = "EMA period/MA sequence is not EMA Yes/No"
        if label == "SIGNAL CANDLE":
            reason = "candlestick description is not journal Pattern range/channel"
        if label in {"RISK", "MAX LOSS", "RISK : REWARD", "RISK:REWARD"}:
            reason = "risk mode/value or settled/calculated fact excluded"
        return dict(**result, status="no_safe_destination", reason=reason)
    header, kind, rationale = RULES[label]
    result.update(header=header, type=kind, rationale=rationale)
    if header not in headers:
        return dict(**result, status="no_safe_destination", reason="canonical destination absent from snapshot")
    text = str(raw).strip()
    value, unit = raw, "text"
    if kind == "yesno":
        value = YES_NO.get(norm(raw))
    elif kind == "timeframe":
        token = re.sub(r"\s", "", norm(raw))
        value = {"1MIN": "1MIN", "5MIN": "5MIN", "15MIN": "15MIN", "30MIN": "30MIN", "1H": "1H", "4H": "4H", "DAILY": "1D", "WEEKLY": "1W", "MONTHLY": "1MO"}.get(token)
    elif kind == "order":
        value = {"MARKET": "Market", "MARKET ORDER": "Market", "LIMIT": "Limit", "BUY LIMIT": "Limit", "SELL LIMIT": "Limit", "LIMIT ORDER": "Limit"}.get(norm(raw))
    elif kind == "pattern":
        value = text.lower() if text.lower() in {"range", "channel"} else None
    elif kind == "extreme":
        value = {"ATH": "All-Time High", "ATHS": "All-Time High", "ALL-TIME HIGH": "All-Time High", "ATL": "All-Time Low", "ATLS": "All-Time Low", "ALL-TIME LOW": "All-Time Low"}.get(norm(raw))
    elif kind in {"decimal", "fraction"}:
        pattern = r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)"
        if kind == "fraction":
            if not text.endswith("%") or number_format is not None and "%" not in number_format:
                value = None
            else:
                text = text[:-1].strip()
                value = str(Decimal(text) / Decimal(100)) if re.fullmatch(pattern, text) else None
            unit = "fraction_of_one"
        else:
            value = str(Decimal(text)) if re.fullmatch(pattern, text) else None
            unit = "price"
    if value is None:
        return dict(**result, status="ambiguous", reason="unknown value/unit, incompatible domain or unsupported conversion")
    return dict(**result, status="supported", canonical=dict(type="decimal" if kind in {"decimal", "fraction"} else "string", value=value, unit=unit))


def plan(sources, snapshot):
    audits, catalog, patch, decisions = [], [], [], []
    seen_sources = set()
    for source in sources:
        require(source["identity"] not in seen_sources, "duplicate source")
        seen_sources.add(source["identity"])
        obs = fields(source)
        match = match_source(source, obs, snapshot)
        audits.append(match)
        groups = collections.defaultdict(list)
        blocked_headers = set()
        for item in obs:
            tentative = mapping(item, snapshot.headers)
            if tentative.get("header") and match["status"] == "proven":
                fmt = snapshot.ws.cell(match["row"], snapshot.headers[tentative["header"]]).number_format
                tentative = mapping(item, snapshot.headers, fmt)
            record = dict(source_identity=source["identity"], source_pointer=item["pointer"], **tentative)
            catalog.append(record)
            if tentative["status"] == "supported":
                groups[tentative["header"]].append((item, tentative))
            else:
                if tentative.get("header") and tentative["status"] != "absent":
                    blocked_headers.add(tentative["header"])
                decisions.append(dict(source_identity=source["identity"], pointer=item["pointer"], status=tentative["status"], reason=tentative.get("reason"), raw=item["raw"], label=item["label"]))
        for header, items in groups.items():
            values = {json.dumps(m["canonical"], sort_keys=True) for _, m in items}
            decision = dict(source_identity=source["identity"], header=header, source_pointers=[i["pointer"] for i, _ in items])
            if header in blocked_headers:
                decisions.append(dict(**decision, status="ambiguous_source_field"))
                continue
            if len(values) != 1:
                decisions.append(dict(**decision, status="conflicting_source_field"))
                continue
            if match["status"] != "proven":
                decisions.append(dict(**decision, status="unresolved_identity", reasons=match["reasons"]))
                continue
            target = snapshot.target(match["row"], header)
            value = items[0][1]["canonical"]
            if not target["structurally_blank"]:
                cell = snapshot.ws.cell(match["row"], snapshot.headers[header])
                equal = str(cell.value) == str(value["value"]) and cell.data_type != "f"
                decisions.append(dict(**decision, status="existing_equal" if equal else "existing_conflicting", target=target, existing=str(cell.value), proposed=value))
                continue
            patch.append(dict(sheet=snapshot.ws.title, coordinate=target["coordinate"], header=header, row=match["row"], row_id=match["row_id"], f_number=match["f_number"], source_identity=source["identity"], source_sha256=source["document"]["source_sha256"], preparation_path=source["preparation"], sources=[dict(pointer=i["pointer"], raw=i["raw"], table_id=i.get("table_id"), coordinates=i.get("coordinates"), label=i["label"]) for i, _ in items], new_value=value, target_precondition=target))
    # Different documents may point at the same row. Agreeing fields are merged
    # with all attribution retained; conflicting proposals are removed entirely.
    unique = {}
    for entry in patch:
        key = (entry["sheet"], entry["coordinate"])
        unique.setdefault(key, []).append(entry)
    patch = []
    for key, entries in sorted(unique.items()):
        if len({json.dumps(x["new_value"], sort_keys=True) for x in entries}) != 1:
            decisions.append(dict(status="conflicting_documents_coordinate", coordinate=key[1], sources=[x["source_identity"] for x in entries]))
        else:
            first = entries[0]
            first["attributions"] = [dict(source_identity=x["source_identity"], source_sha256=x["source_sha256"], preparation_path=x["preparation_path"], sources=x["sources"]) for x in entries]
            patch.append(first)
    return audits, catalog, patch, decisions


def run_preview(handoff, snapshot_path, output, *, approved=None, inventory_directory=None, expected_count=102):
    start = time.monotonic()
    output = Path(output).absolute()
    require(output.is_dir() and not output.is_symlink(), "explicit private output directory required")
    require(Path(snapshot_path).absolute().parent == output and Path(snapshot_path).suffix == ".xlsx", "snapshot must be owned by output directory")
    sources, store, _ = load_sources(handoff, approved=approved, inventory_directory=inventory_directory, expected_count=expected_count)
    capture_path = output / "snapshot_receipt.json"
    capture_raw = capture_path.read_bytes() if capture_path.is_file() else None
    capture = json.loads(capture_raw) if capture_raw else None
    snapshot = Snapshot(snapshot_path, capture)
    audits, catalog, patch, decisions = plan(sources, snapshot)
    adapter_hash = digest(Path(__file__).read_bytes())
    mapping_hash = digest(encoded(catalog))
    verifier_path = Path(__file__).with_name("verify_forex_word_metadata_preview.py")
    verifier_hash = digest(verifier_path.read_bytes())
    receipt = dict(schema_version=1, mode="preview_only", saved_json_hashes=store.hashes, recovery_tool_sha256=RECOVERY_HASH, adapter_sha256=adapter_hash, verifier_sha256=verifier_hash, snapshot=dict(path=str(snapshot.path), sha256=snapshot.sha256, original_workbook=snapshot.original_workbook, hyperlink_base=snapshot.hyperlink_base), schema=dict(sheet=snapshot.ws.title, headers=snapshot.headers, data_start=snapshot.start, worksheet_part=snapshot.part), source_count=len(sources), cache_counters=dict(store.counters))
    if capture:
        require(capture["sha256"] == snapshot.sha256, "snapshot capture hash mismatch")
        receipt["snapshot"]["capture"] = capture
        receipt["snapshot"]["capture_receipt_sha256"] = digest(capture_raw)
    proposal = dict(schema_version=1, mode="preview_only", snapshot_sha256=snapshot.sha256, adapter_sha256=adapter_hash, verifier_sha256=verifier_hash, mapping_sha256=mapping_hash, coordinates_valid_only_for_recorded_row_ids=True, entries=patch, decisions=decisions)
    save(output / "input_receipt.json", receipt)
    save(output / "document_match_audit.json", dict(documents=audits, count=len(audits)))
    save(output / "field_mapping_catalog.json", catalog)
    save(output / "proposed_cell_patch.json", proposal)
    review = ["7C1 PREVIEW ONLY. No live writes. Coordinates bound to snapshot " + snapshot.sha256]
    review.extend("MATCH " + json.dumps(x, ensure_ascii=False, default=str) for x in audits)
    review.extend("PROPOSE " + json.dumps(x, ensure_ascii=False, default=str) for x in patch)
    review.extend("SKIP " + json.dumps(x, ensure_ascii=False, default=str) for x in decisions)
    with (output / "review.txt").open("x", encoding="utf-8") as handle:
        handle.write("\n".join(review) + "\n")
    require(os.fstat(snapshot.handle.fileno()).st_size == len(snapshot.raw), "snapshot modified")
    snapshot.handle.close()
    store.close()
    counts = dict(documents=len(audits), proven_matches=sum(a["status"] == "proven" for a in audits), unresolved_matches=sum(a["status"] != "proven" for a in audits), proposed_cells=len(patch), proposed_rows=len({x["row"] for x in patch}), formula_cells_preserved=snapshot.formula_count, existing_equal=sum(x["status"] == "existing_equal" for x in decisions), existing_conflicting=sum(x["status"] == "existing_conflicting" for x in decisions), mapping_classes=dict(collections.Counter(x["status"] for x in catalog)), decisions=dict(collections.Counter(x["status"] for x in decisions)))
    save(output / "preview_result.json", dict(status="complete", counts=counts, snapshot_sha256=snapshot.sha256, snapshot_unchanged=True, production_writes=0, elapsed_seconds=time.monotonic()-start))
    diagnostic = []
    for source, audit in zip(sources, audits):
        for comparison in audit["comparisons"]:
            if comparison["kind"] in ("symbol", "ENTRY", "ENTRY PRICE", "DATE", "SHORT/LONG") and not comparison["agrees"]:
                observed = fields(source)
                labels = {"symbol": {"PAIR", "SYMBOL"}, "SHORT/LONG": {"SHORT/LONG", "SIDE"}}.get(comparison["kind"], {comparison["kind"]})
                diagnostic.append(dict(source_identity=source["identity"], row=audit["row"], row_id=audit["row_id"], comparison=comparison, context=[x for x in observed if norm(x["label"]) in labels], blocker="unchanged exact identity comparison; no semantic reinterpretation/tolerance"))
    save(output / "conflict_context_diagnostic.json", diagnostic)
    return counts


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--handoff", required=True, type=Path)
    parser.add_argument("--workbook-snapshot", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    print(json.dumps(run_preview(args.handoff, args.workbook_snapshot, args.output)))


if __name__ == "__main__":
    sys.dont_write_bytecode = True
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    main()
