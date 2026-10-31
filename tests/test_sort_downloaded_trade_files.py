import ast
import io
import json
import os
import subprocess
import threading
import time
import types
from pathlib import Path
from urllib.parse import urlparse
from typing import Dict
import pytest
from tools import sort_downloaded_trade_files as m

ROOT=Path(__file__).resolve().parents[1]

def candidate(folder,**extra):
    return dict({'id':'row-398','trade_number':'C398','folder':str(folder),'symbol':'BTCUSDT','account':'BYBIT DEMO','asset_class':'crypto','side':'BUY','open_time':'2026-11-01T09:00:00+10:00','close_time':'2026-11-01T11:00:00+10:00'},**extra)

def aged(path,data=b'chart'):
    path.parent.mkdir(parents=True,exist_ok=True);path.write_bytes(data);past=time.time()-30;os.utime(path,(past,past));return path

def test_matching_evidence_and_ambiguity(tmp_path):
    row=candidate(tmp_path/'Dropbox/CRYPTO/2026/C398');rows=[row]
    ev=m.evidence('C398_BYBIT_DEMO_BTCUSDT.png');assert m.match_file(ev,rows)[0]==row
    assert m.match_file(m.evidence('F398.txt'),rows)[0] is None
    assert m.match_file(m.evidence('cell_C398.txt',b'The worksheet address is C398, not a trade label.'),rows)[0] is None
    assert m.match_file(m.evidence('trade-C398-and-trade-C399.txt'),rows)[0] is None
    assert m.match_file(ev,rows+[dict(row,id='repeat',account='BYBIT LIVE')])[0]==row
    assert m.match_file(m.evidence('C398.png'),rows+[dict(row,id='old',open_time='2025-01-01')])[0] is None
    assert m.match_file(m.evidence('C398_BTCUSDC.png'),rows)[0] is None
    assert m.match_file(m.evidence('C398_OANDA_LIVE_EURUSD.txt'),rows)[0] is None
    assert m.match_file(m.evidence('BINANCE_BTCUSDT.P,60_2026-11-01_10-00-00.png'),rows)[0] is None
    # Explicit offset, exact symbol/account and unique trade interval; UTC day rollover.
    note=b'Symbol: BYBIT:BTCUSDT.P\nAccount: BYBIT DEMO\nCapture time: 2026-10-31T23:30:00Z'
    assert m.match_file(m.evidence('chart.txt',note),rows)[0]==row
    assert m.match_file(m.evidence('chart.txt',note.replace(b'23:30',b'20:30')),rows)[0] is None
    assert m.match_file(m.evidence('chart.txt',note),rows+[dict(row,id='other')])[0] is None
    assert m.match_file(m.evidence('chart.txt',note.replace(b'Z',b'')),rows)[0] is None
    assert m.evidence('BYBIT_BTCUSDT_2026-11-01_10-00-00.png')['assumptions']
    plain=b'Trade Number: C398\nSymbol: BTCUSDT\nAccount: BYBIT DEMO\nEntry time: 2026-11-01T09:00:00+10:00'
    assert m.match_file(m.evidence('notes.txt',plain),rows)[0]==row
    assert m.match_file(m.evidence('C399.txt',plain),rows)[0] is None
    import zipfile
    xml=b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:tbl><w:tr><w:tc><w:p><w:r><w:t>Trade Number: C398</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:body></w:document>'
    buf=io.BytesIO()
    with zipfile.ZipFile(buf,'w') as z:z.writestr('word/document.xml',xml)
    assert m.match_file(m.evidence('notes.docx',buf.getvalue()),rows)[0]==row
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    w=Workbook();ws=w.active;ws.append(['Trade Number','Symbol','Account']);ws.append(['C398','BTCUSDT','BYBIT DEMO']);ws['D2'].comment=Comment('Trade ID: C398','test');ws['C398']='not a trade identity';ws['E2']='=1/0';buf=io.BytesIO();w.save(buf);w.close()
    assert m.match_file(m.evidence('notes.xlsx',buf.getvalue()),rows)[0]==row
    w=Workbook();ws=w.active;ws.append(['Trade Number','Symbol']);ws.append(['C398','BTCUSDT']);ws.append(['C399','BTCUSDT']);buf=io.BytesIO();w.save(buf);w.close()
    assert m.match_file(m.evidence('notes.xlsx',buf.getvalue()),rows)[0] is None
    # Small literal PDF, parsed by the installed safe reader; no external tool/API.
    payload=b'BT /F1 12 Tf 10 50 Td (Trade Number: C398) Tj ET'
    objects=[b'<< /Type /Catalog /Pages 2 0 R >>',b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 100] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',b'<< /Length '+str(len(payload)).encode()+b' >>\nstream\n'+payload+b'\nendstream']
    pdf=b'%PDF-1.4\n';offsets=[0]
    for i,obj in enumerate(objects,1):offsets.append(len(pdf));pdf+=str(i).encode()+b' 0 obj\n'+obj+b'\nendobj\n'
    start=len(pdf);pdf+=b'xref\n0 6\n0000000000 65535 f \n'+b''.join(f'{x:010} 00000 n \n'.encode() for x in offsets[1:])+b'trailer << /Size 6 /Root 1 0 R >>\nstartxref\n'+str(start).encode()+b'\n%%EOF'
    assert m.match_file(m.evidence('note.pdf',pdf),rows)[0]==row
    assert m.match_file(m.evidence('unknown.bin'),rows)[0] is None


def test_verified_moves_collisions_and_recovery(tmp_path,monkeypatch):
    source=tmp_path/'CHARTS';root=tmp_path/'Dropbox'/'CRYPTO';folder=root/'2026'/'C398';source.mkdir();folder.mkdir(parents=True);rows=[candidate(folder)];audit=tmp_path/'audit'
    one=aged(source/'nested'/'C398_chart.txt',b'one');unknown=aged(source/'unknown.txt',b'unknown');part=aged(source/'C398.part',b'incomplete')
    published=[];original=m.publish
    def record_publish(stage,target):assert stage.parent==target.parent and stage.parent!=source;published.append(target);return original(stage,target)
    monkeypatch.setattr(m,'publish',record_publish)
    report=m.run_batch('a'*32,source,[root],audit,rows)
    assert report['counts']['moved']==1 and report['counts']['unresolved']==2 and not one.exists() and unknown.exists() and part.exists() and (source/'nested').is_dir() and source.is_dir()
    assert published and (folder/'C398_chart.txt').read_bytes()==b'one'
    one=aged(source/'nested'/'C398_chart.txt',b'one');m.run_batch('a'*32,source,[root],audit,rows);assert one.exists()
    again=m.run_batch('b'*32,source,[root],audit,rows);assert again['counts']['already_filed']==1 and not one.exists()
    collision=aged(source/'C398_chart.TXT',b'different');r=m.run_batch('c'*32,source,[root],audit,rows)
    destination=next(x['destination'] for x in r['files'] if x['status']=='moved');assert Path(destination).read_bytes()==b'different' and '-' in Path(destination).stem and (folder/'C398_chart.txt').read_bytes()==b'one'
    assert not collision.exists()
    empty=tmp_path/'empty';empty.mkdir();assert m.run_batch('d'*32,empty,[root],tmp_path/'empty-audit',loader=lambda _:pytest.fail('empty must not load journal'))['ok']
    assert m.run_batch('e'*32,tmp_path/'missing',[root],tmp_path/'missing-audit',rows)['status']=='failed'
    # No destination path outside the injected Dropbox tree is ever opened.
    outside=tmp_path/'outside';outside.mkdir();f=aged(source/'C398_escape.bin');r=m.run_batch('f'*32,source,[root],tmp_path/'escape-audit',[candidate(outside)]);assert f.exists() and not list(outside.iterdir())
    f.unlink() # Only this test-owned fixture, never production data.
    # Locks prevent cleanup authority; incomplete download remains untouched.
    locked=aged(source/'C398_locked.bin')
    if os.name=='nt':
        with locked.open('rb') as held:r=m.run_batch('1'*32,source,[root],tmp_path/'locked-audit',rows)
        assert locked.exists() and any('locked' in x['file'] and x['status']=='unresolved' for x in r['files'])
    locked.unlink()
    # Source mutation/identity mismatch at a commit boundary fails closed.
    changing=aged(source/'C398_changing.bin');r=m.run_batch('2'*32,source,[root],tmp_path/'change-audit',rows,hook=lambda phase,*_:(_ for _ in ()).throw(m.Unsafe('Source identity changed')) if phase=='planned' else None);assert changing.exists() and r['status']=='failed'
    changing.unlink()
    # Copy/publication failure stops subsequent work and preserves sources.
    bad=aged(source/'C398_aa.bin');later=aged(source/'C398_zz.bin')
    with monkeypatch.context() as patch:
        patch.setattr(m,'publish',lambda *a:(_ for _ in ()).throw(OSError('injected publication failure')))
        r=m.run_batch('3'*32,source,[root],tmp_path/'copy-audit',rows)
    assert bad.exists() and later.exists() and r['counts']['failed']==1 and r['counts']['unprocessed']>=1
    bad.unlink();later.unlink()
    retained=aged(source/'C398_retained.bin')
    with monkeypatch.context() as patch:
        patch.setattr(m,'remove_source',lambda *a:(_ for _ in ()).throw(PermissionError('injected cleanup failure')))
        r=m.run_batch('4'*32,source,[root],tmp_path/'delete-audit',rows)
    assert retained.exists() and r['counts']['copied_source_retained']==1 and (folder/retained.name).exists();retained.unlink()
    corrupted=aged(source/'C398_hash.bin')
    with monkeypatch.context() as patch:
        patch.setattr(m,'exact_bytes',lambda *a:False);r=m.run_batch('5'*32,source,[root],tmp_path/'hash-audit',rows)
    assert corrupted.exists() and r['counts']['copied_source_retained']==1;corrupted.unlink()
    interrupted=aged(source/'C398_interrupt.bin');crash_audit=tmp_path/'crash-audit'
    with pytest.raises(KeyboardInterrupt):m.run_batch('6'*32,source,[root],crash_audit,rows,hook=lambda phase,*_:(_ for _ in ()).throw(KeyboardInterrupt()) if phase=='copied' else None)
    assert interrupted.exists() and (crash_audit/'active.json').exists()
    blocked=m.run_batch('7'*32,source,[root],crash_audit,rows);assert blocked['status']=='inconclusive' and interrupted.exists();interrupted.unlink()
    # Boundaries retain the unprocessed files and never remove source directories.
    aa=aged(source/'C398_limit_a.bin');bb=aged(source/'C398_limit_b.bin');r=m.run_batch('8'*32,source,[root],tmp_path/'limit-audit',rows,max_files=1)
    assert r['counts']['unprocessed']>=1 and bb.exists()
    # Audit recovery retains verified completed results on worker timeout.
    timeout_audit=tmp_path/'timeout-audit';timeout_audit.mkdir();rid='9'*32;(timeout_audit/(rid+'.json')).write_text(json.dumps({'files':[{'file':'old','status':'moved'}],'status':'running'}))
    with monkeypatch.context() as patch:
        patch.setattr(m.subprocess,'run',lambda *a,**k:(_ for _ in ()).throw(subprocess.TimeoutExpired('owned worker',295)))
        r=m.run_local_batch(rid,tmp_path/'journal.xlsx',timeout_audit)
    assert r['status']=='inconclusive' and r['counts']['moved']==1 and r['ok'] is False
    assert source.is_dir()


def test_local_menu_route_and_ui_contract(tmp_path,monkeypatch):
    from starlette.responses import JSONResponse
    tree=ast.parse((ROOT/'render/master_service.py').read_text(encoding='utf-8'));selected=[];template=''
    for node in tree.body:
        if isinstance(node,ast.Assign) and any(isinstance(x,ast.Name) and x.id=='TRADING_JOURNAL_ACTIONS_TEMPLATE' for x in node.targets):template=ast.literal_eval(node.value)
        if isinstance(node,(ast.FunctionDef,ast.AsyncFunctionDef)) and node.name in {'_run_trading_journal_file_sort','trading_journal_sort_downloaded_files','_is_local_host'}:
            node.decorator_list=[];selected.append(node)
    assert len(selected)==3
    counts={k:0 for k in ['moved','already_filed','unresolved','failed','copied_source_retained','unprocessed']};counts.update(moved=1,unresolved=1)
    payload={'ok':False,'status':'partial','counts':counts,'files':[{'file':'<img src=x>','status':'unresolved','reason':'Ambiguous <trade>'},{'file':'C398.txt','destination':'C:/Dropbox/C398/C398.txt','status':'moved','reason':'Verified'}]}
    calls=[];releases=[]
    (tmp_path/'journal.xlsx').write_bytes(b'isolated saved-journal placeholder; worker mocked')
    monkeypatch.setattr(m,'run_local_batch',lambda *args:calls.append(args) or payload)
    async def immediate_thread(fn,*args):return fn(*args)
    ns={'Dict':Dict,'Request':object,'JSONResponse':JSONResponse,'APP_PROFILE':'local','_render_blocks_path':lambda _:False,'_local_only_disabled_response':lambda path,as_json:JSONResponse({'status':'disabled'},status_code=410),'urlparse':urlparse,'re':__import__('re'),'asyncio':types.SimpleNamespace(to_thread=immediate_thread),'TRADING_JOURNAL_IMPORT_LOCK':threading.Lock(),'TRADING_JOURNAL_RESYNC_LOCK':threading.Lock(),'_reserve_master_journal_workbook_sync':lambda *a,**k:None,'_release_master_journal_workbook_sync':lambda:releases.append(True),'_master_journal_path':lambda:tmp_path/'journal.xlsx','BASE_DIR':tmp_path}
    exec(compile(ast.fix_missing_locations(ast.Module(body=selected,type_ignores=[])),'extracted-journal-sort','exec'),ns)
    def drive(coro):
        try:coro.send(None)
        except StopIteration as stopped:return stopped.value
        finally:coro.close()
        raise AssertionError('Unexpected suspension in mocked route')
    req=types.SimpleNamespace(client=types.SimpleNamespace(host='127.0.0.1'),url=types.SimpleNamespace(hostname='127.0.0.1',scheme='http',netloc='127.0.0.1:8000'),headers={'origin':'http://127.0.0.1:8000','x-ltt-action':'sort-downloaded-files'})
    ns['APP_PROFILE']='render';assert drive(ns['trading_journal_sort_downloaded_files'](req,{'request_id':'a'*32})).status_code==410 and not calls
    ns['APP_PROFILE']='local';bad=types.SimpleNamespace(**vars(req));bad.headers=dict(req.headers,origin='https://evil.invalid');assert drive(ns['trading_journal_sort_downloaded_files'](bad,{'request_id':'a'*32})).status_code==403 and not calls
    assert drive(ns['trading_journal_sort_downloaded_files'](req,{'request_id':'a'*32,'source':'C:/outside'})).status_code==422 and not calls
    lock=ns['TRADING_JOURNAL_IMPORT_LOCK'];lock.acquire()
    try:assert drive(ns['trading_journal_sort_downloaded_files'](req,{'request_id':'a'*32})).status_code==409 and not calls
    finally:lock.release()
    result=drive(ns['trading_journal_sort_downloaded_files'](req,{'request_id':'a'*32}));assert result.status_code==207 and len(calls)==1 and releases==[True]
    # Actual production markup supplies IDs to an isolated fake DOM; actual JS runs.
    fixture=tmp_path/'ui_fixture.json';fixture.write_text(json.dumps({'html':template,'payload':payload,'js':str(ROOT/'render/static/trading_journal_actions.js')}))
    script=tmp_path/'ui.js';script.write_text(r"""const fs=require('fs'),vm=require('vm'),assert=require('assert');const f=JSON.parse(fs.readFileSync(process.argv[2],'utf8'));
const elements={};function el(id){return {id,disabled:false,textContent:'',innerHTML:'',style:{},listeners:{},classList:{add(){},remove(){}},addEventListener(n,fn){this.listeners[n]=fn},after(){},click(){return this.listeners.click?.({preventDefault(){}})}};}
for(const match of f.html.matchAll(/id="([^"]+)"/g))elements[match[1]]=el(match[1]);assert(elements['sort-downloaded-trade-files-btn']);assert(f.html.includes('Sort downloaded trade files'));
let pending,requests=[],timers=new Map(),timer=0,serial=0;const storage=new Map();const window={crypto:{randomUUID:()=>String(++serial).padStart(32,'a')},sessionStorage:{getItem:k=>storage.get(k),setItem:(k,v)=>storage.set(k,v),removeItem:k=>storage.delete(k)},setTimeout:(fn)=>{timers.set(++timer,fn);return timer},clearTimeout:id=>timers.delete(id),addEventListener(){}};
const context={window,document:{getElementById:id=>elements[id]||null,createElement:()=>el('new')},fetch:(url,opts)=>{requests.push({url,opts});return new Promise((resolve,reject)=>pending={resolve,reject})},AbortController:class{constructor(){this.signal={}}abort(){}},console};vm.runInNewContext(fs.readFileSync(f.js,'utf8'),context);
(async()=>{const button=elements['sort-downloaded-trade-files-btn'],status=elements['journal-actions-status'];assert.equal(requests.length,0);const first=button.click();assert(button.disabled);await button.click();assert.equal(requests.length,1);assert.equal(requests[0].url,'/api/trading-journal/sort-downloaded-files');assert.equal(requests[0].opts.method,'POST');assert.deepEqual(Object.keys(JSON.parse(requests[0].opts.body)),['request_id']);pending.resolve({ok:true,json:async()=>f.payload});await first;assert(!button.disabled);assert(!elements['import-journal-btn'].disabled);assert(status.innerHTML.includes('Moved: 1'));assert(status.innerHTML.includes('Unresolved/ambiguous: 1'));assert(status.innerHTML.includes('&lt;img'));assert(!status.innerHTML.includes('<img'));assert.equal(status.style.color,'#fca5a5');assert.equal(timers.size,0);
const failed=button.click();const id=JSON.parse(requests.at(-1).opts.body).request_id;pending.reject(new Error('network lost'));await failed;assert(!button.disabled);assert(status.textContent.includes('No automatic retry'));const retrieve=button.click();assert.equal(JSON.parse(requests.at(-1).opts.body).request_id,id);pending.resolve({ok:true,json:async()=>({...f.payload,ok:true,status:'complete',files:[],counts:{moved:0}})});await retrieve;assert(!button.disabled);assert.equal(timers.size,0);assert.equal(requests.length,3);console.log('production menu handler contract passed');})().catch(e=>{console.error(e);process.exitCode=1});""",encoding='utf-8')
    completed=subprocess.run([r'C:\Program Files\nodejs\node.exe',str(script),str(fixture)],capture_output=True,text=True,timeout=25)
    assert completed.returncode==0,completed.stdout+'\n'+completed.stderr
