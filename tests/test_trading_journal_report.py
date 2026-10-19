"""Two bounded Excel-style print-copy regressions; native adapter is faked here."""
import hashlib
import importlib.util
import io
import sys
import zipfile
from pathlib import Path
from types import SimpleNamespace
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

spec=importlib.util.spec_from_file_location('a4_journal_report',Path(__file__).parents[1]/'tools/build_trading_journal_report.py')
report=importlib.util.module_from_spec(spec);sys.modules[spec.name]=report;spec.loader.exec_module(report)


class Range:
    def __init__(self, sheet):
        self.sheet=sheet;self.Rows=self
    def Merge(self):pass
    def UnMerge(self):pass
    def Copy(self,Destination):self.sheet.excel.events.append('copy-group-context')
    def AutoFit(self):self.sheet.excel.events.append('local-row-expansion')


class Sheet:
    def __init__(self,name,book):
        self.Name=name;self.book=book;self.excel=book.excel;self.PageSetup=SimpleNamespace()
    def Range(self,*args):return Range(self)
    def Cells(self,*args):return Range(self)
    def Copy(self,Before=None,After=None):
        assert Before is None and After is not None
        After.book.Worksheets.items.append(Sheet(self.Name,After.book))
    def Delete(self):self.book.Worksheets.items.remove(self)


class Sheets:
    def __init__(self,book,names):self.items=[Sheet(name,book) for name in names]
    @property
    def Count(self):return len(self.items)
    def __call__(self,key):
        return self.items[key-1] if isinstance(key,int) else next(s for s in self.items if s.Name==key)


class Book:
    def __init__(self,excel,names=('Blank',)):
        self.excel=excel;self.Worksheets=Sheets(self,names)
    def Close(self,SaveChanges):
        assert SaveChanges is False
        self.excel.events.append('close-discard');self.excel.Workbooks.items.remove(self)
    def ExportAsFixedFormat(self,kind,path,OpenAfterPublish):
        assert kind==0 and OpenAfterPublish is False
        self.excel.events.append('native-pdf-export')
        self.excel.export_sheets=list(self.Worksheets.items)


class Books:
    def __init__(self,excel):self.excel=excel;self.items=[]
    @property
    def Count(self):return len(self.items)
    def Add(self,template):
        assert template==-4167
        b=Book(self.excel);self.items.append(b);self.excel.events.append('blank-add');return b
    def Open(self,path,**kwargs):
        assert self.Count>0
        assert Path(path).name=='print-copy.xlsx'
        assert kwargs==dict(UpdateLinks=0,ReadOnly=True,IgnoreReadOnlyRecommended=True,AddToMru=False)
        self.excel.events.append('owned-open')
        if self.excel.fail_open:raise RuntimeError('controlled native open failure')
        b=Book(self.excel,('Grid',));self.items.append(b);return b


class Excel:
    def __init__(self,fail_open=False):
        self.events=[];self.Workbooks=Books(self);self.fail_open=fail_open
        self._printer='physical inherited context'
    @property
    def ActiveSheet(self):return self.Workbooks.items[0].Worksheets(1) if self.Workbooks.Count else None
    @property
    def ActivePrinter(self):
        assert self.Workbooks.Count>0
        return self._printer
    @ActivePrinter.setter
    def ActivePrinter(self,value):
        assert self.Workbooks.Count>0
        assert value=='Microsoft Print to PDF on Ne01:'
        self.events.append('virtual-after-blank');self._printer=value
    def __setattr__(self,key,value):
        if key in ('Calculation','CalculateBeforeSave'):
            assert self.Workbooks.Count>0
            self.events.append(key)
        object.__setattr__(self,key,value)
    def Quit(self):
        assert self.Workbooks.Count==0
        self.events.append('owned-quit')


def test_excel_style_print_plan_preserves_grid_and_continuation_context(tmp_path):
    wb=Workbook();s=wb.active;s.title='Grid'
    s.column_dimensions['A'].width=12
    for col in 'BCDEFGHIJ':s.column_dimensions[col].width=32
    s.merge_cells('D1:F1');s['D1']='Management'
    for r in range(2,11):
        s.row_dimensions[r].height=21
        for c in range(1,11):
            cell=s.cell(r,c,f'R{r}C{c}')
            cell.font=Font(name='Cambria',size=11,bold=r==2)
            cell.fill=PatternFill('solid',fgColor='4472C4' if r==2 else 'FFFFFF')
            cell.border=Border(bottom=Side(style='thin',color='202020'))
            cell.alignment=Alignment(horizontal='left')
    panels=report.plan_panels(wb,[('Grid',10,2,1,'landscape',[(3,6),(7,10)])])
    assert [(p.first_column,p.last_column) for p in panels]==[(2,3),(4,7),(8,10)]*2
    assert [(p.first_row,p.last_row) for p in panels]==[(3,6)]*3+[(7,10)]*3
    expected={(r,c) for r in range(3,11) for c in range(2,11)}
    actual={(r,c) for p in panels for r in range(p.first_row,p.last_row+1) for c in range(p.first_column,p.last_column+1)}
    assert actual==expected
    assert s['B2'].font.name=='Cambria' and s['B2'].font.sz==11 and s['B2'].font.bold
    assert s['B3'].border.bottom.style=='thin' and s['B2'].fill.fgColor.rgb=='004472C4'
    assert s.column_dimensions['B'].width==32 and s.row_dimensions[3].height==21
    excel=Excel();book=excel.Workbooks.Add(-4167)
    writes=[]
    class Settings(SimpleNamespace):
        def __setattr__(self,key,value):
            writes.append(key)
            super().__setattr__(key,value)
    book.Worksheets(1).PageSetup=Settings()
    for i,p in enumerate(panels,1):
        writes.clear()
        report.apply_native_page_setup(excel,book.Worksheets(1),p,i)
        assert writes.index('FitToPagesWide')<writes.index('Zoom')
        assert writes.index('FitToPagesTall')<writes.index('Zoom')
        settings=book.Worksheets(1).PageSetup
        assert settings.PaperSize==9 and settings.Orientation==2 and settings.Zoom==92
        assert settings.PrintTitleRows=='$1:$2' and settings.PrintTitleColumns=='$A:$A'
        assert settings.LeftMargin==28.35 and settings.TopMargin==36
        assert settings.Order==2 and settings.BlackAndWhite is False and settings.PrintHeadings
        assert f'rows {p.first_row}-{p.last_row}' in settings.CenterFooter
        assert f'panel {p.panel}/3' in settings.CenterFooter
        assert 11*settings.Zoom/100>=10
    diagnostics=[]
    settings=book.Worksheets(1).PageSetup
    class WrongZoom(Settings):
        def __setattr__(self,key,value):
            super().__setattr__(key,False if key=='Zoom' else value)
    book.Worksheets(1).PageSetup=WrongZoom()
    import pytest,json
    with pytest.raises(RuntimeError,match='Native A4 settings readback mismatch:') as failure:
        report.apply_native_page_setup(excel,book.Worksheets(1),panels[0],1,diagnostics.append)
    detail=json.loads(str(failure.value).split(': ',1)[1])
    assert detail['panel']=='P001' and detail['sheet']=='Grid'
    assert detail['fresh']['Zoom']=={'value':False,'type':'bool'}
    assert detail['expected']['Zoom']=={'value':92,'type':'int'}
    assert diagnostics[0]['stage']=='before configuration'
    assert diagnostics[-1]['stage']=='after PrintCommunication commit'
    assert str(s.merged_cells)=='D1:F1'
    book.Close(SaveChanges=False);excel.Quit()
    wb.save(tmp_path/'print-copy.xlsx');wb.close()
    adapter=Excel();stages=[]
    report.native_preview(tmp_path/'print-copy.xlsx',tmp_path,panels,
                          'Microsoft Print to PDF on Ne01:',factory=lambda:adapter,progress=stages.append)
    assert len(adapter.export_sheets)==len(panels)
    assert stages.index('read-only print-copy Open')<stages.index('P001 source worksheet Copy')
    assert stages.index('P001 source worksheet Copy')<stages.index('P001 PageSetup configuration and committed readback')
    assert stages[-1]=='ExportAsFixedFormat'


def test_print_copy_preserves_source_values_styles_and_grayscale_structure(tmp_path):
    import re
    from openpyxl.comments import Comment
    wb=Workbook();s=wb.active;s.title='Grid';s.merge_cells('A1:C1');s['A1']='Group heading'
    s['A3']=0;s['B3']=False;s['C3']=-12.345;s['D3']=0.0000000123
    s['C3'].number_format='0.000" AUD"';s['D3'].number_format='0.0000000000'
    s['A4']='=1-1';s['B4']='=2+2';s['C4']='=3+3';s['D4']='Full long note  '*100+'END'
    s['D4'].comment=Comment('Unrelated synthetic comment remains unchanged','Fixture')
    s['A1'].font=Font(name='Cambria',size=11,bold=True,color='FFFFFF')
    s['A1'].fill=PatternFill('solid',fgColor='112288');s['B3'].fill=PatternFill('solid',fgColor='F0D050')
    s['C3'].border=Border(bottom=Side(style='double',color='006600'))
    s.column_dimensions['D'].width=40;s.row_dimensions[4].height=34
    original=tmp_path/'source.xlsx';wb.save(original)
    source=original.read_bytes();buffer=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source)) as z,zipfile.ZipFile(buffer,'w') as out:
        for item in z.infolist():
            data=z.read(item.filename)
            if item.filename=='xl/worksheets/sheet1.xml':
                root=ET.fromstring(data)
                for cell in root.findall('.//s:c',report.NS):
                    if cell.get('r') in ('A4','B4'):cell.find('s:v',report.NS).text='0' if cell.get('r')=='A4' else '4'
                data=ET.tostring(root)
            if item.filename in ('xl/workbook.xml','xl/styles.xml','xl/worksheets/sheet1.xml'):
                # Prefixes are deliberately absent from element/attribute NAMES:
                # they occur only in Office compatibility string values.
                declaration=(b' xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
                    b' xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac"'
                    b' xmlns:future="urn:journal-print-future" mc:Ignorable="x14ac future"'
                    b' mc:PreserveElements="future:*" mc:PreserveAttributes="future:*"'
                    b' mc:ProcessContent="future:wrapper"')
                data=re.sub(rb'(<[^!?][^\s/>]*)',lambda m:m.group(1)+declaration,data,count=1)
                close=data.rfind(b'</')
                data=data[:close]+(b'<mc:AlternateContent><mc:Choice Requires="future"/>'
                                  b'<mc:Fallback/></mc:AlternateContent>')+data[close:]
            out.writestr(item,data)
    original.write_bytes(buffer.getvalue());before=original.read_bytes();mtime=original.stat().st_mtime_ns
    capture=report.capture_workbook(original,tmp_path)
    copied=report.prepare_print_copy(capture,tmp_path)
    assert copied!=original and copied.parent==tmp_path
    values=load_workbook(copied,data_only=True);v=values['Grid']
    assert [v.cell(3,c).value for c in range(1,5)]==[0,False,-12.345,.0000000123]
    assert v['A4'].value==0 and v['B4'].value==4 and v['C4'].value=='[stored result unavailable]'
    assert v['D4'].value=='Full long note  '*100+'END'
    assert v['A1'].font.name=='Cambria' and v['A1'].font.sz==11 and v['A1'].font.bold
    assert v['A1'].font.color.rgb=='FFFFFFFF'
    assert v['C3'].border.bottom.style=='double'
    assert v['C3'].number_format=='0.000" AUD"' and v['D3'].number_format=='0.0000000000'
    assert str(v.merged_cells)=='A1:C1' and v.column_dimensions['D'].width==40 and v.row_dimensions[4].height==34
    colors=[v['A1'].fill.fgColor.rgb,v['B3'].fill.fgColor.rgb]
    assert colors[0]!=colors[1]
    for color in colors:assert color[-6:-4]==color[-4:-2]==color[-2:]
    values.close()
    with zipfile.ZipFile(copied) as z:
        assert not ET.fromstring(z.read('xl/worksheets/sheet1.xml')).findall('.//s:f',report.NS)
        with zipfile.ZipFile(io.BytesIO(before)) as original_zip:
            expected_changed={'xl/workbook.xml','xl/styles.xml','xl/worksheets/sheet1.xml'}
            assert set(z.namelist())==set(original_zip.namelist())
            for name in z.namelist():
                if name not in expected_changed:
                    assert z.read(name)==original_zip.read(name),name
            def compatibility_references(data):
                scopes=[];pending={};references=[]
                for event,item in ET.iterparse(io.BytesIO(data),events=('start-ns','start','end')):
                    if event=='start-ns':pending[item[0]]=item[1]
                    elif event=='start':
                        scope=dict(scopes[-1]) if scopes else {}
                        scope.update(pending);pending={};scopes.append(scope)
                        for attr,value in item.attrib.items():
                            if attr.startswith('{http://schemas.openxmlformats.org/markup-compatibility/2006}') or (attr=='Requires' and item.tag.endswith('}Choice')):
                                for token in value.split():
                                    prefix=token.split(':',1)[0]
                                    assert prefix in scope,(attr,token)
                                    references.append((attr,token,scope[prefix]))
                    else:scopes.pop()
                return references
            for name in expected_changed:
                expected=compatibility_references(original_zip.read(name))
                assert expected and compatibility_references(z.read(name))==expected
    panels=[report.PrintPanel('Grid',2,4,3,4,2,1,'landscape',1,1,1)]
    excel=Excel();report.native_preview(copied,tmp_path,panels,'Microsoft Print to PDF on Ne01:',factory=lambda:excel)
    assert excel.events.index('blank-add')<excel.events.index('Calculation')<excel.events.index('virtual-after-blank')<excel.events.index('owned-open')
    assert excel.events.count('native-pdf-export')==1 and excel.events.count('close-discard')==3
    assert excel.events[-1]=='owned-quit'
    failed=Excel(True)
    try:report.native_preview(copied,tmp_path,panels,'Microsoft Print to PDF on Ne01:',factory=lambda:failed)
    except RuntimeError as error:assert str(error)=='controlled native open failure'
    else:raise AssertionError('Expected controlled failure')
    assert failed.events[-2:]==['close-discard','owned-quit']
    for path,context in [(original,'Microsoft Print to PDF on Ne01:'),(copied,'Brother on Ne03:'),(copied,'network-printer')]:
        try:report.native_preview(path,tmp_path,panels,context,factory=lambda:(_ for _ in ()).throw(AssertionError('Must reject before native creation')))
        except ValueError:pass
        else:raise AssertionError('Unsafe path/context accepted')
    assert original.read_bytes()==before and original.stat().st_mtime_ns==mtime
    assert capture.sha256==hashlib.sha256(before).hexdigest()
