"""Two bounded 7C1 nodes; fixtures contain no private recovery/journal data."""
import copy
import datetime as dt
import json
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from tools import backfill_forex_word_metadata as m


HEADERS = ["Trade Number", "Row ID", "Row Type", "Symbol", "Open Time", "Side", "Entry Price", "Setup", "Timeframe", "EMA", "Round Number", "Move to Break Even Distance From Entry %", "Move to Profit Trigger Price", "Pattern", "Order"]


def workbook(path, rows=1, overrides=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Trade Log"
    ws.append(HEADERS)
    for i in range(1, rows + 1):
        ws.append([f"F{i:03}", f"row-{i}", "TRADE", "EURUSD", dt.datetime(2018, 9, 10, 12), "Buy", 1.23])
        ws.cell(i+1, 12).number_format = "0.00%"
    for (row, header), value in (overrides or {}).items():
        ws.cell(row, HEADERS.index(header)+1, value)
    wb.save(path)
    return wb


def source(number=1, pairs=None):
    pairs = pairs or [("PAIR", "EUR/USD"), ("SHORT/LONG", "LONG"), ("DATE", "10/9/18"), ("ENTRY PRICE", "1.23000"), ("SETUP", "α\nβ"), ("TIMEFRAME", "5 MIN"), ("ROUND NUMBER", "No"), ("MOVE TO BREAK EVEN DISTANCE FROM ENTRY %", "0%"), ("MOVE TO PROFIT TRIGGER PRICE", "1.23000"), ("MOVING AVERAGE", "50"), ("SIGNAL CANDLE", "BULLISH ENGULFING")]
    cells = []
    for row, (label, raw) in enumerate(pairs):
        for col, text in enumerate((label, raw)):
            cells.append(dict(row=row, column=col, rowspan=1, colspan=1, text=text))
    doc = dict(source=dict(relative_path=f"2018/F{number}/sample.docx", parent_relative_path=f"2018/F{number}", f_number_evidence=[f"F{number}"], symbol_evidence=[], year="2018", kind="document", extension=".docx", size=12, mtime_ns=123, availability="metadata_local_candidate", skipped_link=False), source_sha256="a"*64, complete=True, capture="not-opened.docx", before=dict(size=12, mtime_ns=123), after=dict(size=12, mtime_ns=123), outputs=[], coverage=dict(tables=[dict(id="table_01", part="word/document.xml", grid=[100,100], row_count=len(pairs), cells=cells)], prose=[], issues=[], table_coverage=[dict(raw_table=1, output_table="table_01", alternate_alias=False)]))
    return dict(identity=f"2018/f{number}/sample.docx", document=doc, preparation="preparation.json", pointer="/documents/0", provenance="provenance.json")


def rewrite_xml(path, edit):
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
    edit(root)
    parts["xl/worksheets/sheet1.xml"] = ET.tostring(root)
    with zipfile.ZipFile(path, "w") as z:
        for name, raw in parts.items():
            z.writestr(name, raw)


def test_word_metadata_preview_requires_exact_identity_and_preserves_nonblank_cells(tmp_path):
    origin = r"C:\GPT\CODEX-master\journal\Trading Journal.xlsx"
    expected = r"C:\Users\User\Dropbox\FOREX\2018\F1"
    for target, base in [
        ("../../../Users/User/Dropbox/FOREX/2018/F1", None),
        ("../../../Users/User/Dropbox/FOREX/2018/./F1", None),
        (expected, None),
        ("file:///C:/Users/User/Dropbox/FOREX/2018/F1", None),
        ("2018/F1", r"C:\Users\User\Dropbox\FOREX"),
    ]:
        result = m.normalize_folder_link(target, origin, base)
        assert result["status"] == "normalized" and result["normalized_target"] == expected
    unicode_target = "file:///C:/Users/User/Dropbox/FOREX/2018/F1%20%C3%A9"
    assert m.normalize_folder_link(unicode_target, origin)["normalized_target"].endswith("F1 é")
    for target in ["C:relative", "/Users/User/Dropbox/FOREX/2018/F1", "//server/share", "https://host/FOREX/F1", "file://server/share", expected+"?x", expected+"#x", expected+"%2520", expected+"%xy", "../../../outside/F1"]:
        assert m.normalize_folder_link(target, origin)["status"] == "unsupported"
    assert m.normalize_folder_link("2018/F1", origin, "relative-base")["status"] == "unsupported"
    assert m.normalize_folder_link("../../../Users/User/Dropbox/FOREX/2018/F1", r"D:\Temp\copy.xlsx")["status"] == "unsupported"
    p = tmp_path / "snapshot.xlsx"
    wb = workbook(p, rows=10)
    ws = wb["Trade Log"]
    for row, value in [(3, 0), (4, False), (5, " "), (6, ""), (7, '=IF(1,"","")')]:
        ws.cell(row, 8, value)
    ws.cell(2, 8).font = Font(italic=True)  # genuinely style-only
    ws.merge_cells("H8:I8")  # I8 merged continuation
    ws["H9"] = "=1"
    ws["H10"] = "=1"
    wb.save(p)

    def shared(root):
        for c in root.iter(m.NS + "c"):
            if c.get("r") == "H9":
                f = c.find(m.NS + "f")
                f.set("t", "shared"); f.set("si", "0"); f.set("ref", "H9:H10")
            if c.get("r") == "H10":
                f = c.find(m.NS + "f")
                f.text = None; f.set("t", "shared"); f.set("si", "0")
    rewrite_xml(p, shared)
    before = p.read_bytes()
    snap = m.Snapshot(p)
    sources = [source(i) for i in range(1, 11)]
    audit, catalog, patch, decisions = m.plan(sources, snap)
    assert all(x["status"] == "proven" for x in audit)
    assert snap.target(2, "Setup")["structurally_blank"]
    assert snap.target(11, "Setup")["structurally_blank"]  # absent cell
    assert all(not snap.target(row, "Setup")["structurally_blank"] for row in range(3, 8))
    assert snap.target(8, "Timeframe")["reason"] == "merged_continuation"
    assert not snap.target(9, "Setup")["structurally_blank"] and not snap.target(10, "Setup")["structurally_blank"]
    assert not any(x["header"] == "Setup" and x["row"] in range(3, 8) for x in patch)
    assert any(x["new_value"]["value"] == "No" for x in patch)
    assert any(x["new_value"] == dict(type="decimal", value="0", unit="fraction_of_one") for x in patch)
    assert any(x["new_value"]["value"] == "1.23000" for x in patch)
    assert any(x["new_value"]["value"] == "α\nβ" for x in patch)
    assert not any(x["header"] in {"EMA", "Pattern", "Entry Price"} for x in patch)
    assert len({(x["sheet"], x["coordinate"]) for x in patch}) == len(patch)
    assert p.read_bytes() == before
    # No fuzzy fallback: contradictory, duplicated, missing or formula identity.
    observations = m.fields(source())
    conflicting = source(pairs=[("PAIR", "GBPUSD"), ("SETUP", "Pullback")])
    assert m.match_source(conflicting, m.fields(conflicting), snap)["status"] == "unresolved"
    for label, value in [("ENTRY PRICE", "1.23001"), ("DATE", "11/9/18"), ("SHORT/LONG", "SHORT")]:
        c = source(pairs=[("PAIR", "EURUSD"), (label, value), ("SETUP", "Pullback")])
        assert m.match_source(c, m.fields(c), snap)["status"] == "unresolved"
    other = source(99)
    assert m.match_source(other, m.fields(other), snap)["status"] == "unresolved"
    snap.rows.append(copy.deepcopy(snap.rows[0]))
    assert m.match_source(source(), observations, snap)["status"] == "unresolved"
    snap.rows.pop()
    snap.rowids["row-1"] = 2
    assert m.match_source(source(), observations, snap)["status"] == "unresolved"
    snap.rowids["row-1"] = 1
    snap.ws["B2"] = "=1"
    assert m.match_source(source(), observations, snap)["status"] == "unresolved"
    assert m.trade_number('=HYPERLINK("C:\\folder","f0001")') == ("F1", False, "C:\\folder")
    assert m.trade_number('=CONCAT("F",1)')[1]
    # Repeated equal observations retain all origins; alias table coverage is not
    # another table. Differing occurrences never select a convenient value.
    snap = m.Snapshot(p)
    repeated = source(pairs=[("PAIR", "EURUSD"), ("SETUP", "same"), ("SETUP", "same")])
    repeated["document"]["coverage"]["table_coverage"].append(dict(raw_table=2, output_table="table_01", alternate_alias=True))
    planned = m.plan([repeated], snap)[2]
    assert len(planned) == 1 and len(planned[0]["sources"]) == 2
    repeated["document"]["coverage"]["tables"][0]["cells"][-1]["text"] = "different"
    assert not m.plan([repeated], snap)[2]
    assert any(x["status"] == "conflicting_source_field" for x in m.plan([repeated], snap)[3])
    with pytest.raises(ValueError, match="duplicate source"):
        m.plan([source(), source()], snap)
    assert m.mapping(dict(label="EMA", raw=False), HEADERS)["canonical"]["value"] == "No"
    assert m.mapping(dict(label="MOVE TO BREAK EVEN DISTANCE FROM ENTRY %", raw="12 pips"), HEADERS)["status"] == "ambiguous"
    assert m.mapping(dict(label="TIMEFRAME", raw="1/5MIN"), HEADERS)["status"] == "ambiguous"


def receipt_fixture(root):
    directory = root / "receipt"
    directory.mkdir()
    invdir = root / "inventory"
    invdir.mkdir()
    s = source()
    doc = s["document"]
    doc["capture"] = str(directory / "document_01" / "captured.docx")
    inventory = invdir / "manifest.jsonl"
    inventory.write_text(json.dumps(doc["source"]) + "\n", encoding="utf-8")
    selection = dict(root=str(root / "FOREX"), inventory=str(inventory), inventory_sha256=m.digest(inventory.read_bytes()), mode="batch", batch_limit=20, selected_count=1, documents=[doc["source"]])
    m.save(directory / "selection.validated.json", selection)
    selection_hash = m.digest((directory / "selection.validated.json").read_bytes())
    prep = dict(status="complete", selection_sha256=selection_hash, source_count=1, tool_sha256=m.RECOVERY_HASH, documents=[doc])
    m.save(directory / "preparation.json", prep)
    m.save(directory / "preflight.json", dict(tool_sha256=m.RECOVERY_HASH))
    m.save(directory / "publication.json", dict(status="complete", selection_sha256=selection_hash, source_checks=[dict(source=str(root / "FOREX" / "2018" / "F1" / "sample.docx"), sha256=doc["source_sha256"], before=doc["before"], after=doc["after"])], published=[]))
    m.save(directory / "published_verification.json", dict(status="complete", documents=1, published_files=0, counts={"picture": 0, "table": 0, "text": 0}))
    (directory / "document_01").mkdir()
    m.save(directory / "document_01" / "provenance.json", doc)
    refs = [dict(directory=str(directory), tool_sha256=m.RECOVERY_HASH, records={name: m.digest((directory/name).read_bytes()) for name in ["selection.validated.json", "preparation.json", "publication.json", "published_verification.json", "preflight.json"]})]
    context = dict(root=selection["root"], inventory=str(inventory), inventory_sha256=selection["inventory_sha256"], prior_completions=refs)
    m.save(directory / "final_audit_context.json", context)
    coverage = dict(status="complete", exact_set_equality=True, completed_receipts=refs)
    m.save(directory / "final_recovery_coverage.json", coverage)
    m.save(directory / "cumulative_progress.json", dict(remaining=0, completed_receipts=refs))
    handoff = dict(document_count=1, inventory=dict(sha256=selection["inventory_sha256"]), coverage_audit=dict(path=str(directory / "final_recovery_coverage.json"), sha256=m.digest((directory/"final_recovery_coverage.json").read_bytes())), documents=[dict(source_identity=s["identity"], source=doc["source"], source_sha256=doc["source_sha256"], receipt_directory=str(directory), receipt_hashes=refs[0]["records"], historical_tool_sha256=m.RECOVERY_HASH, preparation_reference=dict(path=str(directory/"preparation.json"), sha256=refs[0]["records"]["preparation.json"], document_index=0, json_pointer="/documents/0"), provenance_reference=dict(path=str(directory/"document_01"/"provenance.json")), structured_tables=[dict(table_id="table_01", json_pointer="/documents/0/coverage/tables/0")], coverage_reference=dict(path=str(directory/"preparation.json"), json_pointer="/documents/0/coverage"))])
    m.save(directory / "final_recovery_handoff.json", handoff)
    return directory, invdir, handoff


def test_word_metadata_preview_validates_receipts_and_produces_hash_bound_patch(tmp_path):
    from tools import verify_forex_word_metadata_preview as v
    directory, invdir, h = receipt_fixture(tmp_path)
    handoff = directory / "final_recovery_handoff.json"
    output = tmp_path / "preview"
    output.mkdir()
    snapshot = output / "snapshot.xlsx"
    workbook(snapshot)
    before = {p: p.read_bytes() for p in [*directory.rglob("*.json"), invdir/"manifest.jsonl", snapshot]}
    args = dict(approved=[directory], inventory_directory=invdir, expected_count=1)
    counts = m.run_preview(handoff, snapshot, output, **args)
    assert counts["documents"] == counts["proven_matches"] == 1 and counts["proposed_cells"] > 2
    patch = json.loads((output / "proposed_cell_patch.json").read_bytes())
    assert patch["adapter_sha256"] == m.digest(Path(m.__file__).read_bytes())
    assert patch["snapshot_sha256"] == m.digest(before[snapshot])
    assert patch["mapping_sha256"] == m.digest((output / "field_mapping_catalog.json").read_bytes())
    assert all(e["target_precondition"]["structurally_blank"] for e in patch["entries"])
    assert all(p.read_bytes() == raw for p, raw in before.items())
    verified = v.verify(output, output/"independent.json", patch["adapter_sha256"], **args)
    assert verified["status"] == "complete" and verified["eligible_set_exact"]
    assert verified["max_reads_per_file"] == verified["max_hashes_per_file"] == verified["max_parses_per_file"] == 1
    assert all(max(c.values()) <= 1 for c in verified["cache_counters"].values())
    cache = m.JsonStore([directory])
    actual = m.digest(before[handoff])
    cache.load(handoff, actual); cache.load(handoff, actual)
    with pytest.raises(ValueError, match="hash mismatch"):
        cache.load(handoff, "f"*64)
    assert cache.counters[str(handoff)] == dict(reads=1, hashes=1, parses=1)
    cache.close()
    original_patch = (output/"proposed_cell_patch.json").read_bytes()
    for mutate in [
        lambda p: p.update(entries=[]),
        lambda p: p["entries"].append(copy.deepcopy(p["entries"][0])),
        lambda p: p["entries"][0].update(row_id="wrong"),
        lambda p: p["entries"][0]["target_precondition"].update(style_id=999),
        lambda p: p["entries"][0]["new_value"].update(value="wrong"),
    ]:
        altered = json.loads(original_patch); mutate(altered)
        (output/"proposed_cell_patch.json").write_bytes(m.encoded(altered))
        with pytest.raises(ValueError):
            v.verify(output, tmp_path/"must-not-exist.json", patch["adapter_sha256"], **args)
    (output/"proposed_cell_patch.json").write_bytes(original_patch)
    with pytest.raises(ValueError, match="malformed"):
        v.pointer({}, "/x/~bad")
    # A second independent destination reproduces the proposal without altering
    # inputs; only the private input receipt carries absolute snapshot location.
    other = tmp_path / "other"
    other.mkdir()
    (other / "snapshot.xlsx").write_bytes(before[snapshot])
    m.run_preview(handoff, other/"snapshot.xlsx", other, **args)
    assert (other / "proposed_cell_patch.json").read_bytes() == (output / "proposed_cell_patch.json").read_bytes()
    # Explicit legacy identity: same trusted synthetic sources/snapshot but old
    # conservative relative-link policy, and fixed historical adapter identity.
    legacy_dir = tmp_path/"legacy"
    legacy_dir.mkdir()
    old_book = legacy_dir/"snapshot.xlsx"
    old_wb = workbook(old_book)
    old_wb["Trade Log"]["A2"].hyperlink = "../../../Users/User/Dropbox/FOREX/2018/F1"
    old_wb.save(old_book)
    m.run_preview(handoff, old_book, legacy_dir, **args)
    old_proposal = json.loads((legacy_dir/"proposed_cell_patch.json").read_bytes())
    old_receipt = json.loads((legacy_dir/"input_receipt.json").read_bytes())
    old_proposal["adapter_sha256"] = old_receipt["adapter_sha256"] = v.OLD_HASH
    (legacy_dir/"proposed_cell_patch.json").write_bytes(m.encoded(old_proposal))
    (legacy_dir/"input_receipt.json").write_bytes(m.encoded(old_receipt))
    old_report = v.verify(legacy_dir, legacy_dir/"checked.json", v.OLD_HASH, legacy=True, **args)
    assert old_report["proven_matches"] == old_report["proposed_cells"] == 0
    assert not old_report["legacy_identity_interpretation_fully_validated"]
    corrected = tmp_path/"corrected"
    corrected.mkdir(); (corrected/"snapshot.xlsx").write_bytes(old_book.read_bytes())
    capture = dict(source=r"C:\GPT\CODEX-master\journal\Trading Journal.xlsx", snapshot=str(old_book), before=dict(size=old_book.stat().st_size,mtime_ns=123), after=dict(size=old_book.stat().st_size,mtime_ns=123), sha256=m.digest(old_book.read_bytes()),successful_copies=1,read_only=True,share_read_only=True)
    m.save(corrected/"snapshot_receipt.json", capture)
    m.run_preview(handoff, corrected/"snapshot.xlsx", corrected, **args)
    new_report = v.verify(corrected, corrected/"checked.json", patch["adapter_sha256"], **args)
    assert new_report["proven_matches"] == 1 and new_report["proposed_cells"] > 0
    for mutate, message in [
        (lambda x: x["documents"][0].update(source_sha256="b"*64), "source hash"),
        (lambda x: x["documents"][0]["preparation_reference"].update(json_pointer="/documents/1"), "pointer"),
        (lambda x: x["documents"][0]["provenance_reference"].update(path=str(tmp_path/"escape.json")), "provenance path"),
        (lambda x: x["documents"].append(copy.deepcopy(x["documents"][0])), "count"),
    ]:
        altered = copy.deepcopy(h)
        mutate(altered)
        handoff.write_bytes(m.encoded(altered))
        with pytest.raises(ValueError, match=message):
            m.load_sources(handoff, **args)
    handoff.write_bytes(before[handoff])
    (directory / "preparation.json").write_bytes(before[directory/"preparation.json"] + b" ")
    with pytest.raises(ValueError, match="hash mismatch"):
        m.load_sources(handoff, **args)
    with pytest.raises(ValueError, match="escape"):
        m.JsonStore([directory]).load(invdir/"manifest.jsonl")
    assert len(json.loads((output/"document_match_audit.json").read_bytes())["documents"]) == 1
    assert "SKIP" in (output/"review.txt").read_text(encoding="utf-8")
