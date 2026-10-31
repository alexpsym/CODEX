from pathlib import Path
import json
import pytest
from openpyxl import load_workbook
from tools import backfill_counter_trend as backfill
from tools import master_journal_workbook as m


def test_cached_counter_trend_evidence_and_matching(tmp_path):
    cache=tmp_path/'cache';cache.mkdir()
    rows=[{'id':'one','trade_number':'C398','open_time':'2026-01-01','symbol':'BTCUSDT','account':'A','link':str(tmp_path/'CRYPTO/2026/C398'),'counter_trend':''},
          {'id':'old','trade_number':'C398','open_time':'2025-01-01','symbol':'BTCUSDT','account':'B','link':str(tmp_path/'CRYPTO/2025/C398'),'counter_trend':''}]
    base={'folder':str(tmp_path/'CRYPTO/2026/C398'),'root':str(tmp_path/'CRYPTO'),'trade_number':'C398','year':2026,'folder_label':'C398','status':'captured','sha256':'fixture','filename':'counter-trend.txt','path':'synthetic','text':''}
    records=[base,dict(base,filename='notes.md',text='This trade was counter trend.'),dict(base,folder=str(tmp_path/'CRYPTO/2025/C398'),year=2025,filename='note.txt',text='This trade was not counter trend.')]
    (cache/'notes.jsonl').write_text(''.join(json.dumps(r)+'\n' for r in records));(cache/'coverage.json').write_text(json.dumps({'complete':False,'errors':['one unavailable note']}))
    report=backfill.analyze_capture(cache,rows,tmp_path)
    assert {r['id']:r['counter_trend'] for r in report['changes']}=={'one':'Yes','old':'No'}
    assert report['coverage']['complete'] is False and len(report['c398'])==3
    for text in ['Could this be counter trend?','Maybe counter trend','Example: This trade was counter trend.','> This trade was counter trend.','A lesson about counter trend']:
        assert backfill.classify_note(dict(base,filename='discussion.txt',text=text))[0] is None
    assert backfill.classify_note(dict(base,filename='template.txt',text='Counter Trend: Yes'))[0] is None
    conflict=backfill.classify_note(dict(base,text='This trade was not counter trend.'));assert conflict[0] is None and conflict[2]=='conflicting evidence'
    changed=[dict(rows[0],counter_trend='No'),rows[1]];conflicted=backfill.analyze_capture(cache,changed,tmp_path);assert [r['id'] for r in conflicted['changes']]==['old'];assert conflicted['ledger'][0]['decision']=='conflict'
    duplicates=[dict(rows[0],link='',id='x'),dict(rows[0],link='',id='y')]
    assert backfill.match_folder(dict(base,folder_label='C398 BTCUSDT'),duplicates,tmp_path)[0] is None
    assert backfill.folder_identity(tmp_path/'CRYPTO/2026/OCTOBER/C398/nested',tmp_path/'CRYPTO')['year']==2026
    # Unknown and unreadable content never imply No.
    assert backfill.classify_note(dict(base,filename='missing.pdf',text='',status='unreadable'))[0] is None


    # XLSX is parsed from cached bytes: no original workbook path is reopened.
    import io
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    book=Workbook();sheet=book.active;sheet.title='Trade notes'
    sheet.append(['Trade Number','Counter Trend','Notes'])
    sheet.append(['C398','Yes','known setup']);sheet.append(['C399','No','a different trade'])
    sheet['C2'].comment=Comment('This trade was counter trend.','test')
    sheet['D2']='=1/0'  # Formula is not evaluated or treated as note evidence.
    raw=io.BytesIO();book.save(raw);book.close();parsed=backfill.extract_xlsx(raw.getvalue())
    assert 'Trade notes!B2: Yes' in parsed['text'] and '=1/0' not in parsed['text']
    assert any(f['cell']=='C2#comment' and f['trade_number']=='C398' for f in parsed['xlsx_facts'])
    xlcache=tmp_path/'xlsx-cache';xlcache.mkdir();(xlcache/'coverage.json').write_text('{"complete":true}')
    note=dict(base,filename='notes.xlsx',**parsed);(xlcache/'notes.jsonl').write_text(json.dumps(note)+'\n')
    report=backfill.analyze_capture(xlcache,rows,tmp_path);assert report['changes']==[{'id':'one','counter_trend':'Yes','expected_counter_trend':''}]
    assert all(e.get('sheet')=='Trade notes' and e.get('cell') for e in report['ledger'][0]['evidence'])
    note['filename']='template.xlsx';(xlcache/'notes.jsonl').write_text(json.dumps(note)+'\n');assert not backfill.analyze_capture(xlcache,rows,tmp_path)['changes']
    note['filename']='notes.xlsx';note['xlsx_facts'].append({'trade_number':'C398','stable_id':None,'sheet':'Trade notes','cell':'C2#comment','text':'This trade was not counter trend.'})
    (xlcache/'notes.jsonl').write_text(json.dumps(note)+'\n');assert not backfill.analyze_capture(xlcache,rows,tmp_path)['changes']


def test_counter_trend_combined_candidate_preserves_refresh(tmp_path,monkeypatch):
    from test_master_journal_workbook import sample_snapshot
    monkeypatch.setattr(m,'REPORT_START_YEAR',2026);monkeypatch.setattr(m,'REPORT_MIN_END_YEAR',2026)
    monkeypatch.setattr(m,'assign_trade_numbers_and_create_folders',lambda *a,**k:{})
    snapshot=sample_snapshot();path=tmp_path/'journal.xlsx';m.build_master_journal_workbook(snapshot,path,publish_recommendation_assets=False)
    candidate=tmp_path/'candidate.xlsx'
    result=backfill.prepare_backfill(path,candidate,[{'id':'t1','counter_trend':'Yes','expected_counter_trend':''},{'id':'t2','counter_trend':'No','expected_counter_trend':''}]);assert result['preservation_verified']
    wb=load_workbook(candidate);h=m._trade_log_header_map(wb['Trade Log']);assert h['Counter Trend']==h['Early Close']+1==h['Setup']-1
    stats=wb['STATS1'];row=next(r for r in range(1,stats.max_row+1) if stats.cell(r,1).value=='Counter Trend');assert stats.cell(row,2).value==1;wb.close()
    refreshed=m.update_master_journal_workbook_data_only(candidate,snapshot,preserve_existing_layout=True,publish_recommendation_assets=False)
    assert refreshed['ok'];final=Path(refreshed['candidate_path']);items={r['id']:r for r in m.read_master_journal_source(final)['items']};assert items['t1']['counter_trend']=='Yes' and items['t2']['counter_trend']=='No';assert items['t1']['entry_price']==snapshot['items'][0]['entry_price']
    wb=load_workbook(final);h=m._trade_log_header_map(wb['Trade Log']);assert h['Counter Trend']==h['Early Close']+1==h['Setup']-1;wb.close()
    cache=tmp_path/'cache';cache.mkdir();(cache/'coverage.json').write_text('{"complete":true}')
    note={'folder':str(tmp_path/'C1'),'trade_number':'C1','year':2026,'folder_label':'C1','filename':'counter trend.txt','path':'fixture','status':'captured','text':''};(cache/'notes.jsonl').write_text(json.dumps(note)+'\n')
    assert backfill.analyze_capture(cache,[dict(items['t1'],link=str(tmp_path/'C1'))],tmp_path)['changes']==[]
    with pytest.raises(RuntimeError):backfill.prepare_backfill(final,tmp_path/'stale.xlsx',[{'id':'t1','counter_trend':'No','expected_counter_trend':''}])
