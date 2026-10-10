from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from tools.master_journal_workbook import TRADE_LOG_HEADERS, _TRADE_FOLDER_INDEX_CACHE
from tools import repair_bybit_trade_chart_folder as repairer


ROW_ID = repairer.EXPECTED_ROW_ID


def _fixture(tmp_path: Path):
    workbook = tmp_path / "Trading Journal.xlsx"
    json_path = tmp_path / "trading_journal.json"
    root = tmp_path / "CRYPTO"
    source = root / "2026" / "APRIL" / "C377"
    source.mkdir(parents=True)
    (source / "chart.png").write_bytes(b"original-chart-bytes")
    (source / ".capture-meta").write_bytes(b"hidden-entry")
    nested = source / "notes"
    nested.mkdir()
    (nested / "caption.txt").write_text("manual capture note", encoding="utf-8")

    wb = Workbook()
    ws = wb.active
    ws.title = "Trade Log"
    for col, header in enumerate(TRADE_LOG_HEADERS, start=1):
        ws.cell(1, col, header)
    row = {
        "Trade Number": "C377", "Open Time": datetime(2026, 10, 4, 13, 7, 22),
        "Close Time": datetime(2026, 10, 7, 12, 1, 7), "Account": "Bybit Demo",
        "Symbol": "BTCUSDT", "Side": "Buy", "Qty": 0.004, "Entry Price": 84763.8,
        "Exit Price": 84100.3, "Net P/L": -3.13515409, "Notes": "keep this note",
        "Row ID": ROW_ID,
    }
    for col, header in enumerate(TRADE_LOG_HEADERS, start=1):
        if header in row:
            ws.cell(2, col, row[header])
    number_cell = ws.cell(2, 1)
    number_cell.hyperlink = source.resolve().as_uri()
    number_cell.style = "Hyperlink"
    ws.cell(2, 45).font = Font(bold=True, color="123456")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{ws.cell(2, len(TRADE_LOG_HEADERS)).coordinate}"
    presentation = wb.create_sheet("Presentation")
    presentation["A1"] = "Authored label"
    presentation["A2"] = "manual link"
    presentation["A2"].hyperlink = "https://example.invalid/manual"
    presentation["B2"] = "=1+1"
    presentation.merge_cells("A4:B4")
    presentation["A4"] = "merged authored cell"
    wb.save(workbook)
    json_path.write_text(json.dumps({"items": [{
        "id": ROW_ID, "trade_number": "C377", "account": "Bybit Demo", "symbol": "BTCUSDT",
        "open_time": "2026-10-04T13:07:22+10:00", "close_time": "2026-10-07T12:01:07+10:00",
        "qty": 0.004, "entry_price": 84763.8, "exit_price": 84100.3, "net_profit": -3.13515409,
    }]}), encoding="utf-8")
    return workbook, json_path, root, source


def _workbook_snapshot(path: Path):
    wb = load_workbook(path, data_only=False, read_only=False)
    try:
        trade = wb["Trade Log"]
        row = next(i for i in range(2, trade.max_row + 1) if trade.cell(i, 50).value == ROW_ID)
        return {
            "number": trade.cell(row, 1).value,
            "hyperlink": trade.cell(row, 1).hyperlink.target,
            "style": trade.cell(row, 1).style_id,
            "note": trade.cell(row, 45).value,
            "note_font": trade.cell(row, 45).font.bold,
            "open": trade.cell(row, 2).value,
            "close": trade.cell(row, 3).value,
            "formula": wb["Presentation"]["B2"].value,
            "manual_link": wb["Presentation"]["A2"].hyperlink.target,
            "merged": tuple(str(x) for x in wb["Presentation"].merged_cells.ranges),
            "freeze": trade.freeze_panes,
            "filter": trade.auto_filter.ref,
        }
    finally:
        wb.close()


def test_bybit_chart_folder_repair_moves_and_retargets_once_then_is_verified_noop(tmp_path, monkeypatch):
    workbook, json_path, root, source = _fixture(tmp_path)
    recovery = tmp_path / "recovery"
    monkeypatch.setattr(
        repairer,
        "_select_trade_chart_root",
        lambda prefix: (root, {"selection": "fixture_selected_root", "root": str(root.resolve())}),
    )
    before = _workbook_snapshot(workbook)
    plan = repairer.repair(workbook, json_path, recovery_root=recovery)
    assert plan["applied"] is False
    assert plan["plan"]["action"] == "move_folder_and_retarget_existing_link"
    assert len(plan["plan"]["inventory"]) == 4
    assert source.is_dir()
    assert not (root / "2026" / "OCTOBER" / "C377").exists()

    # A stale, unlocked Excel owner marker must not block the verified repair
    # and is preserved as user-owned temporary metadata.
    owner_file = workbook.with_name("~$" + workbook.name)
    owner_file.write_bytes(b"stale but unlocked owner marker")

    applied = repairer.repair(
        workbook, json_path, recovery_root=recovery,
        apply=True, confirm_plan_sha256=plan["plan_sha256"],
    )
    destination = root / "2026" / "OCTOBER" / "C377"
    assert applied["applied"] is True
    assert applied["verified_workbook"]["row_id"] == ROW_ID
    assert applied["verified_workbook"]["trade_number"] == "C377"
    assert applied["verified_workbook"]["open_time"] == "2026-10-04T13:07:22"
    assert applied["verified_workbook"]["close_time"] == "2026-10-07T12:01:07"
    assert applied["verified_workbook"]["hyperlink"] == destination.resolve().as_uri()
    assert not source.exists()
    assert repairer._inventory(destination) == plan["plan"]["inventory"]
    assert Path(applied["workbook_backup"]).is_file()
    assert Path(applied["chart_folder_backup"]).is_dir()
    assert owner_file.read_bytes() == b"stale but unlocked owner marker"
    after = _workbook_snapshot(workbook)
    assert after == {**before, "hyperlink": destination.resolve().as_uri()}
    assert _TRADE_FOLDER_INDEX_CACHE == {}

    digest_after_move = repairer._sha256(workbook)
    repeat_plan = repairer.repair(workbook, json_path, recovery_root=recovery)
    assert repeat_plan["plan"]["action"] == "verified_no_op"
    repeated = repairer.repair(
        workbook, json_path, recovery_root=recovery,
        apply=True, confirm_plan_sha256=repeat_plan["plan_sha256"],
    )
    assert repeated["verified_no_op"] is True
    assert repeated["verified_workbook"]["hyperlink"] == destination.resolve().as_uri()
    assert repairer._sha256(workbook) == digest_after_move
    assert len(list(recovery.iterdir())) == 1


def test_bybit_chart_folder_destination_conflict_refuses_without_changes(tmp_path):
    workbook, json_path, root, source = _fixture(tmp_path)
    destination = root / "2026" / "OCTOBER" / "C377"
    destination.mkdir(parents=True)
    (destination / "user-file.png").write_bytes(b"do-not-overwrite")
    before_workbook = repairer._sha256(workbook)
    before_source = repairer._inventory(source)
    before_destination = repairer._inventory(destination)
    with pytest.raises(repairer.RepairRefused, match="Both APRIL/C377 and OCTOBER/C377"):
        repairer.repair(workbook, json_path, chart_root=root, recovery_root=tmp_path / "recovery", apply=True, confirm_plan_sha256="unused")
    assert repairer._sha256(workbook) == before_workbook
    assert repairer._inventory(source) == before_source
    assert repairer._inventory(destination) == before_destination
    assert not (tmp_path / "recovery").exists()


def test_bybit_chart_folder_repair_failure_restores_workbook_and_folder(tmp_path, monkeypatch):
    workbook, json_path, root, source = _fixture(tmp_path)
    recovery = tmp_path / "recovery"
    plan = repairer.repair(workbook, json_path, chart_root=root, recovery_root=recovery)
    before_workbook = repairer._sha256(workbook)
    before_inventory = repairer._inventory(source)

    def fail_verification(_workbook, _folder):
        raise repairer.RepairRefused("injected final-verification failure")

    monkeypatch.setattr(repairer, "_verify_saved_workbook", fail_verification)
    with pytest.raises(repairer.RepairRefused, match='"complete": true'):
        repairer.repair(
            workbook, json_path, chart_root=root, recovery_root=recovery,
            apply=True, confirm_plan_sha256=plan["plan_sha256"],
        )
    assert repairer._sha256(workbook) == before_workbook
    assert repairer._inventory(source) == before_inventory
    assert not (root / "2026" / "OCTOBER" / "C377").exists()
